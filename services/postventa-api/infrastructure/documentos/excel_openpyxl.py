# services/postventa-api/infrastructure/documentos/excel_openpyxl.py
"""La plantilla de incidencias en `.xlsx`, con `openpyxl` (F-036, `design.md` §3, §5.2).

**Único** módulo del servicio que importa `openpyxl` (lo vigila
`tests/test_f036_arquitectura.py`), junto a los scripts. Un solo generador
para la plantilla, el Excel de errores y la migración (R55, R62): si cada uno
escribiera la suya, la que vuelve con errores podría no reconocerse.

El libro (§3.1):

- `Incidencias`, primera y activa: la cabecera de R2 en la fila 1, las
  validaciones en las filas 2–1001 (R3, R4), paneles inmovilizados en `A2`,
  filtro en la cabecera y la hoja protegida **sin contraseña** con las celdas de
  datos desbloqueadas salvo `Errores` (§3.3).
- `Instrucciones`: la obra, los textos del YAML (§3.4) y los ejemplos (R5).
- `_catalogos` (`veryHidden`): una columna por lista y un rango con nombre por
  cada una, que es de donde beben los desplegables.
- `_plantilla` (`veryHidden`): los metadatos que reconoce el importador (R6).

**Nada se escribe como fórmula** (R7): toda celda de texto se fuerza a texto,
empiece por lo que empiece, y el XML de las hojas no lleva ningún `<f>`.

**El aspecto** (R121–R125, décima enmienda y enmienda 10 bis; `design.md`
§3.6): colores de la identidad Ruesma de F-035 copiados como constantes,
Calibri en todo, sin imágenes. En «Incidencias», cabecera burdeos (gris acero
en `Errores`); el cuerpo se pinta por **formato condicional**, solo en las
filas *pintadas* (la 2 y las que tienen algo escrito): líneas finas, la
columna `Errores` en gris «no editable» y, solo si ninguna fila lleva errores
(su relleno taparía el rojo), bandas alternas. La celda con error, en rojo
suave con el texto en rojo, como estilo fijo. Una plantilla vacía enseña la
cabecera y una fila. **Sin ajustes de impresión**: este Excel se importa, no
se imprime. En «Instrucciones», la obra en una banda burdeos, el título con
su línea, los párrafos con el alto calculado y la tabla de ejemplos con su
cabecera. Las hojas `veryHidden` no llevan formato. Es **solo aspecto**
(R126): no cambia ni una celda, hoja, validación, protección ni metadato de
los que lee el importador, y la versión sigue en 1.
"""

from __future__ import annotations

import io
import warnings
from collections.abc import Iterator, Mapping
from datetime import UTC, date, datetime, time
from typing import TYPE_CHECKING, Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

try:  # API privada de openpyxl: ver `_filas_presentes`
    from openpyxl.worksheet._reader import WorkSheetParser as _AnalizadorDeFilas
except ImportError:  # pragma: no cover · otra versión de openpyxl: el lector falla cerrado
    _AnalizadorDeFilas = None

from domain.models.errores import FicheroNoEsPlantilla
from domain.models.importacion import CeldaLeida, FilaLeida, LibroLeido
from domain.models.plantilla_incidencias import (
    CABECERA,
    COLUMNA_ERRORES,
    IDENTIFICADOR_PLANTILLA,
    MAX_DESCRIPCION,
    MAX_DETALLE,
    MAX_FILAS,
    VERSION_PLANTILLA,
    CatalogoObra,
    ListasCerradas,
    MensajeColumna,
    OpcionOficio,
    OpcionProveedor,
    TextosPlantilla,
    etiquetas_de_unidades,
)
from domain.ports.hoja_calculo import FilaPlantilla

# Los pasos 1–2 bis (tamaño, ZIP y presupuesto de elementos) viven en el
# lector aislado, que no puede importar este módulo (lleva `openpyxl`); el
# lector en proceso de aquí los usa igual. Los topes se reexportan porque son
# los de este lector también.
from infrastructure.documentos.lector_aislado import (  # noqa: F401 · reexportados
    MAX_BYTES_DESCOMPRIMIDOS,
    PRESUPUESTO_ELEMENTOS_XML,
    _comprobar_tamano,
    _contar_elementos_xml,
    _inspeccionar_zip,
    _no_es_xlsx,
    error_a_json,
    libro_a_json,
)

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

HOJA_INCIDENCIAS = "Incidencias"
HOJA_INSTRUCCIONES = "Instrucciones"
HOJA_CATALOGOS = "_catalogos"
HOJA_METADATOS = "_plantilla"

#: Filas de datos con validación (R3): la cabecera es la 1 y caben MAX_FILAS.
PRIMERA_FILA = 2
ULTIMA_FILA = PRIMERA_FILA + MAX_FILAS - 1

#: Hasta dónde se leen los metadatos de `_plantilla` (R115): las claves son
#: cinco; el tope deja margen sin recorrer la hoja entera.
MAX_FILAS_METADATOS = 10

#: Columnas que se escriben: todas las de la cabecera menos la de errores.
COLUMNAS_DE_DATOS: tuple[str, ...] = tuple(c for c in CABECERA if c != COLUMNA_ERRORES)

#: Las columnas con desplegable y el rango con nombre de cada lista (R3).
LISTAS: Mapping[str, str] = {
    "Unidad": "lista_unidades",
    "Ubicación": "lista_ubicaciones",
    "Oficio": "lista_oficios",
    "Proveedor": "lista_proveedores",
    "Urgencia": "lista_urgencias",
    "Listado": "lista_listados",
}

#: Longitudes de los textos libres (R4).
LONGITUDES: Mapping[str, tuple[int, int]] = {
    "Descripción corta": (1, MAX_DESCRIPCION),
    "Detalle": (0, MAX_DETALLE),
}

#: Anchos de columna de «Incidencias» (R122); los textos largos, además,
#: ajustados. Eran 30, 24, 50, 60, 28, 45, 24, 18 y 60: son los de la muestra
#: que aprobó el humano.
ANCHOS: Mapping[str, int] = {
    "Unidad": 30,
    "Ubicación": 26,
    "Descripción corta": 48,
    "Detalle": 56,
    "Oficio": 28,
    "Proveedor": 40,
    "Urgencia": 24,
    "Listado": 16,
    COLUMNA_ERRORES: 50,
}
AJUSTADAS = frozenset({"Descripción corta", "Detalle", COLUMNA_ERRORES})

#: Formato «texto» de Excel: lo que se teclea no se convierte en fecha,
#: número ni fórmula.
FORMATO_TEXTO = "@"

# --- El aspecto (R121–R125; `design.md` §3.6) ---------------------------------
# Colores de la identidad Ruesma de F-035 (`services/postventa-front/css/
# styles.css` de `feature/F-035-portal-posventa`, bloque `:root`), como ARGB
# opaco. Copiados a mano: el generador es del backend y no lee una hoja de
# estilos del front. Si la identidad cambia, se cambian aquí.
BURDEOS = "FF9F2842"  # --rs-burdeos
BURDEOS_FUERTE = "FF7A1E33"  # --rs-burdeos-fuerte
ACERO = "FF7B868C"  # --rs-acero
ACERO_100 = "FFDFE2E4"  # --rs-acero-100
ACERO_TEXTO = "FF5D676D"  # --rs-acero-texto
TINTA = "FF1D2024"  # --rs-tinta
LIENZO = "FFF3F4F5"  # --rs-lienzo
BLANCO = "FFFFFFFF"  # --rs-papel
ERROR_SUAVE = "FFFEF2F2"  # --rs-error-suave
ERROR = "FFB91C1C"  # --rs-error
BANDA = "FFFAFAFB"  # propio: entre --rs-papel y --rs-lienzo

#: Las pestañas van en RGB de seis cifras, sin el canal alfa.
PESTANA_INCIDENCIAS = BURDEOS[2:]
PESTANA_INSTRUCCIONES = ACERO[2:]

#: Las fuentes de la marca no están en los equipos de quien abre el fichero.
FUENTE = "Calibri"
PUNTOS_CABECERA = 11
PUNTOS_CUERPO = 10.5
PUNTOS_OBRA = 15
PUNTOS_TITULO = 13
PUNTOS_PARRAFO = 11
SANGRIA = 1

ALTO_CABECERA = 34
ALTO_OBRA = 38
ALTO_TITULO = 28
ALTO_CABECERA_EJEMPLOS = 26

#: El alto de un párrafo de «Instrucciones»: Excel no ajusta solo el alto de una
#: fila escrita por programa. 105 caracteres caben en una línea de la columna A
#: (ancho 100, Calibri 11), 16 puntos por línea y 6 de respiro.
ALTO_MINIMO_PARRAFO = 20
PUNTOS_POR_LINEA = 16
CARACTERES_POR_LINEA = 105
RESPIRO_PARRAFO = 6

ANCHO_INSTRUCCIONES = 100
ANCHO_EJEMPLOS = 28

#: La celda con error (R64, R123): era el naranja `FFF4B084`. El burdeos es
#: marca y nunca un estado: el error usa los tokens de error.
RELLENO_ERROR = PatternFill(fill_type="solid", fgColor=ERROR_SUAVE)
FUENTE_ERROR = Font(name=FUENTE, size=PUNTOS_CUERPO, color=ERROR)
AUTOR_COMENTARIO = "Posventa"

#: Fila *pintada* (R122, enmienda 10 bis): la 2, o una con algo escrito de `A`
#: a `I`. Relativa a la primera fila de cada rango, con las columnas fijas: vale
#: igual para `A2:I1001` que para `I2:I1001`. Por formato condicional y no por
#: estilo de celda, para que siga bien tras «ordenar» y tras insertar o borrar
#: filas, y para que Excel no dé por usadas las 1.000 filas vacías.
FORMULA_PINTADA = "OR(ROW()=2,COUNTA($A2:$I2)>0)"

#: Bandas alternas (R122): tiñe las filas impares pintadas (la 2 es par).
FORMULA_BANDAS = "AND(MOD(ROW(),2)=1,COUNTA($A2:$I2)>0)"


def _letra(columna: str) -> str:
    return get_column_letter(CABECERA.index(columna) + 1)


def _alto_de_parrafo(texto: str) -> int:
    """Los puntos de alto de la fila de un párrafo, para que no se vea cortado (R124)."""
    lineas = len(texto) // CARACTERES_POR_LINEA + 1
    return max(ALTO_MINIMO_PARRAFO, PUNTOS_POR_LINEA * lineas + RESPIRO_PARRAFO)


def _solido(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _borde_de_cabecera() -> Border:
    """Laterales finos blancos y una línea inferior media burdeos fuerte (R121)."""
    lateral = Side(style="thin", color=BLANCO)
    return Border(
        left=lateral, right=lateral, bottom=Side(style="medium", color=BURDEOS_FUERTE)
    )


def _borde_fino() -> Border:
    """Los cuatro lados finos en gris claro: el cuerpo y los ejemplos (R122, R124)."""
    linea = Side(style="thin", color=ACERO_100)
    return Border(left=linea, right=linea, top=linea, bottom=linea)


def _texto(celda: Cell, valor: str | None) -> None:
    """Escribe un texto **como texto**, aunque empiece por `=`, `+`, `-` o `@` (R7).

    `openpyxl` toma por fórmula cualquier cadena que empiece por `=`; se le
    devuelve el tipo cadena y se marca con el prefijo de comilla, que es lo
    que hace Excel cuando alguien teclea `'=…`.
    """
    celda.value = valor
    if valor is not None:
        celda.data_type = "s"
        celda.quotePrefix = True


def _utc(instante: datetime) -> str:
    if instante.utcoffset() is None:
        raise ValueError("el instante de generación tiene que llevar zona horaria")
    return instante.astimezone(UTC).isoformat(timespec="seconds")


def _comprobar_filas(filas: tuple[FilaPlantilla, ...]) -> None:
    if len(filas) > MAX_FILAS:
        raise ValueError(f"no caben más de {MAX_FILAS} filas en la plantilla")
    for fila in filas:
        raras = (set(fila.valores) | set(fila.errores)) - set(COLUMNAS_DE_DATOS)
        if raras:
            # Sin nombrar el valor: puede ser texto de la propiedad y acabar en un log.
            raise ValueError("una fila trae una columna que no es de la plantilla")


class GeneradorPlantillaOpenpyxl:
    """Implementa `GeneradorPlantillaPort` (§5.2)."""

    def generar(
        self,
        *,
        catalogo: CatalogoObra,
        opciones_oficio: tuple[OpcionOficio, ...],
        opciones_proveedor: tuple[OpcionProveedor, ...],
        listas: ListasCerradas,
        textos: TextosPlantilla,
        generada_at: datetime,
        filas: tuple[FilaPlantilla, ...] = (),
        importacion_origen: UUID | None = None,
    ) -> bytes:
        generada_utc = _utc(generada_at)
        _comprobar_filas(filas)
        libro = Workbook()
        incidencias = libro.active
        incidencias.title = HOJA_INCIDENCIAS
        instrucciones = libro.create_sheet(HOJA_INSTRUCCIONES)
        catalogos = libro.create_sheet(HOJA_CATALOGOS)
        metadatos = libro.create_sheet(HOJA_METADATOS)

        valores_de_listas = {
            "Unidad": tuple(
                o.etiqueta for o in etiquetas_de_unidades(catalogo.unidades)
            ),
            "Ubicación": listas.ubicaciones,
            "Oficio": tuple(o.etiqueta for o in opciones_oficio),
            "Proveedor": tuple(p.etiqueta for p in opciones_proveedor),
            "Urgencia": tuple(o.etiqueta for o in listas.urgencias),
            "Listado": tuple(o.etiqueta for o in listas.listados),
        }
        self._catalogos(libro, catalogos, valores_de_listas)
        self._incidencias(incidencias, valores_de_listas, textos.mensajes, filas)
        self._instrucciones(
            instrucciones,
            catalogo,
            textos,
            sin_oficios=not opciones_oficio,
            es_excel_de_errores=importacion_origen is not None,
        )
        self._metadatos(
            metadatos, catalogo.obra_codigo, generada_utc, importacion_origen
        )

        libro.active = 0
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    # --- hojas ----------------------------------------------------------------

    @staticmethod
    def _catalogos(
        libro: Workbook, ws: Worksheet, valores_de_listas: Mapping[str, tuple[str, ...]]
    ) -> None:
        """Una columna por lista, con su nombre en la fila 1, y un rango con nombre."""
        for indice, (columna, nombre) in enumerate(LISTAS.items(), start=1):
            letra = get_column_letter(indice)
            _texto(ws.cell(row=1, column=indice), nombre)
            valores = valores_de_listas[columna]
            for fila, valor in enumerate(valores, start=2):
                _texto(ws.cell(row=fila, column=indice), valor)
            if valores:
                referencia = (
                    f"'{HOJA_CATALOGOS}'!${letra}$2:${letra}${len(valores) + 1}"
                )
                libro.defined_names[nombre] = DefinedName(nombre, attr_text=referencia)
        ws.sheet_state = "veryHidden"
        ws.protection.sheet = True

    @staticmethod
    def _incidencias(
        ws: Worksheet,
        valores_de_listas: Mapping[str, tuple[str, ...]],
        mensajes: Mapping[str, MensajeColumna],
        filas: tuple[FilaPlantilla, ...],
    ) -> None:
        # Cabecera (R121): los objetos de estilo se crean una vez y se comparten.
        fuente_cabecera = Font(name=FUENTE, size=PUNTOS_CABECERA, bold=True, color=BLANCO)
        alineacion_cabecera = Alignment(
            horizontal="left", vertical="center", wrap_text=True, indent=SANGRIA
        )
        borde_cabecera = _borde_de_cabecera()
        relleno_cabecera = _solido(BURDEOS)
        relleno_cabecera_errores = _solido(ACERO)
        for indice, columna in enumerate(CABECERA, start=1):
            celda = ws.cell(row=1, column=indice)
            _texto(celda, columna)
            celda.font = fuente_cabecera
            celda.fill = (
                relleno_cabecera_errores if columna == COLUMNA_ERRORES else relleno_cabecera
            )
            celda.alignment = alineacion_cabecera
            celda.border = borde_cabecera
            ws.column_dimensions[get_column_letter(indice)].width = ANCHOS[columna]
        ws.row_dimensions[1].height = ALTO_CABECERA

        # Cuerpo (R122) y columna `Errores` (R123): como estilo fijo, solo lo
        # que no se ve en una celda vacía (fuente y alineación) y lo que no es
        # aspecto (desbloqueo y `@`, R126). Las líneas, el gris y las bandas
        # van por formato condicional, más abajo (enmienda 10 bis).
        desbloqueada = Protection(locked=False)
        fuente_cuerpo = Font(name=FUENTE, size=PUNTOS_CUERPO, color=TINTA)
        fuente_errores = Font(name=FUENTE, size=PUNTOS_CUERPO, color=ACERO_TEXTO)
        arriba = Alignment(vertical="top", indent=SANGRIA)
        ajuste = Alignment(vertical="top", wrap_text=True, indent=SANGRIA)
        for fila in range(PRIMERA_FILA, ULTIMA_FILA + 1):
            for indice, columna in enumerate(CABECERA, start=1):
                celda = ws.cell(row=fila, column=indice)
                celda.alignment = ajuste if columna in AJUSTADAS else arriba
                if columna == COLUMNA_ERRORES:
                    celda.font = fuente_errores
                else:
                    celda.font = fuente_cuerpo
                    celda.protection = desbloqueada
                    celda.number_format = FORMATO_TEXTO

        for numero, fila in enumerate(filas, start=PRIMERA_FILA):
            for columna in COLUMNAS_DE_DATOS:
                _texto(ws[f"{_letra(columna)}{numero}"], fila.valores.get(columna))
            _texto(ws[f"{_letra(COLUMNA_ERRORES)}{numero}"], fila.texto_errores)
            for columna, problema in fila.errores.items():
                celda = ws[f"{_letra(columna)}{numero}"]
                celda.fill = RELLENO_ERROR
                celda.font = FUENTE_ERROR
                comentario = Comment(problema, AUTOR_COMENTARIO)
                comentario.width = 300
                comentario.height = 120
                celda.comment = comentario

        for columna in COLUMNAS_DE_DATOS:
            ws.add_data_validation(
                _validacion(columna, valores_de_listas, mensajes[columna])
            )

        ultima = get_column_letter(len(CABECERA))
        ws.freeze_panes = f"A{PRIMERA_FILA}"
        ws.auto_filter.ref = f"A1:{ultima}{ULTIMA_FILA}"
        proteccion = ws.protection
        proteccion.sheet = True
        # En OOXML, `False` es «no protegido»: se permite (§3.3).
        proteccion.insertRows = False
        proteccion.deleteRows = False
        proteccion.sort = False
        proteccion.autoFilter = False

        # El cuerpo a la vista (R122, R123; enmienda 10 bis): tres reglas de un
        # solo rango cada una, en este orden (`openpyxl` numera la prioridad
        # por orden de alta: 1, 2, 3), sin «detener si es verdad». Sus formatos
        # no se pisan: bordes en A–I, gris solo en `Errores`, bandas solo en
        # las columnas de datos.
        errores = _letra(COLUMNA_ERRORES)
        # 1 · líneas finas en las filas pintadas.
        ws.conditional_formatting.add(
            f"A{PRIMERA_FILA}:{ultima}{ULTIMA_FILA}",
            FormulaRule(formula=[FORMULA_PINTADA], border=_borde_fino()),
        )
        # 2 · el gris «no editable» de `Errores` en las filas pintadas.
        ws.conditional_formatting.add(
            f"{errores}{PRIMERA_FILA}:{errores}{ULTIMA_FILA}",
            FormulaRule(
                formula=[FORMULA_PINTADA],
                fill=PatternFill(fill_type="solid", bgColor=LIENZO),
            ),
        )
        # 3 · bandas alternas, solo si ninguna fila lleva errores: el relleno
        # de la regla se pinta encima del fijo y taparía el rojo de la celda
        # con error en las filas impares.
        if not any(fila.errores for fila in filas):
            ultima_de_datos = get_column_letter(len(COLUMNAS_DE_DATOS))
            ws.conditional_formatting.add(
                f"A{PRIMERA_FILA}:{ultima_de_datos}{ULTIMA_FILA}",
                FormulaRule(
                    formula=[FORMULA_BANDAS],
                    fill=PatternFill(fill_type="solid", bgColor=BANDA),
                ),
            )
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = PESTANA_INCIDENCIAS
        # Sin ajustes de impresión (R125, enmienda 10 bis): no se imprime.

    @staticmethod
    def _instrucciones(
        ws: Worksheet,
        catalogo: CatalogoObra,
        textos: TextosPlantilla,
        *,
        sin_oficios: bool,
        es_excel_de_errores: bool,
    ) -> None:
        obra = catalogo.obra_codigo
        if catalogo.obra_nombre:
            obra = f"{obra} · {catalogo.obra_nombre}"
        parrafos: list[str] = []
        if es_excel_de_errores:
            parrafos.append(textos.excel_errores)
        parrafos.extend(textos.instrucciones)
        if sin_oficios:
            parrafos.append(textos.sin_oficios)

        # «Las columnas de la tabla» son las 8 de datos (A–H). En las filas 1 y
        # 2 las celdas B–H solo llevan estilo, sin valor (R124).
        columnas = range(1, len(COLUMNAS_DE_DATOS) + 1)
        centrada = Alignment(vertical="center", indent=SANGRIA)
        ajuste = Alignment(wrap_text=True, vertical="top", indent=SANGRIA)

        # Fila 1: la obra, en una banda burdeos.
        _texto(ws.cell(row=1, column=1), obra)
        fuente_obra = Font(name=FUENTE, size=PUNTOS_OBRA, bold=True, color=BLANCO)
        relleno_burdeos = _solido(BURDEOS)
        for indice in columnas:
            celda = ws.cell(row=1, column=indice)
            celda.fill = relleno_burdeos
            celda.font = fuente_obra
            celda.alignment = centrada
        ws.row_dimensions[1].height = ALTO_OBRA

        # Fila 2: el título, en burdeos, con una línea debajo.
        titulo = ws.cell(row=2, column=1)
        _texto(titulo, textos.titulo)
        titulo.font = Font(name=FUENTE, size=PUNTOS_TITULO, bold=True, color=BURDEOS)
        titulo.alignment = centrada
        linea = Border(bottom=Side(style="medium", color=BURDEOS))
        for indice in columnas:
            ws.cell(row=2, column=indice).border = linea
        ws.row_dimensions[2].height = ALTO_TITULO

        # Los párrafos, con el alto calculado por su largo.
        fuente_parrafo = Font(name=FUENTE, size=PUNTOS_PARRAFO, color=TINTA)
        fila = 3
        for parrafo in parrafos:
            celda = ws.cell(row=fila, column=1)
            _texto(celda, parrafo)
            celda.font = fuente_parrafo
            celda.alignment = ajuste
            ws.row_dimensions[fila].height = _alto_de_parrafo(parrafo)
            fila += 1

        # Una fila en blanco y la tabla de ejemplos, con su cabecera como la
        # de «Incidencias» y la segunda, cuarta… fila en banda (relleno fijo:
        # esta tabla no se ordena).
        fila += 1
        fuente_cabecera = Font(name=FUENTE, size=PUNTOS_CABECERA, bold=True, color=BLANCO)
        borde_cabecera = _borde_de_cabecera()
        for indice, columna in enumerate(COLUMNAS_DE_DATOS, start=1):
            celda = ws.cell(row=fila, column=indice)
            _texto(celda, columna)
            celda.fill = relleno_burdeos
            celda.font = fuente_cabecera
            celda.border = borde_cabecera
            celda.alignment = centrada
        ws.row_dimensions[fila].height = ALTO_CABECERA_EJEMPLOS
        fuente_ejemplo = Font(name=FUENTE, size=PUNTOS_CUERPO, color=TINTA)
        borde_ejemplo = _borde_fino()
        relleno_banda = _solido(BANDA)
        for numero, ejemplo in enumerate(textos.ejemplos, start=1):
            fila += 1
            for indice, columna in enumerate(COLUMNAS_DE_DATOS, start=1):
                celda = ws.cell(row=fila, column=indice)
                _texto(celda, ejemplo.get(columna))
                celda.font = fuente_ejemplo
                celda.border = borde_ejemplo
                celda.alignment = ajuste
                if numero % 2 == 0:
                    celda.fill = relleno_banda

        ws.column_dimensions["A"].width = ANCHO_INSTRUCCIONES
        for indice in range(2, len(COLUMNAS_DE_DATOS) + 1):
            ws.column_dimensions[get_column_letter(indice)].width = ANCHO_EJEMPLOS
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = PESTANA_INSTRUCCIONES

    @staticmethod
    def _metadatos(
        ws: Worksheet,
        obra_codigo: str,
        generada_utc: str,
        importacion_origen: UUID | None,
    ) -> None:
        pares: list[tuple[str, str | int]] = [
            ("identificador", IDENTIFICADOR_PLANTILLA),
            ("version", VERSION_PLANTILLA),
            ("obra_codigo", obra_codigo),
            ("generada_at_utc", generada_utc),
        ]
        if importacion_origen is not None:
            pares.append(("importacion_origen", str(importacion_origen)))
        for fila, (clave, valor) in enumerate(pares, start=1):
            _texto(ws.cell(row=fila, column=1), clave)
            celda = ws.cell(row=fila, column=2)
            if isinstance(valor, int):
                celda.value = valor
            else:
                _texto(celda, valor)
        ws.sheet_state = "veryHidden"
        ws.protection.sheet = True


def _validacion(
    columna: str,
    valores_de_listas: Mapping[str, tuple[str, ...]],
    mensaje: MensajeColumna,
) -> DataValidation:
    """La validación de datos de una columna, en estilo «detener» (R3, R4).

    Una lista vacía (obra sin oficios, R12, o sin proveedores en `obrofc`) no
    puede ser un desplegable: la columna solo admite quedarse en blanco.
    """
    if columna in LONGITUDES:
        minimo, maximo = LONGITUDES[columna]
        dv = DataValidation(
            type="textLength",
            operator="between",
            formula1=str(minimo),
            formula2=str(maximo),
        )
    elif valores_de_listas[columna]:
        dv = DataValidation(type="list", formula1=LISTAS[columna])
    else:
        dv = DataValidation(type="textLength", operator="equal", formula1="0")
    dv.allow_blank = True
    dv.errorStyle = "stop"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    dv.promptTitle = mensaje.titulo_entrada
    dv.prompt = mensaje.entrada
    dv.errorTitle = mensaje.titulo_error
    dv.error = mensaje.error
    letra = _letra(columna)
    dv.add(f"{letra}{PRIMERA_FILA}:{letra}{ULTIMA_FILA}")
    return dv


# --------------------------------------------------------------------------
# El lector (T11, §5.2)
# --------------------------------------------------------------------------

#: Forma del formato antiguo (R17): texto en A1, A2:A10 vacías, texto en D1 y E1.
FILAS_VACIAS_BAJO_A1 = range(2, 11)


def _relanzar_si_falta_memoria(error: BaseException) -> None:
    """Un `MemoryError` que `openpyxl` ha tapado con otro error sigue siendo memoria.

    Sus descriptores convierten cualquier fallo al construir un atributo en
    `TypeError` (`descriptors.base._convert`, con un `except:` desnudo), y el
    `MemoryError` queda en la cadena (`__context__`). Con el tope de memoria del
    hijo es justo lo que pasa con un `sqref` de millones de rangos (R4-1, T41):
    sin esto sería `no_es_xlsx` en vez de memoria (R118).
    """
    vistos: set[int] = set()
    actual: BaseException | None = error
    while actual is not None and id(actual) not in vistos:
        if isinstance(actual, MemoryError):
            raise MemoryError from error  # noqa: TRY004 · es memoria, no un tipo malo
        vistos.add(id(actual))
        actual = actual.__cause__ or actual.__context__


def _abrir(contenido: bytes) -> Workbook:
    """Paso 3: `openpyxl` **en solo lectura**, con `defusedxml` debajo (§5.2).

    En solo lectura `load_workbook` no enlaza las hojas al cargar: ni crea una
    celda por posición de un rango combinado o de un hipervínculo sobre un
    rango, en ninguna hoja, ni lee las filas hasta que se recorren (R115,
    «Al cargar»). Por eso quien lo llama cierra el libro siempre, y un fallo al
    recorrer una hoja también es `no_es_xlsx` (`LectorPlantillaOpenpyxl.leer`).
    `keep_vba` se queda en su valor por defecto (False): un libro con macros ya
    se ha rechazado en el paso 2. Un `MemoryError` no es «no es un Excel»: sale
    tal cual, para que el hijo acabe con su código de memoria (R118).
    """
    try:
        return load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    except MemoryError:
        raise
    except Exception as error:  # noqa: BLE001 · XML hostil, partes que faltan, tipos raros…
        _relanzar_si_falta_memoria(error)
        raise _no_es_xlsx() from None


def _es_texto(valor: object) -> bool:
    return isinstance(valor, str) and bool(valor.strip())


def _parece_formato_antiguo(ws: ReadOnlyWorksheet) -> bool:
    """R17: la forma del Excel de antes (`docs/referencia/05_…`).

    Sin cabecera reconocible (ninguna celda de la fila 1 es una columna de la
    plantilla), texto en `A1`, `A2:A10` vacías y texto en `D1` y `E1`. En
    **una** pasada por las celdas fijas `A1:I10` (R115): en solo lectura cada
    `ws.cell` volvería a leer la hoja.
    """
    recorrido = ws.iter_rows(
        min_row=1,
        max_row=FILAS_VACIAS_BAJO_A1[-1],
        max_col=len(CABECERA),
        values_only=True,
    )
    # Si la hoja acaba antes de la fila 10, las que faltan están vacías.
    filas = dict(enumerate(recorrido, start=1))
    vacia = (None,) * len(CABECERA)
    primera = filas.get(1, vacia)
    if any(isinstance(v, str) and v.strip() in CABECERA for v in primera):
        return False
    return (
        _es_texto(primera[0])
        and all(filas.get(f, vacia)[0] is None for f in FILAS_VACIAS_BAJO_A1)
        and _es_texto(primera[3])
        and _es_texto(primera[4])
    )


def _entero(valor: object) -> int | None:
    """La versión solo vale como número entero: ni texto, ni booleano, ni decimal."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return None


def _metadatos(ws: ReadOnlyWorksheet) -> tuple[str | None, int | None, str | None]:
    """Paso 4: pares clave/valor de `_plantilla`; con clave repetida, la primera.

    Solo las dos primeras columnas y hasta `MAX_FILAS_METADATOS` filas (R115).
    """
    pares: dict[str, object] = {}
    recorrido = ws.iter_rows(
        min_row=1, max_row=MAX_FILAS_METADATOS, min_col=1, max_col=2, values_only=True
    )
    for clave, valor in recorrido:
        if isinstance(clave, str):
            pares.setdefault(clave, valor)
    identificador = pares.get("identificador")
    obra = pares.get("obra_codigo")
    return (
        identificador if isinstance(identificador, str) else None,
        _entero(pares.get("version")),
        obra if isinstance(obra, str) else None,
    )


def _numero_como_texto(valor: float) -> str:
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _fecha_como_texto(valor: date | time) -> str:
    if isinstance(valor, datetime):
        if valor.time() == time(0, 0):
            return valor.date().isoformat()
        return valor.isoformat(sep=" ")
    return valor.isoformat()


def _celda_leida(valor: object, tipo: str) -> CeldaLeida | None:
    """Paso 5: una celda, por su valor y su tipo, con el texto que se reescribiría (R32, R63)."""
    if valor is None:
        return None
    if tipo == "f":
        texto = (
            valor if isinstance(valor, str) else (getattr(valor, "text", None) or "=")
        )
        return CeldaLeida(texto, True, False, texto)
    if isinstance(valor, bool):
        return CeldaLeida(valor, False, True, "VERDADERO" if valor else "FALSO")
    if isinstance(valor, (date, time)):
        texto = _fecha_como_texto(valor)
        return CeldaLeida(texto, False, True, texto)
    if isinstance(valor, (int, float)):
        return CeldaLeida(valor, False, False, _numero_como_texto(valor))
    texto = str(valor)
    return CeldaLeida(texto, False, False, texto)


def _con_dato(celda: CeldaLeida) -> bool:
    """Lo mismo que el dominio cuenta como no vacía (R23)."""
    return celda.es_formula or not (
        isinstance(celda.valor, str) and not celda.valor.strip()
    )


def _filas_presentes(ws: object) -> Iterator[tuple[int, list[dict[str, Any]]]]:
    """Solo las filas que el XML de la hoja trae, sin rellenar huecos (R115).

    En solo lectura, `iter_rows` sin `max_row` rellena con una fila vacía cada
    número de fila que falta hasta la última: hasta la 1.048.576 es un millón
    de iteraciones. Esto es el analizador que la hoja de solo lectura usa por
    dentro (`WorkSheetParser` sobre `ws._get_source()`, con los mismos
    parámetros que `ReadOnlyWorksheet._cells_by_row`), que da cada fila del XML
    con sus celdas (`column`, `value` ya convertido y `data_type`). Es API
    privada de `openpyxl` (`requirements.txt`: `>=3.1,<4.0`); un test fija su
    forma, y si no está el lector falla cerrado en vez de recorrer la hoja.
    """
    if _AnalizadorDeFilas is None:
        raise _no_es_xlsx()
    try:
        libro = ws.parent  # type: ignore[attr-defined]
        fuente = ws._get_source()  # type: ignore[attr-defined]
        analizador = _AnalizadorDeFilas(
            fuente,
            ws._shared_strings,  # type: ignore[attr-defined]
            data_only=libro.data_only,
            epoch=libro.epoch,
            date_formats=libro._date_formats,
            timedelta_formats=libro._timedelta_formats,
        )
    except AttributeError:
        raise _no_es_xlsx() from None
    with fuente:
        yield from analizador.parse()


def _datos_fuera_del_tope(ws: ReadOnlyWorksheet) -> bool:
    """R115: algún dato en las columnas de datos de la fila 1002 o más abajo.

    Mira solo las filas que el XML trae (`_filas_presentes`); una celda vacía o
    con solo espacios, aunque lleve estilo, no cuenta (el mismo criterio que
    R23 dentro de la plantilla).
    """
    for fila, celdas in _filas_presentes(ws):
        if fila <= ULTIMA_FILA:
            continue
        for celda in celdas:
            if celda["column"] <= len(COLUMNAS_DE_DATOS):
                leida = _celda_leida(celda["value"], celda["data_type"])
                if leida is not None and _con_dato(leida):
                    return True
    return False


def _filas(ws: ReadOnlyWorksheet) -> tuple[FilaLeida, ...]:
    """Las filas con algún dato fuera de `Errores`, que ni se lee (R23, R66).

    De `PRIMERA_FILA` a `ULTIMA_FILA + 1`, nunca hasta `ws.max_row` (R115).
    """
    filas: list[FilaLeida] = []
    recorrido = ws.iter_rows(
        min_row=PRIMERA_FILA, max_row=ULTIMA_FILA + 1, max_col=len(COLUMNAS_DE_DATOS)
    )
    for numero, fila in enumerate(recorrido, start=PRIMERA_FILA):
        celdas: dict[str, CeldaLeida] = {}
        for indice, columna in enumerate(COLUMNAS_DE_DATOS):
            leida = _celda_leida(fila[indice].value, fila[indice].data_type)
            if leida is not None:
                celdas[columna] = leida
        if any(_con_dato(c) for c in celdas.values()):
            filas.append(FilaLeida(numero=numero, celdas=celdas))
    return tuple(filas)


def _cabecera(ws: ReadOnlyWorksheet) -> tuple[str, ...]:
    """La fila 1, en una pasada: las columnas de la cabecera y lo que haya más allá.

    Solo las celdas que la fila trae (R115). Más a la derecha de la cabecera,
    las que traen valor se añaden en orden de columna, para que R19 diga que
    sobran; una celda con estilo y sin valor no añade nada.
    """
    ancho = len(CABECERA)
    leida = [""] * ancho
    lejanas: list[tuple[int, object]] = []
    for numero, celdas in _filas_presentes(ws):
        if numero == 1:
            for celda in celdas:
                columna, valor = celda["column"], celda["value"]
                if valor is None:
                    continue
                if columna <= ancho:
                    leida[columna - 1] = str(valor)
                else:
                    lejanas.append((columna, valor))
        break  # la fila 1, si está, es la primera del XML
    lejanas.sort(key=lambda par: par[0])
    return (*leida, *(str(valor) for _, valor in lejanas))


def _extraer(libro: Workbook) -> LibroLeido:
    """Pasos 4 y 5, sobre el libro ya abierto en solo lectura."""
    if HOJA_METADATOS in libro.sheetnames:
        identificador, version, obra = _metadatos(libro[HOJA_METADATOS])
        antiguo = False
    else:
        identificador = version = obra = None
        antiguo = _parece_formato_antiguo(libro.worksheets[0])
    if HOJA_INCIDENCIAS in libro.sheetnames:
        incidencias = libro[HOJA_INCIDENCIAS]
        fuera = _datos_fuera_del_tope(incidencias)
        cabecera, filas = _cabecera(incidencias), _filas(incidencias)
    else:
        cabecera, filas, fuera = (), (), False
    return LibroLeido(
        identificador=identificador,
        version=version,
        obra_codigo=obra,
        cabecera=cabecera,
        filas=filas,
        parece_formato_antiguo=antiguo,
        datos_fuera_del_tope=fuera,
    )


def _leer_libro(contenido: bytes) -> LibroLeido:
    """Pasos 3–5: abrir en solo lectura, extraer y cerrar, siempre.

    Lo que devuelve son objetos del dominio, ya leídos: el libro se cierra
    después de extraerlo y también si la lectura falla. Un fallo al recorrer
    una hoja es `no_es_xlsx`; un `MemoryError`, no (R118).
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        libro = _abrir(contenido)
        try:
            return _extraer(libro)
        except (FicheroNoEsPlantilla, MemoryError):
            raise
        except Exception as error:  # noqa: BLE001 · en solo lectura el XML de las hojas se lee aquí
            _relanzar_si_falta_memoria(error)
            raise _no_es_xlsx() from None
        finally:
            libro.close()


def leer_en_hijo(contenido: bytes) -> bytes:
    """La lectura que corre en el **proceso hijo** (R118; `lector_aislado.py`).

    Los pasos 3–5 sobre un fichero que ya ha pasado los baratos en el padre, y
    el resultado en JSON: el `LibroLeido` o, si el lector lo rechaza, el error
    con su código y su motivo. Un `MemoryError` no se captura: el hijo sale con
    su código de memoria.
    """
    try:
        return libro_a_json(_leer_libro(contenido))
    except FicheroNoEsPlantilla as error:
        return error_a_json(error)


class LectorPlantillaOpenpyxl:
    """Implementa `LectorPlantillaPort` (§5.2) **en este proceso**.

    En este orden, sin abrir `openpyxl` hasta el paso 3: tamaño (R15), firma y
    ZIP (R16), presupuesto de elementos XML (R117), apertura **en solo
    lectura** con `defusedxml`, metadatos o forma
    antigua (R17), y cabecera y filas de «Incidencias». **No valida** reglas de
    negocio: eso es `reconocer_plantilla` y `validar_filas`. Ni la carga crea
    celdas por posición (solo lectura) ni ningún recorrido va posición a
    posición más allá de lo que la plantilla puede tener (R115): lo de más
    abajo o más a la derecha se mira solo entre las celdas que el fichero trae.

    Desde la octava enmienda, un fichero **subido** lo lee
    `LectorPlantillaAislado` (en un proceso hijo, R118); este lo sigue usando
    el script de migración, que lee en local ficheros conocidos.
    """

    def leer(self, *, contenido: bytes) -> LibroLeido:
        _comprobar_tamano(contenido)
        _inspeccionar_zip(contenido)
        _contar_elementos_xml(contenido)
        return _leer_libro(contenido)

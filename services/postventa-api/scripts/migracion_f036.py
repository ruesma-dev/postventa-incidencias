# services/postventa-api/scripts/migracion_f036.py
"""La migración del Excel actual de incidencias a la plantilla (F-036, T23).

`specs/F-036-importar-excel/design.md` §10. Lógica **sin disco**: todo entra y
sale en memoria (bytes, el YAML ya cargado, los JSON ya cargados), y la línea
de órdenes (`scripts/migrar_excel_f036.py`) solo lee, escribe e imprime.

En el orden de §10.3, parando en el primer fallo:

1. **La tabla de correcciones** (§10.2): su forma, y lo que el script hace
   cumplir de R57 (ni «peligro de seguridad» ni remisiones a otras filas en
   descripción y detalle).
2. **El original** (`leer_original`), leído con `openpyxl` en modo lectura: el
   lector de la plantilla lo rechazaría, a propósito. Se coteja con la tabla
   (R54): cada fila una vez, en su orden, con su ubicación y su texto; una
   fila que cambia dice por qué (`cambios`).
3. **El catálogo** del JSON de T2 se compone con **el mismo** código que el de
   Sigrid (`leer_catalogo` sobre un puerto que lee el JSON) y los grupos del
   fichero de R98 con `opciones_de_la_obra` (§7.2) sobre un puerto que da sus
   decisiones: las mismas opciones que la plantilla y la importación. Cada
   oficio de la tabla es un **nombre exacto** de Sigrid (R97), que se lleva a
   su código y el código a su grupo, cuya etiqueta es la que se escribe; el
   proveedor, según R91. *(Undécima enmienda, §10.5.)* Si el nombre no es de
   la obra pero sí de `oficios_catalogo` (todo `auxofc`), el oficio ha salido
   de la obra en Sigrid: la fila va con `Oficio` y `Proveedor` vacíos, se
   cuenta y se lista en el informe, y no se para (R128, R130); solo para la
   errata, que no es de ningún oficio de Sigrid (R97). El proveedor sale
   **solo** de R91 con el catálogo recibido (R129).
4. **El v2** sale del generador de la plantilla (R55) y vuelve por el lector,
   `reconocer_plantilla` y `validar_filas`: cero errores o no hay v2 (R56).

`informe` escribe el Markdown de revisión de §10.4 (R58), **sin nombres ni
códigos de proveedor** (R116): un proveedor completado se dice, y cuántas filas
lo llevan y por qué regla; el proveedor solo va en el v2, que no se versiona.

## Decisiones que no fija la spec (T23, 2026-09-30)

- **Cotejo con los blancos colapsados** (R54): la ubicación y el texto del
  original casan con los de la tabla si son iguales tras colapsar blancos,
  saltos de línea incluidos, y quitar los de los extremos. Mayúsculas, tildes
  y puntuación se comparan tal cual: un carácter cambiado no casa.
- **`origen` es el número de fila de Excel** de una fila con algún dato en las
  columnas A–F; las vacías no cuentan. En el original, sin cabecera, la 1 es la
  primera incidencia.
- **Una sola unidad** (la del YAML): si la columna A del original trae dos
  distintas, se para.
- **`urgencia` y `listado` van por código** en la tabla (`urgente`,
  `seguridad`, `primero`, `segundo`), como el oficio va por nombre de Sigrid y
  no por etiqueta: las etiquetas del YAML de la plantilla pueden cambiar.
- **R91 exige además que el par resuelva al mismo código** que dio el nombre
  de la tabla: si no, completarlo cambiaría el oficio elegido (el único par
  del grupo de `0033`/`0133` es de `0033`, y `0133` se queda sin proveedor,
  como dice `tasks.md` T23).
- **El nombre de Sigrid se compara recortado** por los extremos: la tabla lo
  escribe sin los blancos que pueda arrastrar `auxofc.res`. *(Undécima
  enmienda: igual con los de `oficios_catalogo`; el de la tabla, tal cual.)*

## Decisiones de la undécima enmienda (T62, 2026-10-03)

- **La obra manda**: un nombre que es de la obra y de `oficios_catalogo` se
  resuelve como siempre; `oficios_catalogo` solo se mira cuando el nombre no
  es de ningún oficio de la obra.
- **El número de la fila nueva** de la sección de R130 es su posición dentro
  de su fila del original, como en los mensajes de parada («fila nueva n»).
"""

from __future__ import annotations

import io
import sys
import warnings
from collections import Counter
from collections.abc import Collection, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

# Lanzado como `python scripts/…`, la raíz del servicio no está en la ruta.
sys.path.append(str(Path(__file__).resolve().parent.parent))

from application.pipelines.catalogo_obra import leer_catalogo
from application.pipelines.plantilla import OpcionesDeLaObra, opciones_de_la_obra
from domain.models.equivalencias import MISMO, Catalogo, DecisionPar
from domain.models.errores import (
    CatalogoSinVerificar,
    CodigoDeObraInvalido,
    FicheroDemasiadoGrande,
    FicheroNoEsPlantilla,
    ObraAmbigua,
    ObraSinUnidades,
)
from domain.models.importacion import (
    DESCRIPCION,
    DETALLE,
    LISTADO,
    OFICIO,
    PROVEEDOR,
    UBICACION,
    UNIDAD,
    URGENCIA,
    reconocer_plantilla,
    resolver_oficio_y_proveedor,
    validar_filas,
)
from domain.models.plantilla_incidencias import (
    CatalogoObra,
    Listado,
    ListasCerradas,
    OpcionOficio,
    OpcionProveedor,
    TextosPlantilla,
    Urgencia,
    etiquetas_de_unidades,
    normalizar_codigo_obra,
    plegar,
)
from domain.ports.catalogo_obra import LecturaCatalogo
from domain.ports.hoja_calculo import (
    FilaPlantilla,
    GeneradorPlantillaPort,
    LectorPlantillaPort,
)
from infrastructure.documentos.excel_openpyxl import (
    PRIMERA_FILA,
    GeneradorPlantillaOpenpyxl,
    LectorPlantillaOpenpyxl,
)
from infrastructure.sigrid.consultas_catalogo import (
    fila_a_oficio_catalogo,
    fila_a_unidad_catalogo,
)

#: Las columnas del original que se leen (doc. 05): A, D, E y F.
COLUMNAS_ORIGINAL = 6

#: Textos que no pueden quedar en la descripción ni en el detalle (R57), plegados.
TEXTOS_PROHIBIDOS: tuple[str, ...] = (
    "peligro de seguridad",
    "como en los otros",
    "como en las otras",
    "como en todos",
    "como en todas",
    "igual que",
)

#: Los campos de una fila nueva, que son también lo que puede ser `propuesto`.
CAMPOS_NUEVA: tuple[str, ...] = (
    "ubicacion",
    "descripcion",
    "detalle",
    "oficio",
    "urgencia",
    "listado",
)

PROVEEDOR_COMPLETADO = "proveedor completado (único de la obra para ese oficio)"

#: Las filas nuevas sin oficio porque ese oficio ha salido de la obra (R130).
SIN_OFICIO_FUERA_DE_LA_OBRA = (
    "sin oficio porque ese oficio ya no está en la obra en Sigrid"
)

_CLAVES_TABLA = ("obra", "unidad", "filas")
_CLAVES_FILA = (
    "origen",
    "ubicacion_original",
    "texto_original",
    "resultado",
    "cambios",
    "propuesto",
)
_OBLIGATORIAS_FILA = ("origen", "ubicacion_original", "texto_original", "resultado")
_FORMA_CATALOGO = "el JSON del catálogo no tiene la forma de T2 (`design.md` §10.3)"
_FORMA_GRUPOS = "el fichero de grupos no tiene la forma de R98"
#: La referencia de la obra del JSON: hay una sola y no sale de aquí.
_OBRA_REF = "json"
#: Las decisiones del fichero de grupos no tienen fecha: cualquiera vale.
_CUANDO = datetime.min.replace(tzinfo=UTC)


class MigracionDetenida(Exception):
    """La migración se para sin escribir nada; `problemas` dice por qué."""

    def __init__(self, problemas: Sequence[str]) -> None:
        super().__init__("; ".join(problemas))
        self.problemas = tuple(problemas)


# --------------------------------------------------------------------------
# La tabla de correcciones (§10.2)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class FilaNueva:
    ubicacion: str | None
    descripcion: str
    detalle: str | None
    oficio: str | None  # nombre exacto de Sigrid (R97)
    urgencia: str | None  # código de `Urgencia`
    listado: str | None  # código de `Listado`


@dataclass(frozen=True)
class Correccion:
    origen: int
    ubicacion_original: str | None
    texto_original: str | None
    resultado: tuple[FilaNueva, ...]
    cambios: tuple[str, ...]
    propuesto: tuple[str, ...]


@dataclass(frozen=True)
class TablaCorrecciones:
    obra: str
    unidad: str
    filas: tuple[Correccion, ...]


def _para(problema: str) -> MigracionDetenida:
    return MigracionDetenida([problema])


def _campos(
    crudo: object, donde: str, permitidas: Sequence[str], obligatorias: Sequence[str]
) -> Mapping:
    if not isinstance(crudo, Mapping):
        raise _para(f"{donde} tiene que ser un mapping con sus campos")
    for clave in obligatorias:
        if clave not in crudo:
            raise _para(f"{donde}: falta «{clave}»")
    for clave in crudo:
        if clave not in permitidas:
            raise _para(f"{donde}: sobra «{clave}»")
    return crudo


def _texto(valor: object, donde: str) -> str:
    if not isinstance(valor, str) or not valor.strip():
        raise _para(f"{donde}: tiene que ser un texto")
    return valor


def _texto_o_nada(valor: object, donde: str) -> str | None:
    return None if valor is None else _texto(valor, donde)


def _textos(
    valor: object, donde: str, admitidos: Collection[str] | None = None
) -> tuple[str, ...]:
    if not isinstance(valor, list):
        raise _para(f"{donde}: tiene que ser una lista de textos")
    textos = tuple(_texto(v, donde) for v in valor)
    if admitidos is not None and not set(textos) <= set(admitidos):
        raise _para(f"{donde}: solo admite {', '.join(admitidos)}")
    return textos


def _codigo(valor: object, donde: str, codigos: Collection[str]) -> str | None:
    if valor is not None and valor not in codigos:
        raise _para(f"{donde}: tiene que ser uno de {', '.join(codigos)} o nada")
    return valor  # type: ignore[return-value]


def _fila_nueva(crudo: object, donde: str) -> FilaNueva:
    campos = _campos(crudo, donde, CAMPOS_NUEVA, ("descripcion",))
    return FilaNueva(
        ubicacion=_texto_o_nada(campos.get("ubicacion"), f"{donde}, ubicacion"),
        descripcion=_texto(campos["descripcion"], f"{donde}, descripcion"),
        detalle=_texto_o_nada(campos.get("detalle"), f"{donde}, detalle"),
        oficio=_texto_o_nada(campos.get("oficio"), f"{donde}, oficio"),
        urgencia=_codigo(
            campos.get("urgencia"), f"{donde}, urgencia", [u.value for u in Urgencia]
        ),
        listado=_codigo(
            campos.get("listado"), f"{donde}, listado", [x.value for x in Listado]
        ),
    )


def _correccion(crudo: object, donde: str) -> Correccion:
    campos = _campos(crudo, donde, _CLAVES_FILA, _OBLIGATORIAS_FILA)
    origen = campos["origen"]
    if isinstance(origen, bool) or not isinstance(origen, int) or origen < 1:
        raise _para(f"{donde}, origen: tiene que ser un número de fila desde 1")
    resultado = campos["resultado"]
    if not isinstance(resultado, list):
        raise _para(
            f"{donde}, resultado: tiene que ser una lista (vacía si se descarta)"
        )
    return Correccion(
        origen=origen,
        ubicacion_original=_texto_o_nada(
            campos["ubicacion_original"], f"{donde}, ubicacion_original"
        ),
        texto_original=_texto_o_nada(
            campos["texto_original"], f"{donde}, texto_original"
        ),
        resultado=tuple(
            _fila_nueva(r, f"{donde}, resultado {n}")
            for n, r in enumerate(resultado, start=1)
        ),
        cambios=_textos(campos.get("cambios", []), f"{donde}, cambios"),
        propuesto=_textos(
            campos.get("propuesto", []), f"{donde}, propuesto", CAMPOS_NUEVA
        ),
    )


def leer_correcciones(crudo: object) -> TablaCorrecciones:
    """La tabla de §10.2 ya cargada del YAML, validada en su forma."""
    donde = "la tabla de correcciones"
    campos = _campos(crudo, donde, _CLAVES_TABLA, _CLAVES_TABLA)
    filas = campos["filas"]
    if not isinstance(filas, list) or not filas:
        raise _para(f"{donde}, filas: tiene que ser una lista con una entrada por fila")
    return TablaCorrecciones(
        obra=_texto(campos["obra"], f"{donde}, obra"),
        unidad=_texto(campos["unidad"], f"{donde}, unidad"),
        filas=tuple(
            _correccion(f, f"{donde}, entrada {n}")
            for n, f in enumerate(filas, start=1)
        ),
    )


# --------------------------------------------------------------------------
# El original y el cotejo (R54)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class FilaOriginal:
    numero: int  # fila de Excel; sin cabecera, la 1 es la primera incidencia
    unidad: str | None  # columna A
    ubicacion: str | None  # columna D
    texto: str | None  # columna E
    oficio: str | None  # columna F


def _valor(valor: object) -> str | None:
    if valor is None:
        return None
    texto = valor if isinstance(valor, str) else str(valor)
    return texto if texto.strip() else None


def leer_original(contenido: bytes) -> tuple[FilaOriginal, ...]:
    """Las filas con algún dato de la primera hoja del original (doc. 05)."""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 · un fichero que no es un .xlsx falla de muchas formas
        raise _para("el original no se puede abrir como un Excel (.xlsx)") from None
    try:
        filas: list[FilaOriginal] = []
        recorrido = libro.worksheets[0].iter_rows(
            max_col=COLUMNAS_ORIGINAL, values_only=True
        )
        for numero, celdas in enumerate(recorrido, start=1):
            valores = [_valor(v) for v in (*celdas, *[None] * COLUMNAS_ORIGINAL)]
            if any(valores[:COLUMNAS_ORIGINAL]):
                filas.append(
                    FilaOriginal(numero, valores[0], valores[3], valores[4], valores[5])
                )
    finally:
        libro.close()
    return tuple(filas)


def _colapsado(texto: str | None) -> str:
    return " ".join((texto or "").split())


def _sin_cambios(correccion: Correccion, original: FilaOriginal) -> bool:
    if len(correccion.resultado) != 1:
        return False
    nueva = correccion.resultado[0]
    return (
        nueva.ubicacion == original.ubicacion
        and nueva.descripcion == original.texto
        and nueva.detalle is None
        and nueva.oficio == original.oficio
        and nueva.urgencia is None
        and nueva.listado is None
    )


def cotejar(
    originales: tuple[FilaOriginal, ...], tabla: TablaCorrecciones
) -> list[str]:
    """Lo que no casa entre el original y la tabla (R54); vacío si todo casa."""
    problemas: list[str] = []
    unidades = {_colapsado(o.unidad) for o in originales if o.unidad is not None}
    if len(unidades) > 1:
        problemas.append(
            f"el original trae {len(unidades)} unidades distintas en la columna A y "
            "la tabla de correcciones solo admite una"
        )
    por_numero = {o.numero: o for o in originales}
    origenes = [c.origen for c in tabla.filas]
    veces = Counter(origenes)
    for numero, n in veces.items():
        if n > 1:
            problemas.append(
                f"la fila {numero} del original aparece {n} veces en la tabla"
            )
    for numero in sorted(set(por_numero) - set(veces)):
        problemas.append(
            f"falta la fila {numero} del original en la tabla de correcciones"
        )
    for numero in sorted(set(veces) - set(por_numero)):
        problemas.append(
            f"la fila {numero} de la tabla no es una fila con datos del original"
        )
    if len(veces) == len(origenes) and origenes != sorted(origenes):
        problemas.append("las filas de la tabla no van en el orden del original")
    for correccion in tabla.filas:
        original = por_numero.get(correccion.origen)
        if original is None:
            continue
        numero = correccion.origen
        for que, suyo, dice in (
            ("ubicación", original.ubicacion, correccion.ubicacion_original),
            ("texto", original.texto, correccion.texto_original),
        ):
            if _colapsado(suyo) != _colapsado(dice):
                problemas.append(
                    f"la fila {numero} del original no casa: su {que} es "
                    f"«{_colapsado(suyo)}» y la tabla dice «{_colapsado(dice)}»"
                )
        if not correccion.cambios and not _sin_cambios(correccion, original):
            problemas.append(
                f"la fila {numero} del original cambia y no dice por qué (`cambios`)"
            )
    return problemas


def revisar_contenido(tabla: TablaCorrecciones) -> list[str]:
    """Lo que el script hace cumplir de R57 en descripción y detalle (§10.2)."""
    problemas: list[str] = []
    for correccion in tabla.filas:
        for n, nueva in enumerate(correccion.resultado, start=1):
            for nombre, texto in (
                ("la descripción", nueva.descripcion),
                ("el detalle", nueva.detalle),
            ):
                plegado = plegar(texto or "")
                for prohibido in TEXTOS_PROHIBIDOS:
                    if prohibido in plegado:
                        problemas.append(
                            f"fila {correccion.origen} del original, fila nueva {n}: "
                            f"{nombre} lleva «{prohibido}»"
                        )
    return problemas


# --------------------------------------------------------------------------
# El catálogo del JSON y los grupos del fichero (R55, R97)
# --------------------------------------------------------------------------


class CatalogoDeFichero:
    """`CatalogoObraPort` sobre el JSON de T2: las mismas filas que da Sigrid.

    Mapea con las funciones del adaptador de Sigrid (`fila_a_unidad_catalogo`,
    `fila_a_oficio_catalogo`) para que `leer_catalogo` componga el catálogo
    igual que en la plantilla del portal. Una sola obra, nunca al techo.
    """

    def __init__(self, datos: Mapping) -> None:
        self._datos = datos

    def leer_unidades(self, *, codigo_obra: str) -> LecturaCatalogo:
        obra = self._datos["obra"]
        return LecturaCatalogo(
            tuple(
                fila_a_unidad_catalogo(
                    [
                        _OBRA_REF,
                        obra["codigo"],
                        obra["nombre"],
                        u["codigo"],
                        u["nombre"],
                    ]
                )
                for u in self._datos["unidades"]
            ),
            False,
        )

    def leer_oficios(self, *, obra_ref: str) -> LecturaCatalogo:
        return LecturaCatalogo(
            tuple(
                fila_a_oficio_catalogo(
                    [
                        f["oficio_codigo"],
                        f["oficio_nombre"],
                        f["proveedor_codigo"],
                        f["proveedor_nombre"],
                    ]
                )
                for f in self._datos["oficios_obra"]
            ),
            False,
        )


def catalogo_de_la_obra(datos: object, obra: str) -> CatalogoObra:
    """El catálogo del JSON de T2, que tiene que ser de `obra` (ya normalizada)."""
    try:
        codigo = normalizar_codigo_obra(datos["obra"]["codigo"])  # type: ignore[index]
    except (KeyError, TypeError, CodigoDeObraInvalido):
        raise _para(_FORMA_CATALOGO) from None
    if codigo != obra:
        raise _para(f"el catálogo es de la obra «{codigo}» y la tabla, de la «{obra}»")
    try:
        return leer_catalogo(CatalogoDeFichero(datos), obra)  # type: ignore[arg-type]
    except (KeyError, TypeError, ValueError):
        raise _para(_FORMA_CATALOGO) from None
    except (ObraSinUnidades, ObraAmbigua, CatalogoSinVerificar) as error:
        raise _para(error.motivo) from None


def _fila_de_auxofc(fila: object) -> bool:
    return (
        isinstance(fila, Mapping)
        and "codigo" in fila
        and "nombre" in fila
        and (fila["nombre"] is None or isinstance(fila["nombre"], str))
    )


def nombres_de_sigrid(datos: object) -> frozenset[str]:
    """Los nombres de `oficios_catalogo` del JSON de T2 (todo `auxofc`), recortados
    por los extremos y sin los nulos ni los vacíos (R128, §10.5).

    `oficios_catalogo` tiene que ser una lista de `{codigo, nombre}`, con el
    nombre texto o `null`; si no, para con el error de forma del catálogo.
    """
    try:
        filas = datos["oficios_catalogo"]  # type: ignore[index]
    except (KeyError, TypeError):
        raise _para(_FORMA_CATALOGO) from None
    if not isinstance(filas, list) or not all(_fila_de_auxofc(f) for f in filas):
        raise _para(_FORMA_CATALOGO)
    return frozenset(nombre for f in filas if (nombre := (f["nombre"] or "").strip()))


def decisiones_de_grupos(datos: object, obra: str) -> tuple[DecisionPar, ...]:
    """Los grupos de R98 como decisiones «mismo» entre todos sus pares."""
    if not isinstance(datos, Mapping) or set(datos) != {"obra", "oficio"}:
        raise _para(f"{_FORMA_GRUPOS}: `{{obra, oficio}}`")
    try:
        codigo = normalizar_codigo_obra(datos["obra"])
    except CodigoDeObraInvalido:
        raise _para(f"{_FORMA_GRUPOS}: el código de obra no es válido") from None
    grupos = datos["oficio"]
    if not isinstance(grupos, list) or not all(
        isinstance(g, list)
        and len(g) > 1
        and all(isinstance(c, str) and c.strip() for c in g)
        for g in grupos
    ):
        raise _para(
            f"{_FORMA_GRUPOS}: `oficio` es una lista de grupos de dos o más códigos"
        )
    if codigo != obra:
        raise _para(
            f"el fichero de grupos es de la obra «{codigo}» y la tabla, de la «{obra}»"
        )
    vistos: set[str] = set()
    decisiones: list[DecisionPar] = []
    for grupo in grupos:
        for c in grupo:
            if c in vistos:
                raise _para(f"el fichero de grupos repite el código «{c}»")
            vistos.add(c)
        codigos = sorted(grupo)
        decisiones.extend(
            DecisionPar(
                catalogo=Catalogo.OFICIO,
                codigo_a=a,
                codigo_b=b,
                decision=MISMO,
                motivos=frozenset(),
                obra_codigo=obra,
                decidido_por="fichero de grupos",
                decidido_at_utc=_CUANDO,
            )
            for i, a in enumerate(codigos)
            for b in codigos[i + 1 :]
        )
    return tuple(decisiones)


class EquivalenciasDeFichero:
    """`EquivalenciasPort` sobre el fichero de grupos: solo lee (§10.3)."""

    def __init__(self, decisiones: tuple[DecisionPar, ...]) -> None:
        self._decisiones = decisiones

    def ultimas_decisiones(
        self, *, catalogo: Catalogo, codigos: Collection[str]
    ) -> tuple[DecisionPar, ...]:
        """Como el repositorio: las del catálogo con **los dos** códigos dados."""
        return tuple(
            d
            for d in self._decisiones
            if d.catalogo is catalogo
            and d.codigo_a in codigos
            and d.codigo_b in codigos
        )

    def registrar(self, *, decisiones: tuple[DecisionPar, ...]) -> None:
        raise RuntimeError("la migración no registra decisiones")


# --------------------------------------------------------------------------
# La composición (R91, R97)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class FilaCompuesta:
    origen: int
    nueva: FilaNueva
    oficio_codigo: str | None  # el del nombre de la tabla (R97)
    oficio_etiqueta: str | None  # la que va en el v2: la del grupo
    codigos_en_obra: int  # los del grupo en la obra (0 sin oficio)
    proveedor_codigo: str | None  # completado por R91
    # El nombre de la tabla si la fila va sin oficio porque ha salido de la
    # obra en Sigrid (R128); `None` en los demás casos, también sin oficio.
    oficio_fuera_de_la_obra: str | None
    fila: FilaPlantilla


def _par_de_r91(
    opcion: OpcionOficio, codigo: str, pares: Sequence[OpcionProveedor]
) -> tuple[OpcionProveedor, str] | None:
    """El único par del grupo, si resuelve a este oficio y a un proveedor (R91)."""
    if len(pares) != 1:
        return None
    oficio, proveedor = resolver_oficio_y_proveedor(opcion, pares[0])
    if oficio.codigo != codigo or proveedor.codigo is None:  # type: ignore[union-attr]
        return None
    return pares[0], proveedor.codigo


def componer(
    tabla: TablaCorrecciones,
    catalogo: CatalogoObra,
    opciones: OpcionesDeLaObra,
    listas: ListasCerradas,
    *,
    nombres_sigrid: Collection[str],
) -> tuple[FilaCompuesta, ...]:
    """Las filas del v2, en el orden de la tabla (R91, R97, R128).

    El nombre de oficio de la tabla, tal cual, frente a los de Sigrid
    recortados; la obra manda (§10.5): de **un** oficio de la obra, su grupo y
    el proveedor por R91; de **varios**, para; de ninguno pero sí de
    `nombres_sigrid` (todo `auxofc`), oficio y proveedor vacíos; de ninguno de
    Sigrid, para (la errata). Los problemas se dicen todos a la vez.
    """
    codigos_de: dict[str, list[str]] = {}
    for oficio in catalogo.oficios:
        codigos_de.setdefault((oficio.nombre or "").strip(), []).append(oficio.codigo)
    opcion_de = {c: o for o in opciones.oficios for c in o.codigos_en_obra}
    pares_de: dict[OpcionOficio, list[OpcionProveedor]] = {}
    for par in opciones.proveedores:
        pares_de.setdefault(par.oficio, []).append(par)
    urgencias = {o.codigo: o.etiqueta for o in listas.urgencias}
    listados = {o.codigo: o.etiqueta for o in listas.listados}

    problemas: list[str] = []
    compuestas: list[FilaCompuesta] = []
    for correccion in tabla.filas:
        for n, nueva in enumerate(correccion.resultado, start=1):
            codigo = opcion = par = proveedor = fuera = None
            if nueva.oficio is not None:
                codigos = codigos_de.get(nueva.oficio, [])
                donde = f"fila {correccion.origen} del original, fila nueva {n}"
                if len(codigos) > 1:
                    problemas.append(
                        f"{donde}: el oficio «{nueva.oficio}» es el nombre de varios "
                        f"oficios de la obra en Sigrid ({', '.join(codigos)}): no se "
                        "puede elegir"
                    )
                    continue
                if not codigos and nueva.oficio not in nombres_sigrid:
                    problemas.append(
                        f"{donde}: el oficio «{nueva.oficio}» no es exactamente el "
                        "nombre de ningún oficio de Sigrid, ni de la obra ni del resto "
                        "del catálogo: corrígelo en la tabla de correcciones (R97)"
                    )
                    continue
                if not codigos:
                    # Ha salido de la obra en Sigrid: sin oficio ni proveedor (R128).
                    fuera = nueva.oficio
                else:
                    (codigo,) = codigos
                    opcion = opcion_de[codigo]
                    elegido = _par_de_r91(opcion, codigo, pares_de.get(opcion, []))
                    if elegido is not None:
                        par, proveedor = elegido
            valores = {
                UNIDAD: tabla.unidad,
                UBICACION: nueva.ubicacion,
                DESCRIPCION: nueva.descripcion,
                DETALLE: nueva.detalle,
                OFICIO: None if opcion is None else opcion.etiqueta,
                PROVEEDOR: None if par is None else par.etiqueta,
                URGENCIA: None if nueva.urgencia is None else urgencias[nueva.urgencia],
                LISTADO: None if nueva.listado is None else listados[nueva.listado],
            }
            compuestas.append(
                FilaCompuesta(
                    origen=correccion.origen,
                    nueva=nueva,
                    oficio_codigo=codigo,
                    oficio_etiqueta=None if opcion is None else opcion.etiqueta,
                    codigos_en_obra=0
                    if opcion is None
                    else len(opcion.codigos_en_obra),
                    proveedor_codigo=proveedor,
                    oficio_fuera_de_la_obra=fuera,
                    fila=FilaPlantilla(valores=valores),
                )
            )
    if problemas:
        raise MigracionDetenida(problemas)
    return tuple(compuestas)


# --------------------------------------------------------------------------
# La ida y vuelta por el importador (R56)
# --------------------------------------------------------------------------


def avisos_sin_proveedor(avisos: Iterable[str]) -> tuple[str, ...]:
    """Los avisos del importador sin los que nombran un proveedor (R58)."""
    return tuple(a for a in avisos if not a.startswith("el proveedor "))


def ida_y_vuelta(
    v2: bytes,
    *,
    catalogo: CatalogoObra,
    opciones: OpcionesDeLaObra,
    listas: ListasCerradas,
    compuestas: tuple[FilaCompuesta, ...],
    lector: LectorPlantillaPort,
) -> tuple[tuple[str, ...], ...]:
    """El v2 por el lector, `reconocer_plantilla` y `validar_filas`; sus avisos."""
    try:
        libro = lector.leer(contenido=v2)
        obra = reconocer_plantilla(libro)
    except (FicheroNoEsPlantilla, FicheroDemasiadoGrande) as error:
        raise _para(f"el v2 no pasa el lector del importador: {error.motivo}") from None
    if obra != catalogo.obra_codigo:
        raise _para(f"el v2 se reconoce como la plantilla de otra obra («{obra}»)")
    validas, con_error = validar_filas(
        libro.filas,
        catalogo=catalogo,
        opciones_oficio=opciones.oficios,
        opciones_proveedor=opciones.proveedores,
        listas=listas,
    )
    if con_error:
        raise MigracionDetenida(
            [
                f"fila {compuestas[e.fila - PRIMERA_FILA].origen} del original "
                f"(fila {e.fila} del v2), {e.columna}: {e.problema}"
                for fila in con_error
                for e in fila.errores
            ]
        )
    if len(validas) != len(compuestas):
        raise _para(
            f"el v2 trae {len(validas)} filas válidas al volver a leerlo y se "
            f"escribieron {len(compuestas)}"
        )
    return tuple(avisos_sin_proveedor(v.avisos) for v in validas)


# --------------------------------------------------------------------------
# De punta a punta
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Recuentos:
    filas_originales: int
    filas_nuevas: int
    separadas: int
    descartadas: int
    oficios_completados: int  # la muestra no traía oficio y la fila nueva sí
    oficios_no_exactos: int  # el de la muestra no es un nombre de Sigrid de la obra
    oficios_en_grupo_de_varios: int
    # Filas nuevas sin oficio porque ha salido de la obra (R128, R130); no
    # cuenta las filas a las que la tabla no pone oficio.
    oficios_fuera_de_la_obra: int
    proveedores_completados: int
    urgencias: Mapping[str, int]  # por código de `Urgencia`
    con_avisos: int


@dataclass(frozen=True)
class Migracion:
    obra: str
    unidad: str
    con_grupos: bool
    grupos_de_varios: int
    originales: tuple[FilaOriginal, ...]
    tabla: TablaCorrecciones
    catalogo: CatalogoObra
    opciones: OpcionesDeLaObra
    listas: ListasCerradas
    compuestas: tuple[FilaCompuesta, ...]
    avisos: tuple[tuple[str, ...], ...]  # uno por fila compuesta
    v2: bytes
    recuentos: Recuentos


def _recuentos(
    originales: tuple[FilaOriginal, ...],
    tabla: TablaCorrecciones,
    catalogo: CatalogoObra,
    compuestas: tuple[FilaCompuesta, ...],
    avisos: tuple[tuple[str, ...], ...],
) -> Recuentos:
    muestra = {o.numero: o.oficio for o in originales}
    nombres = {(o.nombre or "").strip() for o in catalogo.oficios}
    urgencias = Counter(c.nueva.urgencia for c in compuestas)
    return Recuentos(
        filas_originales=len(originales),
        filas_nuevas=len(compuestas),
        separadas=sum(1 for c in tabla.filas if len(c.resultado) > 1),
        descartadas=sum(1 for c in tabla.filas if not c.resultado),
        oficios_completados=sum(
            1 for c in compuestas if c.oficio_codigo and muestra[c.origen] is None
        ),
        oficios_no_exactos=sum(
            1 for o in originales if o.oficio is not None and o.oficio not in nombres
        ),
        oficios_en_grupo_de_varios=sum(1 for c in compuestas if c.codigos_en_obra > 1),
        oficios_fuera_de_la_obra=sum(
            1 for c in compuestas if c.oficio_fuera_de_la_obra is not None
        ),
        proveedores_completados=sum(1 for c in compuestas if c.proveedor_codigo),
        urgencias={u.value: urgencias[u.value] for u in Urgencia},
        con_avisos=sum(1 for a in avisos if a),
    )


def migrar(
    *,
    original: bytes,
    correcciones: object,
    catalogo: object,
    grupos: object | None,
    listas: ListasCerradas,
    textos: TextosPlantilla,
    ahora: datetime,
    generador: GeneradorPlantillaPort | None = None,
    lector: LectorPlantillaPort | None = None,
) -> Migracion:
    """Los pasos 2–4 de §10.3, en memoria. `MigracionDetenida` si algo no casa.

    `grupos` es el JSON de R98 ya cargado, o `None`: sin él, cada código es su
    propio grupo (la línea de órdenes solo lo admite con `--solo-informe`).
    """
    tabla = leer_correcciones(correcciones)
    try:
        obra = normalizar_codigo_obra(tabla.obra)
    except CodigoDeObraInvalido:
        raise _para(
            "la tabla de correcciones: el código de obra no es válido"
        ) from None
    originales = leer_original(original)
    problemas = cotejar(originales, tabla) + revisar_contenido(tabla)
    if problemas:
        raise MigracionDetenida(problemas)

    cat = catalogo_de_la_obra(catalogo, obra)
    nombres_sigrid = nombres_de_sigrid(catalogo)
    decisiones = () if grupos is None else decisiones_de_grupos(grupos, obra)
    try:
        opciones = opciones_de_la_obra(cat, EquivalenciasDeFichero(decisiones))
    except ValueError as error:
        raise _para(
            f"las opciones de la plantilla no se pueden componer: {error}"
        ) from None
    if tabla.unidad not in {o.etiqueta for o in etiquetas_de_unidades(cat.unidades)}:
        raise _para(
            f"la unidad «{tabla.unidad}» no es la etiqueta de ninguna unidad de la "
            "obra en el desplegable de la plantilla"
        )
    compuestas = componer(tabla, cat, opciones, listas, nombres_sigrid=nombres_sigrid)

    try:
        v2 = (generador or GeneradorPlantillaOpenpyxl()).generar(
            catalogo=cat,
            opciones_oficio=opciones.oficios,
            opciones_proveedor=opciones.proveedores,
            listas=listas,
            textos=textos,
            generada_at=ahora,
            filas=tuple(c.fila for c in compuestas),
            importacion_origen=None,
        )
    except ValueError as error:
        raise _para(f"el v2 no se puede generar: {error}") from None
    avisos = ida_y_vuelta(
        v2,
        catalogo=cat,
        opciones=opciones,
        listas=listas,
        compuestas=compuestas,
        lector=lector or LectorPlantillaOpenpyxl(),
    )
    return Migracion(
        obra=obra,
        unidad=tabla.unidad,
        con_grupos=grupos is not None,
        grupos_de_varios=0 if grupos is None else len(grupos["oficio"]),  # type: ignore[index]
        originales=originales,
        tabla=tabla,
        catalogo=cat,
        opciones=opciones,
        listas=listas,
        compuestas=compuestas,
        avisos=avisos,
        v2=v2,
        recuentos=_recuentos(originales, tabla, cat, compuestas, avisos),
    )


# --------------------------------------------------------------------------
# El informe (R58, §10.4)
# --------------------------------------------------------------------------


def _celda(texto: str | None) -> str:
    """Un texto dentro de una celda de tabla Markdown, sin romperla."""
    if not texto:
        return "—"
    escapado = (
        texto.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("|", "\\|")
    )
    return "<br>".join(escapado.splitlines())


def _descripcion_de(
    compuesta: FilaCompuesta, avisos: tuple[str, ...], listas: ListasCerradas
) -> str:
    nueva = compuesta.nueva
    urgencias = {o.codigo: o.etiqueta for o in listas.urgencias}
    listados = {o.codigo: o.etiqueta for o in listas.listados}
    partes = [
        f"Ubicación: {nueva.ubicacion or '—'}",
        f"Descripción: {nueva.descripcion}",
    ]
    if nueva.detalle is not None:
        partes.append(f"Detalle: {nueva.detalle}")
    if compuesta.oficio_fuera_de_la_obra is not None:
        partes.append(
            "Oficio: — (ya no está en la obra en Sigrid: "
            f"«{compuesta.oficio_fuera_de_la_obra}»)"
        )
    elif compuesta.oficio_codigo is None:
        partes.append("Oficio: —")
    else:
        oficio = f"Oficio: {nueva.oficio} (código `{compuesta.oficio_codigo}`"
        if compuesta.oficio_etiqueta != nueva.oficio or compuesta.codigos_en_obra > 1:
            oficio += f"; en el v2, «{compuesta.oficio_etiqueta}»"
            if compuesta.codigos_en_obra > 1:
                oficio += f", grupo con {compuesta.codigos_en_obra} códigos en la obra"
        partes.append(oficio + ")")
    if compuesta.proveedor_codigo is not None:
        # Sin el código (R116): el de un autónomo lleva a su nombre en Sigrid.
        partes.append(f"Proveedor: {PROVEEDOR_COMPLETADO}")
    partes.append(f"Urgencia: {urgencias[nueva.urgencia] if nueva.urgencia else '—'}")
    if nueva.listado is not None:
        partes.append(f"Listado: {listados[nueva.listado]}")
    partes.extend(f"Aviso: {a}" for a in avisos)
    return " · ".join(partes)


def informe(
    m: Migracion, *, sha256_original: str, nombre: str = "migracion_F-036.md"
) -> str:
    """El informe de revisión de §10.4, en Markdown, sin nombres ni códigos de
    proveedor (R58, R116)."""
    r = m.recuentos
    etiquetas_urgencia = {o.codigo: o.etiqueta for o in m.listas.urgencias}
    grupos = (
        f"fichero de grupos vigentes ({m.grupos_de_varios} grupos de varios códigos)"
        if m.con_grupos
        else "sin fichero de grupos: cada código es su propio grupo (`--solo-informe`)"
    )
    lineas = [
        f"<!-- progress/{nombre} -->",
        "# F-036 · Migración del Excel de incidencias: propuesta para revisar",
        "",
        (
            "> Generado por `services/postventa-api/scripts/migrar_excel_f036.py` desde "
            "la tabla de correcciones (`scripts/migracion_f036_correcciones.yaml`). **No "
            "se edita a mano**: los cambios van a la tabla y se regenera (T25). Sin "
            "nombres ni códigos de proveedor (R58, R116): un proveedor completado se "
            "dice sin decir cuál; va en el v2, que no se versiona."
        ),
        "",
        f"- Obra `{m.obra}` · unidad «{m.unidad}»",
        f"- Grupos de oficio: {grupos}",
        f"- `sha256` del original: `{sha256_original}`",
        "- Ida y vuelta por el importador: **0 errores**",
        "",
        "## Recuentos",
        "",
        "| Qué | Cuántas |",
        "|---|---:|",
        f"| Filas del original | {r.filas_originales} |",
        f"| Filas nuevas | {r.filas_nuevas} |",
        f"| Filas del original separadas en varias | {r.separadas} |",
        f"| Filas del original descartadas | {r.descartadas} |",
        f"| Oficios completados (la muestra no lo traía) | {r.oficios_completados} |",
        (
            "| Oficios de la muestra que no son exactamente un nombre de Sigrid de la "
            f"obra | {r.oficios_no_exactos} |"
        ),
        (
            "| Filas nuevas con un oficio cuyo grupo tiene varios códigos en la obra | "
            f"{r.oficios_en_grupo_de_varios} |"
        ),
        f"| Filas nuevas {SIN_OFICIO_FUERA_DE_LA_OBRA} | {r.oficios_fuera_de_la_obra} |",
        (
            "| Filas nuevas con proveedor, completado por R91 (el único de la obra "
            f"para ese oficio) | {r.proveedores_completados} |"
        ),
        *(
            f"| Urgencias marcadas: {etiquetas_urgencia[codigo]} | {n} |"
            for codigo, n in r.urgencias.items()
        ),
        f"| Filas nuevas con algún aviso del importador | {r.con_avisos} |",
    ]

    originales = {o.numero: o for o in m.originales}
    por_origen: dict[int, list[int]] = {}
    for indice, compuesta in enumerate(m.compuestas):
        por_origen.setdefault(compuesta.origen, []).append(indice)

    if r.oficios_fuera_de_la_obra:
        # R130: justo después de los recuentos, sin nada de proveedores (R116).
        lineas += [
            "",
            f"## Filas {SIN_OFICIO_FUERA_DE_LA_OBRA}",
            "",
            "| Fila del original | Fila nueva | Oficio en la tabla de correcciones |",
            "|---:|---:|---|",
        ]
        for indices in por_origen.values():
            for n, i in enumerate(indices, start=1):
                fuera = m.compuestas[i].oficio_fuera_de_la_obra
                if fuera is not None:
                    lineas.append(
                        f"| {m.compuestas[i].origen} | {n} | {_celda(fuera)} |"
                    )

    por_ubicacion: dict[str | None, list[Correccion]] = {}
    for correccion in m.tabla.filas:
        ubicacion = originales[correccion.origen].ubicacion
        por_ubicacion.setdefault(ubicacion, []).append(correccion)

    for ubicacion, correcciones in por_ubicacion.items():
        lineas += [
            "",
            f"## Ubicación original «{_celda(ubicacion)}»",
            "",
            "| Fila | Texto original | Oficio en la muestra | Filas nuevas | Cambios | Propuesto |",
            "|---:|---|---|---|---|---|",
        ]
        for correccion in correcciones:
            original = originales[correccion.origen]
            indices = por_origen.get(correccion.origen, [])
            nuevas = "<br>".join(
                f"{n}. "
                + _celda(_descripcion_de(m.compuestas[i], m.avisos[i], m.listas))
                for n, i in enumerate(indices, start=1)
            )
            propuesto = list(correccion.propuesto)
            if any(m.compuestas[i].proveedor_codigo for i in indices):
                propuesto.append("proveedor")
            lineas.append(
                f"| {correccion.origen} | {_celda(original.texto)} | "
                f"{_celda(original.oficio)} | {nuevas or '— (descartada)'} | "
                f"{'<br>'.join(_celda(c) for c in correccion.cambios) or '—'} | "
                f"{', '.join(propuesto) or '—'} |"
            )
    return "\n".join(lineas) + "\n"

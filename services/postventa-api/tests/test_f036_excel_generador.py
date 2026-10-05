# services/postventa-api/tests/test_f036_excel_generador.py
"""El generador de la plantilla y del Excel de errores (F-036, T10).

`design.md` §3 y §5.2; R2–R7, R12, R62–R65 y R92. Cada test genera el libro
**en memoria** y lo reabre con `openpyxl`: se comprueba lo que verá Excel, no
lo que el generador cree haber escrito. Ningún `.xlsx` toca el disco (R59) y
los nombres de proveedor son inventados.

Las hojas, rangos y nombres van escritos aquí **literales**, no importados del
generador: si el generador cambia uno, el test tiene que enterarse.
"""

from __future__ import annotations

import base64
import dataclasses
import functools
import io
import re
import zipfile
from datetime import datetime, timedelta, timezone
from xml.etree import ElementTree

import pytest
from openpyxl.utils import get_column_letter

from domain.models.plantilla_incidencias import (
    CABECERA,
    COLUMNA_ERRORES,
    IDENTIFICADOR_PLANTILLA,
    MAX_DESCRIPCION,
    MAX_DETALLE,
    MAX_FILAS,
    OficioObra,
    ProveedorEnObra,
    UnidadPosventa,
    etiquetas_de_unidades,
)
from domain.ports.hoja_calculo import FilaPlantilla
from infrastructure.documentos import excel_openpyxl
from infrastructure.documentos.excel_openpyxl import GeneradorPlantillaOpenpyxl
from tests.utiles_plantilla import (
    IMPORTACION,
    INSTANTE,
    abrir,
    catalogo,
    configuracion,
    generar,
    mismo,
    opciones,
)

HOJAS = ["Incidencias", "Instrucciones", "_catalogos", "_plantilla"]
COLUMNAS_TASADAS = ("Unidad", "Ubicación", "Oficio", "Proveedor", "Urgencia", "Listado")
COLUMNAS_DE_DATOS = tuple(c for c in CABECERA if c != COLUMNA_ERRORES)


def _letra(columna: str) -> str:
    return get_column_letter(CABECERA.index(columna) + 1)


def _rango(columna: str) -> str:
    letra = _letra(columna)
    return f"{letra}2:{letra}1001"


def _validaciones(ws, columna: str) -> list:
    """Las validaciones de datos que cubren exactamente la columna, filas 2–1001."""
    return [
        dv
        for dv in ws.data_validations.dataValidation
        if str(dv.sqref) == _rango(columna)
    ]


def _validacion(ws, columna: str):
    encontradas = _validaciones(ws, columna)
    assert len(encontradas) == 1, f"{columna}: {len(encontradas)} validaciones"
    return encontradas[0]


def _lista_del_desplegable(wb, columna: str) -> tuple[str, ...]:
    """Los valores que ofrece el desplegable, siguiendo el rango con nombre."""
    dv = _validacion(wb["Incidencias"], columna)
    assert dv.type == "list"
    nombre = dv.formula1
    assert nombre in wb.defined_names, f"{columna}: {nombre} no es un rango con nombre"
    destinos = list(wb.defined_names[nombre].destinations)
    assert len(destinos) == 1
    hoja, rango = destinos[0]
    assert hoja == "_catalogos"
    celdas = wb[hoja][rango.replace("$", "")]
    if not isinstance(celdas, tuple):
        celdas = ((celdas,),)
    elif celdas and not isinstance(celdas[0], tuple):
        celdas = (celdas,)
    return tuple(fila[0].value for fila in celdas)


def _metadatos(wb) -> dict[str, object]:
    ws = wb["_plantilla"]
    return {fila[0].value: fila[1].value for fila in ws.iter_rows(min_col=1, max_col=2)}


def _textos_de(ws) -> list[str]:
    return [
        c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)
    ]


def _esperadas(columna: str, cat=None, decisiones=()) -> tuple[str, ...]:
    cat = cat or catalogo()
    oficios, pares = opciones(cat, decisiones)
    listas = configuracion().listas
    return {
        "Unidad": tuple(o.etiqueta for o in etiquetas_de_unidades(cat.unidades)),
        "Ubicación": listas.ubicaciones,
        "Oficio": tuple(o.etiqueta for o in oficios),
        "Proveedor": tuple(p.etiqueta for p in pares),
        "Urgencia": tuple(o.etiqueta for o in listas.urgencias),
        "Listado": tuple(o.etiqueta for o in listas.listados),
    }[columna]


def _fila(**valores: str | None) -> dict[str, str | None]:
    """Valores por columna, con los nombres de la cabecera (sin tildes en el código)."""
    nombres = {
        "unidad": "Unidad",
        "ubicacion": "Ubicación",
        "descripcion": "Descripción corta",
        "detalle": "Detalle",
        "oficio": "Oficio",
        "proveedor": "Proveedor",
        "urgencia": "Urgencia",
        "listado": "Listado",
    }
    return {nombres[k]: v for k, v in valores.items()}


def _fuente(celda) -> tuple[str | None, float | None, bool, str | None]:
    """Nombre, tamaño, negrita y color ARGB de la fuente, como los verá Excel."""
    fuente = celda.font
    color = fuente.color.rgb if fuente.color is not None else None
    return (
        fuente.name,
        fuente.sz,
        fuente.b is True,
        color if isinstance(color, str) else None,
    )


def _relleno(celda) -> tuple[str | None, str | None]:
    relleno = celda.fill
    if relleno.fill_type is None:
        return None, None
    return relleno.fill_type, relleno.fgColor.rgb


def _lado(lado) -> tuple[str | None, str | None]:
    if lado is None or lado.style is None:
        return None, None
    return lado.style, lado.color.rgb if lado.color is not None else None


def _bordes(celda) -> dict[str, tuple[str | None, str | None]]:
    borde = celda.border
    return {
        "izquierdo": _lado(borde.left),
        "derecho": _lado(borde.right),
        "superior": _lado(borde.top),
        "inferior": _lado(borde.bottom),
    }


def _alineacion(celda) -> tuple[str | None, str | None, bool, float]:
    """Horizontal, vertical, ajuste de texto y sangría."""
    a = celda.alignment
    return a.horizontal, a.vertical, a.wrap_text is True, a.indent


FILAS_CON_ERROR = (
    FilaPlantilla(
        valores=_fila(unidad="villa 1", descripcion="Sellar sifón", oficio="Mamparas"),
        errores={"Unidad": "ese valor no está en la lista de unidades"},
        texto_errores="Fila 7 del fichero subido · Unidad: ese valor no está en la lista",
    ),
    FilaPlantilla(
        valores=_fila(
            unidad="Viviendas Bloque Villa 2",
            ubicacion="Cocina",
            descripcion="=SUMA(A1:A3)",
            detalle=" con espacios \n y salto ",
            urgencia="urgente",
        ),
        errores={
            "Descripción corta": "la celda lleva una fórmula: escribe el texto",
            "Urgencia": "ese valor no está en la lista de urgencias",
        },
        texto_errores="Fila 9 del fichero subido · Descripción corta: fórmula · Urgencia: …",
    ),
)


@functools.cache
def _excel_errores() -> bytes:
    return generar(filas=FILAS_CON_ERROR, importacion_origen=IMPORTACION)


# Generar y reabrir cuesta casi un segundo: los tests que solo **leen** el libro
# por defecto comparten uno. Ninguno lo modifica.
@functools.cache
def _plantilla():
    return abrir(generar())


@functools.cache
def _errores():
    return abrir(_excel_errores())


# --------------------------------------------------------------------------
# R2 · hojas, cabecera, paneles, filtro y protección
# --------------------------------------------------------------------------


def test_f036_r2_hojas_en_orden_con_su_estado_e_incidencias_activa() -> None:
    wb = _plantilla()
    assert wb.sheetnames == HOJAS
    assert [wb[h].sheet_state for h in HOJAS] == [
        "visible",
        "visible",
        "veryHidden",
        "veryHidden",
    ]
    assert wb.active.title == "Incidencias"


def test_f036_r2_cabecera_exacta_en_la_fila_1_y_nada_mas() -> None:
    ws = _plantilla()["Incidencias"]
    assert tuple(c.value for c in ws[1]) == CABECERA
    assert ws.max_column == len(CABECERA)


def test_f036_r2_paneles_inmovilizados_bajo_la_cabecera_y_filtro() -> None:
    ws = _plantilla()["Incidencias"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:I1001"


def test_f036_r2_la_plantilla_no_trae_filas_de_datos() -> None:
    ws = _plantilla()["Incidencias"]
    valores = [c.value for fila in ws.iter_rows(min_row=2) for c in fila]
    assert all(v is None for v in valores)


def test_f036_r2_proteccion_sin_contrasena_con_filas_orden_y_filtro_permitidos() -> (
    None
):
    ws = _plantilla()["Incidencias"]
    proteccion = ws.protection
    assert proteccion.sheet is True
    assert not proteccion.password
    # En OOXML «False» es «no protegido»: se permite.
    assert proteccion.insertRows is False
    assert proteccion.deleteRows is False
    assert proteccion.sort is False
    assert proteccion.autoFilter is False
    # Lo demás sigue protegido: la cabecera no se toca por accidente.
    assert proteccion.formatCells is True
    assert proteccion.insertColumns is True
    assert proteccion.deleteColumns is True


@pytest.mark.parametrize("columna", COLUMNAS_DE_DATOS)
def test_f036_r2_celdas_de_datos_desbloqueadas_y_como_texto(columna: str) -> None:
    ws = _plantilla()["Incidencias"]
    letra = _letra(columna)
    for fila in (2, 500, 1001):
        celda = ws[f"{letra}{fila}"]
        assert celda.protection.locked is False, f"{letra}{fila}"
        # Formato texto: Excel no convierte lo tecleado en fecha, número ni fórmula.
        assert celda.number_format == "@", f"{letra}{fila}"
    assert ws[f"{letra}1"].protection.locked is True
    assert ws[f"{letra}1002"].protection.locked is True


def test_f036_r2_la_columna_errores_y_la_cabecera_quedan_bloqueadas() -> None:
    ws = _plantilla()["Incidencias"]
    for fila in (1, 2, 1001):
        assert ws[f"I{fila}"].protection.locked is True
    assert all(c.protection.locked is True for c in ws[1])


def test_f036_r2_las_hojas_ocultas_tambien_van_protegidas() -> None:
    wb = _plantilla()
    assert wb["_catalogos"].protection.sheet is True
    assert wb["_plantilla"].protection.sheet is True


@pytest.mark.parametrize("columna", ["Descripción corta", "Detalle", COLUMNA_ERRORES])
def test_f036_r2_textos_largos_con_ancho_y_ajuste(columna: str) -> None:
    ws = _plantilla()["Incidencias"]
    letra = _letra(columna)
    assert ws.column_dimensions[letra].width >= 40
    assert ws[f"{letra}2"].alignment.wrap_text is True
    assert ws[f"{letra}1001"].alignment.wrap_text is True


def test_f036_r2_la_cabecera_se_distingue() -> None:
    ws = _plantilla()["Incidencias"]
    assert all(c.font.bold for c in ws[1])


# --------------------------------------------------------------------------
# R3 · desplegables
# --------------------------------------------------------------------------


@pytest.mark.parametrize("columna", COLUMNAS_TASADAS)
def test_f036_r3_desplegable_detener_en_filas_2_a_1001_con_mensajes(
    columna: str,
) -> None:
    ws = _plantilla()["Incidencias"]
    dv = _validacion(ws, columna)
    mensaje = configuracion().textos.mensajes[columna]
    assert dv.type == "list"
    assert dv.errorStyle == "stop"
    assert dv.showErrorMessage is True
    assert dv.showInputMessage is True
    assert dv.allow_blank is True
    assert (dv.promptTitle, dv.prompt) == (mensaje.titulo_entrada, mensaje.entrada)
    assert (dv.errorTitle, dv.error) == (mensaje.titulo_error, mensaje.error)


@pytest.mark.parametrize("columna", COLUMNAS_TASADAS)
def test_f036_r3_el_desplegable_sale_de_la_hoja_de_catalogos(columna: str) -> None:
    wb = _plantilla()
    assert _lista_del_desplegable(wb, columna) == _esperadas(columna)


def test_f036_r3_cada_desplegable_su_propio_rango_con_nombre() -> None:
    wb = _plantilla()
    ws = wb["Incidencias"]
    nombres = [_validacion(ws, c).formula1 for c in COLUMNAS_TASADAS]
    assert len(set(nombres)) == len(COLUMNAS_TASADAS)


def test_f036_r3_la_hoja_de_catalogos_solo_lleva_texto() -> None:
    ws = _plantilla()["_catalogos"]
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is not None:
                assert celda.data_type == "s", celda.coordinate


def test_f036_r3_las_columnas_libres_y_errores_sin_desplegable() -> None:
    ws = _plantilla()["Incidencias"]
    listas = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    cubiertas = {str(dv.sqref)[0] for dv in listas}
    assert cubiertas == {_letra(c) for c in COLUMNAS_TASADAS}
    assert not _validaciones(ws, COLUMNA_ERRORES)


# --------------------------------------------------------------------------
# R4 · longitudes
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("columna", "minimo", "maximo"),
    [("Descripción corta", 1, MAX_DESCRIPCION), ("Detalle", 0, MAX_DETALLE)],
)
def test_f036_r4_longitud_de_los_textos_libres(
    columna: str, minimo: int, maximo: int
) -> None:
    ws = _plantilla()["Incidencias"]
    dv = _validacion(ws, columna)
    mensaje = configuracion().textos.mensajes[columna]
    assert dv.type == "textLength"
    assert dv.operator == "between"
    assert (dv.formula1, dv.formula2) == (str(minimo), str(maximo))
    assert dv.errorStyle == "stop"
    assert dv.showErrorMessage is True
    assert dv.showInputMessage is True
    assert (dv.promptTitle, dv.prompt) == (mensaje.titulo_entrada, mensaje.entrada)
    assert (dv.errorTitle, dv.error) == (mensaje.titulo_error, mensaje.error)


def test_f036_r4_topes_del_contrato() -> None:
    assert (MAX_DESCRIPCION, MAX_DETALLE) == (128, 2000)


# --------------------------------------------------------------------------
# R5 · Instrucciones
# --------------------------------------------------------------------------


def test_f036_r5_instrucciones_con_la_obra_y_todos_los_parrafos() -> None:
    ws = _plantilla()["Instrucciones"]
    textos = configuracion().textos
    assert ws["A1"].value == "0677 · Obra Ejemplo"
    presentes = _textos_de(ws)
    assert textos.titulo in presentes
    for parrafo in textos.instrucciones:
        assert parrafo in presentes
    # Solo en una obra sin oficios y en el Excel de errores.
    assert textos.sin_oficios not in presentes
    assert textos.excel_errores not in presentes


def test_f036_r5_obra_sin_nombre_titulo_con_el_codigo() -> None:
    ws = abrir(generar(catalogo(obra_nombre=None)))["Instrucciones"]
    assert ws["A1"].value == "0677"


def test_f036_r5_los_ejemplos_van_en_instrucciones_como_tabla() -> None:
    ws = _plantilla()["Instrucciones"]
    textos = configuracion().textos
    filas = [tuple(c.value for c in fila) for fila in ws.iter_rows()]
    cabecera = tuple(COLUMNAS_DE_DATOS)
    assert cabecera in [f[: len(cabecera)] for f in filas]
    inicio = [f[: len(cabecera)] for f in filas].index(cabecera)
    for desplazamiento, ejemplo in enumerate(textos.ejemplos, start=1):
        esperada = tuple(ejemplo.get(c) for c in cabecera)
        assert filas[inicio + desplazamiento][: len(cabecera)] == esperada


def test_f036_r5_los_parrafos_se_leen_sin_cortar() -> None:
    ws = _plantilla()["Instrucciones"]
    assert ws.column_dimensions["A"].width >= 80
    for fila in ws.iter_rows(min_row=2, max_col=1):
        celda = fila[0]
        if celda.value in configuracion().textos.instrucciones:
            assert celda.alignment.wrap_text is True


# --------------------------------------------------------------------------
# R6 · metadatos
# --------------------------------------------------------------------------


def test_f036_r6_metadatos_de_la_plantilla() -> None:
    metadatos = _metadatos(_plantilla())
    assert metadatos == {
        "identificador": IDENTIFICADOR_PLANTILLA,
        "version": 1,
        "obra_codigo": "0677",
        "generada_at_utc": "2026-09-29T08:30:00+00:00",
    }


def test_f036_r6_el_instante_se_escribe_en_utc() -> None:
    madrid = datetime(2026, 9, 29, 10, 30, tzinfo=timezone(timedelta(hours=2)))
    metadatos = _metadatos(abrir(generar(generada_at=madrid)))
    assert metadatos["generada_at_utc"] == "2026-09-29T08:30:00+00:00"


def test_f036_r6_un_instante_sin_zona_se_rechaza() -> None:
    with pytest.raises(ValueError, match="zona"):
        generar(generada_at=datetime(2026, 9, 29, 8, 30))  # noqa: DTZ001 (a propósito)


def test_f036_r6_el_codigo_de_obra_va_como_texto() -> None:
    ws = _plantilla()["_plantilla"]
    celdas = {f[0].value: f[1] for f in ws.iter_rows(min_col=1, max_col=2)}
    assert celdas["obra_codigo"].data_type == "s"
    assert celdas["identificador"].data_type == "s"
    assert celdas["version"].data_type == "n"


# --------------------------------------------------------------------------
# R7 · ninguna fórmula
# --------------------------------------------------------------------------

PELIGROSOS = ('=HYPERLINK("x")', "+34 600", "-3", "@SUMA(A1)")


def _catalogo_peligroso():
    return catalogo(
        unidades=tuple(
            UnidadPosventa(f"U{i}", nombre) for i, nombre in enumerate(PELIGROSOS)
        ),
        oficios=(OficioObra("0046", "=Carpintería"),),
        filas=(ProveedorEnObra("0046", "P1", "@Proveedor Ejemplo"),),
        obra_nombre="=Obra",
    )


def _filas_peligrosas() -> tuple[FilaPlantilla, ...]:
    return tuple(
        FilaPlantilla(
            valores=_fila(unidad=v, descripcion=v, detalle=v), texto_errores=v
        )
        for v in PELIGROSOS
    )


def test_f036_r7_ninguna_celda_de_ninguna_hoja_es_formula() -> None:
    wb = abrir(
        generar(
            _catalogo_peligroso(),
            filas=_filas_peligrosas(),
            importacion_origen=IMPORTACION,
        )
    )
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                assert celda.data_type != "f", f"{ws.title}!{celda.coordinate}"


def test_f036_r7_ningun_xml_de_hoja_lleva_formulas() -> None:
    contenido = generar(
        _catalogo_peligroso(), filas=_filas_peligrosas(), importacion_origen=IMPORTACION
    )
    with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
        hojas = [n for n in libro.namelist() if n.startswith("xl/worksheets/sheet")]
        assert len(hojas) == 4
        for nombre in hojas:
            assert not re.search(rb"<f[ >]", libro.read(nombre)), nombre


def test_f036_r7_lo_que_empieza_por_igual_se_conserva_como_texto() -> None:
    wb = abrir(
        generar(
            _catalogo_peligroso(),
            filas=_filas_peligrosas(),
            importacion_origen=IMPORTACION,
        )
    )
    ws = wb["Incidencias"]
    for numero, valor in enumerate(PELIGROSOS, start=2):
        for letra in ("A", "C", "D", "I"):
            celda = ws[f"{letra}{numero}"]
            assert celda.value == valor
            assert celda.data_type == "s"
    assert _lista_del_desplegable(wb, "Unidad") == tuple(sorted_unidades_peligrosas())
    assert _lista_del_desplegable(wb, "Oficio") == ("=Carpintería",)
    assert _lista_del_desplegable(wb, "Proveedor") == (
        "=Carpintería · @Proveedor Ejemplo",
    )
    assert wb["Instrucciones"]["A1"].value == "0677 · =Obra"


def sorted_unidades_peligrosas() -> list[str]:
    return [o.etiqueta for o in etiquetas_de_unidades(_catalogo_peligroso().unidades)]


# --------------------------------------------------------------------------
# R12 · obra sin oficios; R92 · un oficio por grupo
# --------------------------------------------------------------------------


def _solo_en_blanco(ws, columna: str) -> None:
    dv = _validacion(ws, columna)
    assert dv.type == "textLength"
    assert dv.operator == "equal"
    assert dv.formula1 == "0"
    assert dv.errorStyle == "stop"
    assert dv.showErrorMessage is True


def test_f036_r12_obra_sin_oficios_oficio_y_proveedor_sin_opciones() -> None:
    wb = abrir(generar(catalogo(oficios=(), filas=())))
    ws = wb["Incidencias"]
    _solo_en_blanco(ws, "Oficio")
    _solo_en_blanco(ws, "Proveedor")
    assert configuracion().textos.sin_oficios in _textos_de(wb["Instrucciones"])
    # Los demás desplegables siguen.
    assert _lista_del_desplegable(wb, "Unidad") == _esperadas("Unidad")


def test_f036_r12_oficios_sin_ningun_proveedor_solo_proveedor_en_blanco() -> None:
    cat = catalogo(filas=(ProveedorEnObra("0133", None, None),))
    wb = abrir(generar(cat))
    _solo_en_blanco(wb["Incidencias"], "Proveedor")
    assert _lista_del_desplegable(wb, "Oficio") == _esperadas("Oficio", cat)
    assert configuracion().textos.sin_oficios not in _textos_de(wb["Instrucciones"])


def test_f036_r12_obra_sin_unidades_unidad_sin_opciones() -> None:
    # La aplicación da antes 404 (R10); el generador no inventa una lista.
    wb = abrir(generar(catalogo(unidades=())))
    _solo_en_blanco(wb["Incidencias"], "Unidad")


def test_f036_r92_una_entrada_por_grupo_de_oficio() -> None:
    decisiones = (mismo("0085", "0166"),)
    wb = abrir(generar(decisiones=decisiones))
    oficios = _lista_del_desplegable(wb, "Oficio")
    assert oficios == ("Carpintería de madera", "Mamparas", "Albañilería")
    assert oficios == _esperadas("Oficio", decisiones=decisiones)


def test_f036_r92_sin_grupos_un_oficio_por_codigo() -> None:
    wb = _plantilla()
    assert _lista_del_desplegable(wb, "Oficio") == (
        "Carpintería de madera",
        "Mamparas",
        "Albañilería",
        "Mampara",
    )


def test_f036_r71_r92_los_pares_empiezan_por_la_etiqueta_del_oficio() -> None:
    decisiones = (mismo("0085", "0166"),)
    wb = abrir(generar(decisiones=decisiones))
    oficios = _lista_del_desplegable(wb, "Oficio")
    pares = _lista_del_desplegable(wb, "Proveedor")
    assert pares == _esperadas("Proveedor", decisiones=decisiones)
    assert pares == (
        "Carpintería de madera · Carpinterías Ejemplo S.L.",
        "Carpintería de madera · Juan Ejemplo Ejemplo",
        "Mamparas · Mamparas Ejemplo S.A.",
    )
    assert all(p.split(" · ")[0] in oficios for p in pares)


# --------------------------------------------------------------------------
# R62–R65 · el Excel de errores
# --------------------------------------------------------------------------


def test_f036_r62_mismo_libro_que_la_plantilla_con_los_mismos_desplegables() -> None:
    plantilla = _plantilla()
    errores = _errores()
    assert errores.sheetnames == HOJAS
    assert tuple(c.value for c in errores["Incidencias"][1]) == CABECERA
    for columna in COLUMNAS_TASADAS:
        assert _lista_del_desplegable(errores, columna) == _lista_del_desplegable(
            plantilla, columna
        )
        dv = _validacion(errores["Incidencias"], columna)
        assert dv.errorStyle == "stop"


def test_f036_r62_mismos_metadatos_mas_la_importacion_de_origen() -> None:
    plantilla = _metadatos(_plantilla())
    errores = _metadatos(_errores())
    assert errores == {**plantilla, "importacion_origen": str(IMPORTACION)}


def test_f036_r62_solo_las_filas_dadas_en_su_orden() -> None:
    ws = _errores()["Incidencias"]
    assert ws["C2"].value == "Sellar sifón"
    assert ws["C3"].value == "=SUMA(A1:A3)"
    assert all(c.value is None for fila in ws.iter_rows(min_row=4) for c in fila)


def test_f036_r63_valores_tal_como_llegaron_y_texto_de_errores() -> None:
    ws = _errores()["Incidencias"]
    fila_2 = {CABECERA[i]: c.value for i, c in enumerate(ws[2])}
    fila_3 = {CABECERA[i]: c.value for i, c in enumerate(ws[3])}
    assert fila_2 == {
        **dict.fromkeys(CABECERA),
        "Unidad": "villa 1",
        "Descripción corta": "Sellar sifón",
        "Oficio": "Mamparas",
        COLUMNA_ERRORES: FILAS_CON_ERROR[0].texto_errores,
    }
    assert fila_3["Detalle"] == " con espacios \n y salto "
    assert fila_3["Urgencia"] == "urgente"
    assert fila_3[COLUMNA_ERRORES] == FILAS_CON_ERROR[1].texto_errores
    assert ws["C3"].data_type == "s"


def test_f036_r63_la_columna_errores_sigue_bloqueada_en_el_excel_de_errores() -> None:
    ws = _errores()["Incidencias"]
    assert ws["I2"].protection.locked is True
    assert ws["A2"].protection.locked is False
    assert ws["I2"].alignment.wrap_text is True


def test_f036_r64_celdas_con_error_marcadas_con_relleno_y_comentario() -> None:
    ws = _errores()["Incidencias"]
    marcadas = {"A2": FILAS_CON_ERROR[0].errores["Unidad"]}
    marcadas["C3"] = FILAS_CON_ERROR[1].errores["Descripción corta"]
    marcadas["G3"] = FILAS_CON_ERROR[1].errores["Urgencia"]
    for coordenada, problema in marcadas.items():
        celda = ws[coordenada]
        assert celda.fill.fill_type == "solid", coordenada
        assert celda.comment is not None, coordenada
        assert celda.comment.text == problema
    for fila in ws.iter_rows(min_row=2, max_row=3):
        for celda in fila:
            if celda.coordinate not in marcadas:
                assert celda.comment is None, celda.coordinate
                # Enmienda 10 bis (R123): ninguna celda sin error lleva relleno
                # fijo, tampoco las de `Errores` (su gris va por la regla 2).
                assert celda.fill.fill_type is None, celda.coordinate


def test_f036_r64_las_marcadas_siguen_desbloqueadas_y_como_texto() -> None:
    ws = _errores()["Incidencias"]
    assert ws["A2"].protection.locked is False
    assert ws["A2"].number_format == "@"


def test_f036_r65_el_excel_de_errores_viaja_en_base64_y_se_reabre() -> None:
    contenido = _excel_errores()
    assert contenido.startswith(b"PK\x03\x04")
    ida_y_vuelta = base64.b64decode(base64.b64encode(contenido))
    assert abrir(ida_y_vuelta)["Incidencias"]["C2"].value == "Sellar sifón"


def test_f036_s3_5_el_excel_de_errores_explica_que_trae_en_instrucciones() -> None:
    presentes = _textos_de(_errores()["Instrucciones"])
    assert configuracion().textos.excel_errores in presentes


def test_f036_s5_2_filas_sin_origen_son_la_migracion_sin_parrafo_de_errores() -> None:
    wb = abrir(generar(filas=FILAS_CON_ERROR))
    assert "importacion_origen" not in _metadatos(wb)
    assert configuracion().textos.excel_errores not in _textos_de(wb["Instrucciones"])
    assert wb["Incidencias"]["C2"].value == "Sellar sifón"


@pytest.mark.parametrize(
    "fila",
    [
        FilaPlantilla(valores={"Columna rara": "x"}),
        FilaPlantilla(valores={COLUMNA_ERRORES: "x"}),
        FilaPlantilla(valores=_fila(unidad="x"), errores={"Columna rara": "x"}),
        FilaPlantilla(valores=_fila(unidad="x"), errores={COLUMNA_ERRORES: "x"}),
    ],
    ids=[
        "valor-columna-rara",
        "valor-en-errores",
        "error-columna-rara",
        "error-en-errores",
    ],
)
def test_f036_s5_2_una_columna_que_no_es_de_la_plantilla_se_rechaza(fila) -> None:
    with pytest.raises(ValueError, match="columna"):
        generar(filas=(fila,))


def test_f036_s5_2_no_mas_filas_que_el_tope() -> None:
    filas = tuple(
        FilaPlantilla(valores=_fila(unidad="x")) for _ in range(MAX_FILAS + 1)
    )
    with pytest.raises(ValueError, match="filas"):
        generar(filas=filas)


def test_f036_s5_2_mil_filas_caben() -> None:
    filas = tuple(
        FilaPlantilla(valores=_fila(unidad=f"u{i}")) for i in range(MAX_FILAS)
    )
    ws = abrir(generar(filas=filas))["Incidencias"]
    assert ws["A1001"].value == f"u{MAX_FILAS - 1}"


def test_f036_s5_2_fila_plantilla_inmutable_y_sin_errores_por_defecto() -> None:
    fila = FilaPlantilla(valores=_fila(unidad="x"))
    assert dict(fila.errores) == {}
    assert fila.texto_errores is None
    with pytest.raises(dataclasses.FrozenInstanceError):
        fila.texto_errores = "y"  # type: ignore[misc]
    assert dataclasses.fields(FilaPlantilla)[0].name == "valores"


def test_f036_s5_2_la_misma_entrada_da_el_mismo_contenido() -> None:
    # Determinista salvo lo que ponga openpyxl: las hojas son idénticas.
    def hojas(contenido: bytes) -> dict[str, bytes]:
        with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
            return {
                n: libro.read(n)
                for n in libro.namelist()
                if n.startswith("xl/worksheets/")
            }

    assert hojas(generar()) == hojas(generar())


# --------------------------------------------------------------------------
# Disposición (§3.1): lo que la mutación vio sin vigilar
# --------------------------------------------------------------------------

# Décima enmienda (R122): eran 30, 24, 50, 60, 28, 45, 24, 18 y 60.
ANCHOS = {
    "A": 30,
    "B": 26,
    "C": 48,
    "D": 56,
    "E": 28,
    "F": 40,
    "G": 24,
    "H": 16,
    "I": 50,
}
LISTAS = (
    "lista_unidades",
    "lista_ubicaciones",
    "lista_oficios",
    "lista_proveedores",
    "lista_urgencias",
    "lista_listados",
)


def test_f036_s3_1_anchos_de_las_columnas_de_incidencias() -> None:
    ws = _plantilla()["Incidencias"]
    assert {letra: ws.column_dimensions[letra].width for letra in ANCHOS} == ANCHOS


def test_f036_r7_los_textos_llevan_el_prefijo_de_comilla() -> None:
    # Si alguien edita la celda en Excel, sigue siendo texto y no una fórmula.
    wb = _errores()
    assert wb["Incidencias"]["C3"].quotePrefix is True
    assert wb["Incidencias"]["A1"].quotePrefix is True
    assert wb["_catalogos"]["A2"].quotePrefix is True
    assert wb["_plantilla"]["B3"].quotePrefix is True


def test_f036_s3_1_catalogos_una_columna_por_lista_con_su_nombre() -> None:
    wb = _plantilla()
    assert tuple(c.value for c in wb["_catalogos"][1]) == LISTAS
    for indice, nombre in enumerate(LISTAS):
        ((hoja, rango),) = list(wb.defined_names[nombre].destinations)
        letra = get_column_letter(indice + 1)
        assert hoja == "_catalogos"
        assert rango.startswith(f"${letra}$2:${letra}$")


def test_f036_r64_el_comentario_tiene_sitio_para_leerse() -> None:
    with zipfile.ZipFile(io.BytesIO(_excel_errores())) as libro:
        (vml,) = [libro.read(n) for n in libro.namelist() if n.endswith(".vml")]
    assert vml.decode("utf-8").count("width:300px;height:120px") == 3


def test_f036_r5_disposicion_de_instrucciones() -> None:
    ws = _plantilla()["Instrucciones"]
    textos = configuracion().textos
    columna_a = [ws.cell(row=f, column=1).value for f in range(1, ws.max_row + 1)]
    assert columna_a == [
        "0677 · Obra Ejemplo",
        textos.titulo,
        *textos.instrucciones,
        None,
        "Unidad",
        *[e.get("Unidad") for e in textos.ejemplos],
    ]
    parrafos = len(textos.instrucciones) + 2
    assert all(ws.cell(row=f, column=2).value is None for f in range(1, parrafos + 2))
    # Décima enmienda (R124): el título de la obra pasa de 14 a 15 puntos.
    assert (ws["A1"].font.bold, ws["A1"].font.size) == (True, 15)
    assert ws["A2"].font.bold is True
    assert ws["B2"].font.bold is not True
    assert ws["A3"].font.bold is not True
    cabecera_ejemplos = parrafos + 2
    assert all(ws.cell(row=cabecera_ejemplos, column=c).font.bold for c in range(1, 9))
    assert ws.cell(row=cabecera_ejemplos + 1, column=1).font.bold is not True


def test_f036_r5_anchos_de_instrucciones() -> None:
    ws = _plantilla()["Instrucciones"]
    assert ws.column_dimensions["A"].width == 100
    # Décima enmienda (R124): de la B a la H, 28 (eran 30).
    assert [ws.column_dimensions[get_column_letter(c)].width for c in range(2, 9)] == [
        28
    ] * 7
    assert ws.column_dimensions["I"].width != 28


# --------------------------------------------------------------------------
# R121–R125 · el formato visual (décima enmienda; `design.md` §3.6)
#
# Los valores van **literales**, copiados de las tablas de §3.6: si el
# generador cambia un color, un tamaño o un ancho, el test tiene que enterarse.
# --------------------------------------------------------------------------

SIN_LADO = (None, None)
LADOS = ("izquierdo", "derecho", "superior", "inferior")
BORDES_DE_CABECERA = {
    "izquierdo": ("thin", "FFFFFFFF"),
    "derecho": ("thin", "FFFFFFFF"),
    "superior": SIN_LADO,
    "inferior": ("medium", "FF7A1E33"),
}
BORDES_DEL_CUERPO = dict.fromkeys(LADOS, ("thin", "FFDFE2E4"))
SIN_BORDES = dict.fromkeys(LADOS, SIN_LADO)
LETRAS_DE_DATOS = "ABCDEFGH"
#: Las columnas con ajuste de texto en el cuerpo de «Incidencias».
LETRAS_AJUSTADAS = "CDI"
LIBROS = pytest.mark.parametrize(
    "libro", [_plantilla, _errores], ids=["plantilla", "errores"]
)


def _reglas(ws) -> list[tuple[str, list]]:
    """El rango y las reglas de cada formato condicional de la hoja."""
    return [(str(cf.sqref), list(cf.rules)) for cf in ws.conditional_formatting]


def _generar_con(textos, *, filas=(), importacion_origen=None) -> bytes:
    """Como `generar`, con otros textos (párrafos y ejemplos de prueba)."""
    cat = catalogo()
    oficios, pares = opciones(cat)
    return GeneradorPlantillaOpenpyxl().generar(
        catalogo=cat,
        opciones_oficio=oficios,
        opciones_proveedor=pares,
        listas=configuracion().listas,
        textos=textos,
        generada_at=INSTANTE,
        filas=filas,
        importacion_origen=importacion_origen,
    )


@LIBROS
@pytest.mark.parametrize("letra", LETRAS_DE_DATOS + "I")
def test_f036_r121_cabecera_de_incidencias(libro, letra: str) -> None:
    celda = libro()["Incidencias"][f"{letra}1"]
    # Burdeos en las columnas de datos; gris acero en `Errores`.
    color = "FF7B868C" if letra == "I" else "FF9F2842"
    assert _relleno(celda) == ("solid", color)
    assert _fuente(celda) == ("Calibri", 11, True, "FFFFFFFF")
    assert _alineacion(celda) == ("left", "center", True, 1)
    assert _bordes(celda) == BORDES_DE_CABECERA


def test_f036_r121_la_fila_de_la_cabecera_mide_34_puntos() -> None:
    for libro in (_plantilla(), _errores()):
        ws = libro["Incidencias"]
        assert ws.row_dimensions[1].height == 34
        # Solo la cabecera: las filas de datos se quedan con el alto por defecto.
        assert ws.row_dimensions[2].height is None
        assert ws.row_dimensions[1001].height is None


def test_f036_r121_el_texto_de_la_cabecera_no_cambia() -> None:
    assert tuple(c.value for c in _plantilla()["Incidencias"][1]) == (
        "Unidad",
        "Ubicación",
        "Descripción corta",
        "Detalle",
        "Oficio",
        "Proveedor",
        "Urgencia",
        "Listado",
        "Errores (lo rellena el sistema)",
    )


@pytest.mark.parametrize("fila", [2, 3, 500, 1001])
@pytest.mark.parametrize("letra", LETRAS_DE_DATOS)
def test_f036_r122_cuerpo_de_incidencias(letra: str, fila: int) -> None:
    celda = _plantilla()["Incidencias"][f"{letra}{fila}"]
    assert _fuente(celda) == ("Calibri", 10.5, False, "FF1D2024")
    # Enmienda 10 bis: ni bordes ni relleno fijos; las líneas y las bandas van
    # por formato condicional y solo se ven en las filas pintadas.
    assert _bordes(celda) == SIN_BORDES
    assert _alineacion(celda) == (None, "top", letra in LETRAS_AJUSTADAS, 1)
    assert _relleno(celda) == (None, None)


FILAS_RELLENAS = (
    FilaPlantilla(valores=_fila(unidad="Viviendas Bloque Villa 1", descripcion="Roza")),
    FilaPlantilla(valores=_fila(unidad="Viviendas Bloque Villa 2", detalle="Gotea")),
)


@functools.cache
def _rellena():
    """Una plantilla con dos filas sin errores, sin importación de origen."""
    return abrir(generar(filas=FILAS_RELLENAS))


#: Las celdas con error del Excel de errores (FILAS_CON_ERROR).
CON_ERROR = frozenset({"A2", "C3", "G3"})


@pytest.mark.parametrize(
    "libro", [_plantilla, _rellena, _errores], ids=["vacia", "rellena", "errores"]
)
@pytest.mark.parametrize("fila", [2, 3, 500, 1001])
@pytest.mark.parametrize("letra", ["A", "H", "I"])
def test_f036_r122_sin_bordes_ni_relleno_estaticos_en_el_cuerpo(
    libro, fila: int, letra: str
) -> None:
    # Enmienda 10 bis (R122, R123): lo que se ve del cuerpo va por formato
    # condicional; el estilo de la celda solo lleva lo que no se ve vacío.
    coordenada = f"{letra}{fila}"
    celda = libro()["Incidencias"][coordenada]
    assert _bordes(celda) == SIN_BORDES, coordenada
    if libro is _errores and coordenada in CON_ERROR:
        # Salvo el rojo de la celda con error, que sigue siendo estático.
        assert _relleno(celda) == ("solid", "FFFEF2F2")
        assert _fuente(celda) == ("Calibri", 10.5, False, "FFB91C1C")
    else:
        assert celda.fill.fill_type is None, coordenada
        tinta = "FF5D676D" if letra == "I" else "FF1D2024"
        assert _fuente(celda) == ("Calibri", 10.5, False, tinta), coordenada
    # Lo estático que sigue: alineación con su ajuste, y el desbloqueo y `@`
    # de las columnas de datos (R126), en las 1.000 filas.
    assert _alineacion(celda) == (None, "top", letra in LETRAS_AJUSTADAS, 1)
    if letra == "I":
        assert celda.protection.locked is True
        assert celda.number_format == "General"
    else:
        assert celda.protection.locked is False
        assert celda.number_format == "@"


@pytest.mark.parametrize(
    "libro", [_plantilla, _rellena, _errores], ids=["vacia", "rellena", "errores"]
)
def test_f036_r122_ninguna_celda_del_cuerpo_lleva_borde_relleno_ni_fuente_de_mas(
    libro,
) -> None:
    # T60 (mutación de la enmienda 10 bis): el test de arriba muestrea las
    # columnas A, H e I; un estilo fijo puesto solo en B–G de las filas con
    # datos se le escapaba. Aquí se recorren las 9.000 celdas del cuerpo: sin
    # borde, relleno ni fuente de más, y con lo estático que se conserva.
    ws = libro()["Incidencias"]
    distintas = []
    for fila in ws.iter_rows(min_row=2, max_row=1001, max_col=9):
        for celda in fila:
            letra = celda.column_letter
            if libro is _errores and celda.coordinate in CON_ERROR:
                aspecto = (("solid", "FFFEF2F2"), ("Calibri", 10.5, False, "FFB91C1C"))
            else:
                tinta = "FF5D676D" if letra == "I" else "FF1D2024"
                aspecto = ((None, None), ("Calibri", 10.5, False, tinta))
            fijo = (
                (None, "top", letra in LETRAS_AJUSTADAS, 1),
                letra == "I",
                "General" if letra == "I" else "@",
            )
            visto = (
                _bordes(celda),
                (_relleno(celda), _fuente(celda)),
                (_alineacion(celda), celda.protection.locked, celda.number_format),
            )
            if visto != (SIN_BORDES, aspecto, fijo):
                distintas.append(celda.coordinate)
    assert distintas == []


def test_f036_r122_el_formato_no_crea_celdas_fuera_de_la_plantilla() -> None:
    # Libros recién abiertos: en los compartidos, leer `ws["A1002"]` crea la celda.
    for libro in (abrir(generar()), abrir(_excel_errores())):
        ws = libro["Incidencias"]
        assert (ws.max_row, ws.max_column) == (1001, 9)
    for contenido in (generar(), _excel_errores()):
        with zipfile.ZipFile(io.BytesIO(contenido)) as libro_zip:
            xml = libro_zip.read("xl/worksheets/sheet1.xml")
        assert xml.count(b"<row ") == 1001
        assert xml.count(b"<c ") == 1001 * 9


@pytest.mark.parametrize("hoja", ["Incidencias", "Instrucciones"])
def test_f036_r122_r124_las_hojas_visibles_no_ensenan_la_cuadricula(hoja: str) -> None:
    for libro in (_plantilla(), _errores()):
        assert libro[hoja].sheet_view.showGridLines is False


def test_f036_r122_r124_las_hojas_ocultas_no_llevan_formato() -> None:
    wb = _plantilla()
    for hoja in ("_catalogos", "_plantilla"):
        ws = wb[hoja]
        assert ws.sheet_view.showGridLines is not False
        assert ws.sheet_properties.tabColor is None
        assert _relleno(ws["A1"]) == (None, None)
        assert _bordes(ws["A1"]) == SIN_BORDES
        assert _reglas(ws) == []


def test_f036_r122_r124_color_de_las_pestanas() -> None:
    # RGB de seis cifras; `openpyxl` lo guarda con el canal alfa a cero.
    for libro in (_plantilla(), _errores()):
        assert libro["Incidencias"].sheet_properties.tabColor.rgb == "009F2842"
        assert libro["Instrucciones"].sheet_properties.tabColor.rgb == "007B868C"


#: Enmienda 10 bis: una fila está pintada si es la 2 o si tiene algo de A a I.
FORMULA_PINTADA = "OR(ROW()=2,COUNTA($A2:$I2)>0)"
FORMULA_BANDAS = "AND(MOD(ROW(),2)=1,COUNTA($A2:$I2)>0)"


def _comprobar_reglas(ws, *, bandas: bool) -> None:
    """Las reglas del cuerpo (R122, R123), en su orden y con su formato.

    1 · líneas en `A2:I1001`; 2 · gris de `Errores` en `I2:I1001`; y, solo sin
    errores, 3 · bandas en `A2:H1001`. Cada una, un rango y una regla.
    """
    reglas = _reglas(ws)
    rangos = ["A2:I1001", "I2:I1001", *(["A2:H1001"] if bandas else [])]
    assert [rango for rango, _ in reglas] == rangos
    unicas = []
    for prioridad, (rango, lista) in enumerate(reglas, start=1):
        (regla,) = lista
        assert regla.type == "expression", rango
        assert regla.priority == prioridad, rango
        assert regla.stopIfTrue is None, rango
        unicas.append(regla)
    lineas, gris = unicas[0], unicas[1]
    # 1 · los cuatro lados finos en gris claro; sin relleno ni fuente.
    assert lineas.formula == ["OR(ROW()=2,COUNTA($A2:$I2)>0)"]
    borde = lineas.dxf.border
    assert {
        "izquierdo": _lado(borde.left),
        "derecho": _lado(borde.right),
        "superior": _lado(borde.top),
        "inferior": _lado(borde.bottom),
    } == BORDES_DEL_CUERPO
    assert lineas.dxf.fill is None
    assert lineas.dxf.font is None
    # 2 · el gris «no editable»; en un formato diferencial, el color de un
    # relleno sólido va en `bgColor`.
    assert gris.formula == ["OR(ROW()=2,COUNTA($A2:$I2)>0)"]
    assert (gris.dxf.fill.fill_type, gris.dxf.fill.bgColor.rgb) == ("solid", "FFF3F4F5")
    assert gris.dxf.border is None
    assert gris.dxf.font is None
    if bandas:
        banda = unicas[2]
        # 3 · las filas impares pintadas.
        assert banda.formula == ["AND(MOD(ROW(),2)=1,COUNTA($A2:$I2)>0)"]
        assert (banda.dxf.fill.fill_type, banda.dxf.fill.bgColor.rgb) == (
            "solid",
            "FFFAFAFB",
        )
        assert banda.dxf.border is None
        assert banda.dxf.font is None


def test_f036_r122_r123_la_plantilla_lleva_las_tres_reglas_condicionales() -> None:
    # Vacía y rellena (sin errores): las tres reglas, y ninguna en «Instrucciones».
    for libro in (_plantilla(), _rellena()):
        _comprobar_reglas(libro["Incidencias"], bandas=True)
        assert _reglas(libro["Instrucciones"]) == []


def test_f036_r122_el_v2_sin_errores_lleva_las_bandas() -> None:
    # La migración: filas sin errores y sin importación de origen.
    filas = (FilaPlantilla(valores=_fila(unidad="x")), FilaPlantilla(valores={}))
    _comprobar_reglas(abrir(generar(filas=filas))["Incidencias"], bandas=True)


def test_f036_r122_las_bandas_dependen_de_los_errores_y_no_del_origen() -> None:
    # Filas sin errores marcados, aunque el libro diga de qué importación sale.
    filas = (FilaPlantilla(valores=_fila(unidad="x"), texto_errores="Fila 7 · …"),)
    ws = abrir(generar(filas=filas, importacion_origen=IMPORTACION))["Incidencias"]
    _comprobar_reglas(ws, bandas=True)


#: Los `dxf` de las tres reglas, en su orden, tal como los escribe el
#: generador (§3.6): líneas finas, gris de `Errores` y banda.
DXF_DE_LAS_REGLAS = (
    (
        b'<dxf><border><left style="thin"><color rgb="FFDFE2E4" /></left>'
        b'<right style="thin"><color rgb="FFDFE2E4" /></right>'
        b'<top style="thin"><color rgb="FFDFE2E4" /></top>'
        b'<bottom style="thin"><color rgb="FFDFE2E4" /></bottom></border></dxf>'
    ),
    (
        b'<dxf><fill><patternFill patternType="solid"><bgColor rgb="FFF3F4F5" />'
        b"</patternFill></fill></dxf>"
    ),
    (
        b'<dxf><fill><patternFill patternType="solid"><bgColor rgb="FFFAFAFB" />'
        b"</patternFill></fill></dxf>"
    ),
)


def _canonico(xml: bytes) -> str:
    """El XML en forma canónica (C14N 2.0), sin espacios entre etiquetas."""
    return ElementTree.canonicalize(xml.decode("utf-8"), strip_text=True)


@pytest.mark.parametrize(
    ("contenido", "rangos"),
    [
        (generar, [b"A2:I1001", b"I2:I1001", b"A2:H1001"]),
        (lambda: generar(filas=FILAS_RELLENAS), [b"A2:I1001", b"I2:I1001", b"A2:H1001"]),
        (_excel_errores, [b"A2:I1001", b"I2:I1001"]),
    ],
    ids=["vacia", "rellena", "errores"],
)
def test_f036_r122_r123_las_reglas_en_el_xml(contenido, rangos: list[bytes]) -> None:
    with zipfile.ZipFile(io.BytesIO(contenido())) as libro:
        hoja = libro.read("xl/worksheets/sheet1.xml")
        estilos = libro.read("xl/styles.xml")
        otras = [libro.read(f"xl/worksheets/sheet{n}.xml") for n in (2, 3, 4)]
    # Un `<conditionalFormatting>` por regla, cada `sqref` de un solo rango.
    assert re.findall(rb'<conditionalFormatting sqref="([^"]*)"', hoja) == rangos
    assert hoja.count(b"<conditionalFormatting") == len(rangos)
    assert hoja.count(b"<cfRule") == len(rangos)
    # Tantos formatos diferenciales como reglas, y el bloque `<dxfs>` entero
    # (review 8, R8-1): cualquier atributo de más en un `dxf` (una diagonal,
    # un `fgColor`, un `numFmt`...) cae aquí. Se compara canonizado, para no
    # depender de los espacios ni del orden de los atributos.
    (dxfs,) = re.findall(rb"<dxfs\b.*?</dxfs>", estilos, re.DOTALL)
    esperado = b'<dxfs count="%d">' % len(rangos) + b"".join(DXF_DE_LAS_REGLAS[: len(rangos)])
    assert _canonico(dxfs) == _canonico(esperado + b"</dxfs>")
    assert all(b"<conditionalFormatting" not in xml for xml in otras)


SIN_ERROR = FilaPlantilla(valores=_fila(unidad="Viviendas Bloque Villa 1"))


@pytest.mark.parametrize(
    ("filas", "origen"),
    [
        (FILAS_CON_ERROR, IMPORTACION),
        (FILAS_CON_ERROR, None),
        ((SIN_ERROR, FILAS_CON_ERROR[0]), IMPORTACION),
        ((FILAS_CON_ERROR[0], SIN_ERROR), IMPORTACION),
        ((SIN_ERROR, FILAS_CON_ERROR[0], SIN_ERROR), None),
    ],
    ids=["errores", "errores-sin-origen", "la-ultima", "la-primera", "la-de-en-medio"],
)
def test_f036_r122_con_alguna_fila_con_error_no_hay_regla_de_bandas(
    filas, origen
) -> None:
    # El relleno de la regla de bandas taparía el rojo de la celda con error
    # (R123): solo las reglas 1 (líneas) y 2 (gris de `Errores`).
    wb = abrir(generar(filas=filas, importacion_origen=origen))
    _comprobar_reglas(wb["Incidencias"], bandas=False)
    assert _reglas(wb["Instrucciones"]) == []


@LIBROS
@pytest.mark.parametrize("fila", [2, 3, 500, 1001])
def test_f036_r123_la_columna_errores_va_en_gris_no_editable(libro, fila: int) -> None:
    ws = libro()["Incidencias"]
    celda = ws[f"I{fila}"]
    # Enmienda 10 bis: el gris va por la regla 2, no como relleno fijo; la
    # fuente gris, la alineación y el bloqueo siguen siendo estáticos.
    assert _relleno(celda) == (None, None)
    assert _bordes(celda) == SIN_BORDES
    assert _fuente(celda) == ("Calibri", 10.5, False, "FF5D676D")
    assert _alineacion(celda) == (None, "top", True, 1)
    assert celda.protection.locked is True
    rango, (regla,) = _reglas(ws)[1]
    assert rango == "I2:I1001"
    assert regla.formula == ["OR(ROW()=2,COUNTA($A2:$I2)>0)"]
    assert regla.dxf.fill.bgColor.rgb == "FFF3F4F5"
    # Frontera con la última columna de datos: ni el gris ni el bloqueo.
    vecina = ws[f"H{fila}"]
    assert _relleno(vecina) == (None, None)
    assert _fuente(vecina) == ("Calibri", 10.5, False, "FF1D2024")
    assert vecina.protection.locked is False


@pytest.mark.parametrize("coordenada", ["A2", "C3", "G3"])
def test_f036_r123_celda_con_error_en_rojo_suave_con_el_texto_en_rojo(
    coordenada: str,
) -> None:
    celda = _errores()["Incidencias"][coordenada]
    assert _relleno(celda) == ("solid", "FFFEF2F2")
    assert _fuente(celda) == ("Calibri", 10.5, False, "FFB91C1C")
    # Lo demás, como cualquier celda del cuerpo en su columna; sin borde fijo
    # (enmienda 10 bis): sus líneas son las de la regla 1.
    assert _bordes(celda) == SIN_BORDES
    assert _alineacion(celda) == (None, "top", coordenada[0] in LETRAS_AJUSTADAS, 1)
    assert celda.protection.locked is False
    assert celda.number_format == "@"
    assert celda.quotePrefix is True
    assert celda.comment is not None


@pytest.mark.parametrize("coordenada", ["B2", "C2", "A3", "H3", "A4"])
def test_f036_r123_la_vecina_sin_error_no_lleva_ni_relleno_ni_rojo(
    coordenada: str,
) -> None:
    celda = _errores()["Incidencias"][coordenada]
    assert _relleno(celda) == (None, None)
    assert _fuente(celda) == ("Calibri", 10.5, False, "FF1D2024")
    # Enmienda 10 bis: sin borde fijo, como todo el cuerpo (R122).
    assert _bordes(celda) == SIN_BORDES
    assert celda.comment is None


def _zonas_de_instrucciones() -> dict[str, int]:
    """Dónde cae cada zona de «Instrucciones» en la plantilla por defecto."""
    parrafos = len(configuracion().textos.instrucciones)
    return {
        "primer_parrafo": 3,
        "en_blanco": 3 + parrafos,
        "cabecera_ejemplos": 4 + parrafos,
    }


@LIBROS
def test_f036_r124_fila_1_la_obra_en_una_banda_burdeos(libro) -> None:
    ws = libro()["Instrucciones"]
    assert ws.row_dimensions[1].height == 38
    for letra in LETRAS_DE_DATOS:
        celda = ws[f"{letra}1"]
        assert _relleno(celda) == ("solid", "FF9F2842"), letra
        assert _fuente(celda) == ("Calibri", 15, True, "FFFFFFFF"), letra
        assert _alineacion(celda) == (None, "center", False, 1), letra
        assert _bordes(celda) == SIN_BORDES, letra
    # Solo estilo: el texto de la obra sigue en A1 y nada más (R126).
    assert ws["A1"].value == "0677 · Obra Ejemplo"
    assert [ws[f"{letra}1"].value for letra in "BCDEFGH"] == [None] * 7
    assert _relleno(ws["I1"]) == (None, None)


@LIBROS
def test_f036_r124_fila_2_el_titulo_en_burdeos_con_su_linea(libro) -> None:
    ws = libro()["Instrucciones"]
    assert ws.row_dimensions[2].height == 28
    assert ws["A2"].value == configuracion().textos.titulo
    assert _fuente(ws["A2"]) == ("Calibri", 13, True, "FF9F2842")
    assert _alineacion(ws["A2"]) == (None, "center", False, 1)
    linea = {**SIN_BORDES, "inferior": ("medium", "FF9F2842")}
    for letra in LETRAS_DE_DATOS:
        celda = ws[f"{letra}2"]
        assert _bordes(celda) == linea, letra
        assert _relleno(celda) == (None, None), letra
    assert [ws[f"{letra}2"].value for letra in "BCDEFGH"] == [None] * 7
    assert _bordes(ws["I2"]) == SIN_BORDES


def test_f036_r124_parrafos_en_calibri_11_con_el_alto_calculado() -> None:
    ws = _plantilla()["Instrucciones"]
    zonas = _zonas_de_instrucciones()
    textos = configuracion().textos
    for fila, parrafo in enumerate(textos.instrucciones, start=zonas["primer_parrafo"]):
        celda = ws.cell(row=fila, column=1)
        assert celda.value == parrafo
        assert _fuente(celda) == ("Calibri", 11, False, "FF1D2024"), fila
        assert _alineacion(celda) == (None, "top", True, 1), fila
        assert _relleno(celda) == (None, None), fila
        assert _bordes(celda) == SIN_BORDES, fila
        alto = max(20, 16 * (len(parrafo) // 105 + 1) + 6)
        assert ws.row_dimensions[fila].height == alto, fila
    # La fila en blanco que separa los párrafos de los ejemplos, sin alto propio.
    assert ws.row_dimensions[zonas["en_blanco"]].height is None
    assert ws.cell(row=zonas["en_blanco"], column=1).value is None


def test_f036_r124_el_parrafo_del_excel_de_errores_y_el_de_sin_oficios_tambien() -> None:
    textos = configuracion().textos
    wb = abrir(
        generar(
            catalogo(oficios=(), filas=()),
            filas=FILAS_CON_ERROR,
            importacion_origen=IMPORTACION,
        )
    )
    ws = wb["Instrucciones"]
    ultima = 3 + len(textos.instrucciones) + 1
    assert ws["A3"].value == textos.excel_errores
    assert ws.cell(row=ultima, column=1).value == textos.sin_oficios
    for fila in (3, ultima):
        celda = ws.cell(row=fila, column=1)
        assert _fuente(celda) == ("Calibri", 11, False, "FF1D2024"), fila
        assert _alineacion(celda) == (None, "top", True, 1), fila
        alto = 16 * (len(celda.value) // 105 + 1) + 6
        assert ws.row_dimensions[fila].height == alto, fila
    assert ws.row_dimensions[ultima + 1].height is None
    # La cabecera de los ejemplos baja con los párrafos, y su alto con ella.
    assert ws.row_dimensions[ultima + 2].height == 26
    assert ws.cell(row=ultima + 2, column=1).value == "Unidad"


@pytest.mark.parametrize(
    ("largo", "alto"),
    [
        (0, 22),
        (1, 22),
        (104, 22),
        (105, 38),
        (106, 38),
        (209, 38),
        (210, 54),
        (1050, 182),
    ],
)
def test_f036_r124_alto_de_un_parrafo_en_sus_fronteras(largo: int, alto: int) -> None:
    assert excel_openpyxl._alto_de_parrafo("x" * largo) == alto


def test_f036_r124_los_cuatro_numeros_del_alto_de_un_parrafo() -> None:
    # `max(20, 16 × (largo // 105 + 1) + 6)`: con estos números el mínimo no
    # llega a verse (una sola línea ya mide 22), así que se fija aquí.
    assert excel_openpyxl.ALTO_MINIMO_PARRAFO == 20
    assert excel_openpyxl.PUNTOS_POR_LINEA == 16
    assert excel_openpyxl.CARACTERES_POR_LINEA == 105
    assert excel_openpyxl.RESPIRO_PARRAFO == 6


def test_f036_r124_el_alto_del_parrafo_no_baja_del_minimo(monkeypatch) -> None:
    monkeypatch.setattr(excel_openpyxl, "PUNTOS_POR_LINEA", 2)
    assert excel_openpyxl._alto_de_parrafo("x") == 20
    assert excel_openpyxl._alto_de_parrafo("x" * 1050) == 28


def test_f036_r124_parrafo_corto_y_largo_en_el_libro() -> None:
    textos = dataclasses.replace(
        configuracion().textos, instrucciones=("x" * 104, "y" * 105, "z" * 210)
    )
    ws = abrir(_generar_con(textos))["Instrucciones"]
    assert [ws.row_dimensions[f].height for f in (3, 4, 5)] == [22, 38, 54]
    assert ws.row_dimensions[6].height is None


@LIBROS
def test_f036_r124_cabecera_de_los_ejemplos(libro) -> None:
    ws = libro()["Instrucciones"]
    fila = _zonas_de_instrucciones()["cabecera_ejemplos"]
    if libro is _errores:
        fila += 1  # el párrafo del Excel de errores
    assert ws.row_dimensions[fila].height == 26
    valores = tuple(ws.cell(row=fila, column=c).value for c in range(1, 9))
    assert valores == COLUMNAS_DE_DATOS
    for columna in range(1, 9):
        celda = ws.cell(row=fila, column=columna)
        assert _relleno(celda) == ("solid", "FF9F2842"), columna
        assert _fuente(celda) == ("Calibri", 11, True, "FFFFFFFF"), columna
        assert _bordes(celda) == BORDES_DE_CABECERA, columna
        assert _alineacion(celda) == (None, "center", False, 1), columna
    assert _relleno(ws.cell(row=fila, column=9)) == (None, None)


def test_f036_r124_filas_de_ejemplo_con_la_segunda_y_la_cuarta_en_banda() -> None:
    ejemplo = {"Unidad": "Vivienda 1", "Descripción corta": "La puerta roza"}
    textos = dataclasses.replace(configuracion().textos, ejemplos=(ejemplo,) * 5)
    ws = abrir(_generar_con(textos))["Instrucciones"]
    cabecera = _zonas_de_instrucciones()["cabecera_ejemplos"]
    rellenos = []
    for desplazamiento in range(1, 6):
        fila = cabecera + desplazamiento
        for columna in range(1, 9):
            celda = ws.cell(row=fila, column=columna)
            assert _fuente(celda) == ("Calibri", 10.5, False, "FF1D2024"), celda.coordinate
            assert _bordes(celda) == BORDES_DEL_CUERPO, celda.coordinate
            assert _alineacion(celda) == (None, "top", True, 1), celda.coordinate
        assert _relleno(ws.cell(row=fila, column=9)) == (None, None)
        assert _bordes(ws.cell(row=fila, column=9)) == SIN_BORDES
        por_columna = {_relleno(ws.cell(row=fila, column=c)) for c in range(1, 9)}
        assert len(por_columna) == 1, fila
        rellenos.append(por_columna.pop())
        assert ws.row_dimensions[fila].height is None
    banda = ("solid", "FFFAFAFB")
    assert rellenos == [(None, None), banda, (None, None), banda, (None, None)]
    assert ws.max_row == cabecera + 5


def test_f036_r124_los_ejemplos_del_yaml_con_su_banda() -> None:
    ws = _plantilla()["Instrucciones"]
    cabecera = _zonas_de_instrucciones()["cabecera_ejemplos"]
    assert len(configuracion().textos.ejemplos) == 3
    for columna in (1, 8):
        rellenos = [_relleno(ws.cell(row=cabecera + d, column=columna)) for d in (1, 2, 3)]
        assert rellenos == [(None, None), ("solid", "FFFAFAFB"), (None, None)]


@LIBROS
@pytest.mark.parametrize("hoja", HOJAS)
def test_f036_r125_ninguna_hoja_lleva_ajustes_de_impresion(libro, hoja: str) -> None:
    # Enmienda 10 bis: este Excel no se imprime. Se fija la **ausencia** de
    # todo ajuste, también de los que nadie pidió (R7-1: área, escala,
    # centrado, márgenes).
    ws = libro()[hoja]
    assert ws.print_area == ""
    assert ws.print_title_rows is None
    assert ws.print_title_cols is None
    assert list(ws.defined_names) == []
    pagina = ws.page_setup
    assert pagina.orientation is None
    assert pagina.paperSize is None
    assert pagina.scale is None
    assert pagina.fitToWidth is None
    assert pagina.fitToHeight is None
    # Review 8, R8-2: ni `fitToPage` ni ningún otro atributo de `pageSetUpPr`
    # (`autoPageBreaks`...). `openpyxl` crea siempre un `PageSetupProperties`
    # vacío al leer (y lo escribe vacío), así que «ausente» es «sin ningún
    # atributo puesto»; en el XML lo fija el test de abajo.
    ajuste = ws.sheet_properties.pageSetUpPr
    assert ajuste is None or dict(ajuste) == {}
    # Ni «Diseño de página» ni «Vista previa de salto de página»: la vista normal.
    assert ws.sheet_view.view in (None, "normal")
    # Review 8, R8-3: las hojas visibles enseñan los encabezados de fila y
    # columna, y el libro calcula en automático.
    if hoja in ("Incidencias", "Instrucciones"):
        assert ws.sheet_view.showRowColHeaders is not False
    assert ws.parent.calculation.calcMode in (None, "auto")
    opciones_de_impresion = ws.print_options
    assert opciones_de_impresion.horizontalCentered is None
    assert opciones_de_impresion.verticalCentered is None
    assert opciones_de_impresion.gridLines is None
    assert opciones_de_impresion.headings is None
    margenes = ws.page_margins
    assert (
        margenes.left,
        margenes.right,
        margenes.top,
        margenes.bottom,
        margenes.header,
        margenes.footer,
    ) == (0.75, 0.75, 1.0, 1.0, 0.5, 0.5)
    for nombre in (
        "oddHeader",
        "oddFooter",
        "evenHeader",
        "evenFooter",
        "firstHeader",
        "firstFooter",
    ):
        trozos = getattr(ws, nombre)
        assert (trozos.left.text, trozos.center.text, trozos.right.text) == (
            None,
            None,
            None,
        ), nombre
    assert not ws.row_breaks.brk
    assert not ws.col_breaks.brk


@pytest.mark.parametrize("contenido", [generar, _excel_errores], ids=["plantilla", "errores"])
def test_f036_r125_ni_las_hojas_ni_el_libro_llevan_impresion_en_el_xml(contenido) -> None:
    with zipfile.ZipFile(io.BytesIO(contenido())) as libro:
        hojas = {
            nombre: libro.read(nombre)
            for nombre in libro.namelist()
            if re.fullmatch(r"xl/worksheets/sheet\d\.xml", nombre)
        }
        cuaderno = libro.read("xl/workbook.xml")
    assert len(hojas) == 4
    for nombre, xml in hojas.items():
        for etiqueta in (
            b"<pageSetup",
            b"<printOptions",
            b"<headerFooter",
            b"<rowBreaks",
            b"<colBreaks",
            b"fitToPage",
        ):
            assert etiqueta not in xml, (nombre, etiqueta)
        # `openpyxl` escribe siempre un `<pageSetUpPr />` vacío; ningún
        # atributo en él (review 8, R8-2).
        assert set(re.findall(rb"<pageSetUpPr\b[^>]*>", xml)) <= {b"<pageSetUpPr />"}, nombre
    assert b"_xlnm.Print_" not in cuaderno
    # Los nombres definidos: los seis de las listas y el del filtro, nada más.
    nombres = re.findall(rb'<definedName name="([^"]+)"', cuaderno)
    assert sorted(nombres) == sorted(
        [n.encode() for n in LISTAS] + [b"_xlnm._FilterDatabase"]
    )


# --------------------------------------------------------------------------
# R126 · invariantes: el formato no cambia lo que lee el importador. Estos
# nacen en verde sobre el generador de antes de la enmienda y siguen en verde.
# --------------------------------------------------------------------------

PARTES_DE_LA_PLANTILLA = {
    "[Content_Types].xml",
    "_rels/.rels",
    "docProps/app.xml",
    "docProps/core.xml",
    "xl/_rels/workbook.xml.rels",
    "xl/styles.xml",
    "xl/theme/theme1.xml",
    "xl/workbook.xml",
    "xl/worksheets/sheet1.xml",
    "xl/worksheets/sheet2.xml",
    "xl/worksheets/sheet3.xml",
    "xl/worksheets/sheet4.xml",
}
#: Lo que añaden los comentarios de las celdas con error (R64): no es formato.
PARTES_DE_LOS_COMENTARIOS = {
    "xl/comments/comment1.xml",
    "xl/drawings/commentsDrawing1.vml",
    "xl/worksheets/_rels/sheet1.xml.rels",
}


def _partes(contenido: bytes) -> set[str]:
    with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
        nombres = libro.namelist()
    assert len(nombres) == len(set(nombres))
    return set(nombres)


def test_f036_r126_las_mismas_partes_en_el_zip_y_ninguna_imagen() -> None:
    assert _partes(generar()) == PARTES_DE_LA_PLANTILLA
    errores = _partes(_excel_errores())
    assert errores == PARTES_DE_LA_PLANTILLA | PARTES_DE_LOS_COMENTARIOS
    assert not [n for n in errores if "media" in n or n.endswith((".png", ".jpg", ".emf"))]


def test_f036_r126_las_mismas_cuatro_hojas_con_incidencias_activa() -> None:
    for wb in (_plantilla(), _errores()):
        assert wb.sheetnames == ["Incidencias", "Instrucciones", "_catalogos", "_plantilla"]
        assert [ws.sheet_state for ws in wb.worksheets] == [
            "visible",
            "visible",
            "veryHidden",
            "veryHidden",
        ]
        assert wb.active.title == "Incidencias"


def test_f036_r126_cabecera_en_la_fila_1_sin_celdas_combinadas_ni_desplazadas() -> None:
    # Libros recién abiertos: en los compartidos, leer `ws["A1002"]` crea la celda.
    for wb in (abrir(generar()), abrir(_excel_errores())):
        ws = wb["Incidencias"]
        assert tuple(c.value for c in ws[1]) == CABECERA
        assert (ws.min_row, ws.min_column, ws.max_row, ws.max_column) == (1, 1, 1001, 9)
        assert not ws.merged_cells.ranges
        assert not wb["Instrucciones"].merged_cells.ranges
    # Los datos siguen empezando en la fila 2.
    assert _errores()["Incidencias"]["C2"].value == "Sellar sifón"


def test_f036_r126_los_rangos_con_nombre_de_las_listas_resuelven_a_lo_mismo() -> None:
    for wb in (_plantilla(), _errores()):
        for indice, nombre in enumerate(LISTAS):
            ((hoja, rango),) = list(wb.defined_names[nombre].destinations)
            letra = get_column_letter(indice + 1)
            assert hoja == "_catalogos"
            assert rango.startswith(f"${letra}$2:${letra}$")
        # Los seis de las listas y ninguno más a nivel de libro.
        assert set(wb.defined_names) == set(LISTAS)


def test_f036_r126_validaciones_proteccion_filtro_paneles_y_formato_de_texto() -> None:
    for wb in (_plantilla(), _errores()):
        ws = wb["Incidencias"]
        validaciones = ws.data_validations.dataValidation
        assert sorted(str(dv.sqref) for dv in validaciones) == [
            f"{letra}2:{letra}1001" for letra in LETRAS_DE_DATOS
        ]
        assert all(dv.errorStyle == "stop" for dv in validaciones)
        assert ws.protection.sheet is True
        assert not ws.protection.password
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:I1001"
        for fila in (2, 3, 1001):
            for letra in LETRAS_DE_DATOS:
                celda = ws[f"{letra}{fila}"]
                assert celda.number_format == "@", celda.coordinate
                assert celda.protection.locked is False, celda.coordinate
            assert ws[f"I{fila}"].protection.locked is True
            assert ws[f"I{fila}"].number_format == "General"
        assert all(c.protection.locked is True for c in ws[1])
        assert wb["Instrucciones"].protection.sheet is False


def test_f036_r126_misma_version_e_identificador() -> None:
    for wb in (_plantilla(), _errores()):
        metadatos = _metadatos(wb)
        assert metadatos["version"] == 1
        assert metadatos["identificador"] == IDENTIFICADOR_PLANTILLA


def test_f036_r126_ningun_xml_de_hoja_lleva_formulas_de_celda() -> None:
    # La fórmula de la regla condicional va en `<formula>`, no en un `<f>`.
    for contenido in (generar(), _excel_errores()):
        with zipfile.ZipFile(io.BytesIO(contenido)) as libro:
            for nombre in libro.namelist():
                if nombre.startswith("xl/worksheets/sheet"):
                    assert not re.search(rb"<f[ >/]", libro.read(nombre)), nombre

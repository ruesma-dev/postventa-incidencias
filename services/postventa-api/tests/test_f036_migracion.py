# services/postventa-api/tests/test_f036_migracion.py
"""La migración del Excel actual a la plantilla (F-036, T23; `design.md` §10).

`scripts/migracion_f036.py` (la lógica) y `scripts/migrar_excel_f036.py` (la
línea de órdenes), con un original, un catálogo y un fichero de grupos
**falsos**, construidos en memoria: ningún `.xlsx` toca el disco (R59) y la
línea de órdenes se prueba con un disco **en memoria**. Los nombres de
proveedor son inventados («Proveedor Inventado …») y nunca de Sigrid.

R53 (el original no se toca), R54 (cotejo), R55 (el mismo generador), R56 (ida
y vuelta por el importador), R57 en lo que hace cumplir el script (§10.2),
R58 (el informe, sin nombres de proveedor), R91 (proveedor completado) y R97
(nombre exacto de Sigrid → código → grupo; sin `--grupos`, solo
`--solo-informe`).

Undécima enmienda (`design.md` §10.5): R97 enmendado (solo para la errata) y
R128–R131 (un oficio que ha salido de la obra en Sigrid queda vacío y se
cuenta; el proveedor, solo del catálogo con que se ejecuta), con el catálogo
de prueba `CATALOGO_SIGRID`.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import subprocess
import sys
import unicodedata
from collections.abc import Callable
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
import yaml
from application.pipelines.plantilla import opciones_de_la_obra
from domain.models.equivalencias import MISMO, Catalogo
from domain.models.errores import FicheroDemasiadoGrande, FicheroNoEsPlantilla
from domain.models.importacion import (
    DESCRIPCION,
    OFICIO,
    PROVEEDOR,
    UBICACION,
    LibroLeido,
    reconocer_plantilla,
    validar_filas,
)
from infrastructure.documentos.excel_openpyxl import (
    PRIMERA_FILA,
    GeneradorPlantillaOpenpyxl,
    LectorPlantillaOpenpyxl,
)
from infrastructure.documentos.plantilla_yaml import cargar_plantilla_yaml
from openpyxl import Workbook

from scripts import migracion_f036 as mig
from scripts import migrar_excel_f036 as cli
from tests.utiles_plantilla import abrir

SERVICIO = Path(__file__).resolve().parents[1]
CONFIG = cargar_plantilla_yaml()
AHORA = datetime(2026, 9, 30, 9, 0, tzinfo=UTC)

NOMBRES_DE_PROVEEDOR = (
    "Proveedor Inventado Uno",
    "Proveedor Inventado Dos",
    "Proveedor Inventado Tres",
    "Proveedor Inventado Cuatro",
    "Proveedor Inventado Seis",
)

#: El JSON de T2, inventado: la forma de `design.md` §10.3.
CATALOGO: dict = {
    "obra": {"codigo": "0677", "nombre": "Obra inventada"},
    "unidades": [
        {"codigo": "0677.03VILLA 1.", "nombre": "Villa 1 inventada"},
        {"codigo": "0677.03VILLA 2.", "nombre": "Villa 2 inventada"},
    ],
    "oficios_obra": [
        # Un solo proveedor: R91 lo completa.
        {
            "oficio_codigo": "0010",
            "oficio_nombre": "Pintura",
            "proveedor_codigo": "P001",
            "proveedor_nombre": "Proveedor Inventado Uno",
            "marca_cif": None,
        },
        # Dos proveedores: R91 no elige.
        {
            "oficio_codigo": "0020",
            "oficio_nombre": "Albañilería",
            "proveedor_codigo": "P002",
            "proveedor_nombre": "Proveedor Inventado Dos",
            "marca_cif": None,
        },
        {
            "oficio_codigo": "0020",
            "oficio_nombre": "Albañilería",
            "proveedor_codigo": "P003",
            "proveedor_nombre": "Proveedor Inventado Tres",
            "marca_cif": None,
        },
        # Un grupo: 0033 con proveedor y 0133 sin él (como en la 0677).
        {
            "oficio_codigo": "0033",
            "oficio_nombre": "Solados y Alicatados",
            "proveedor_codigo": "P004",
            "proveedor_nombre": "Proveedor Inventado Cuatro",
            "marca_cif": None,
        },
        {
            "oficio_codigo": "0133",
            "oficio_nombre": "Solados y Alicatados M.O.",
            "proveedor_codigo": None,
            "proveedor_nombre": None,
            "marca_cif": None,
        },
        # Sin proveedor y sin grupo (como 0144 en la 0677).
        {
            "oficio_codigo": "0144",
            "oficio_nombre": "Mamparas",
            "proveedor_codigo": None,
            "proveedor_nombre": None,
            "marca_cif": None,
        },
        # Un grupo cuyo único par junta dos códigos de oficio.
        {
            "oficio_codigo": "0060",
            "oficio_nombre": "Fontanería",
            "proveedor_codigo": "P006",
            "proveedor_nombre": "Proveedor Inventado Seis",
            "marca_cif": None,
        },
        {
            "oficio_codigo": "0160",
            "oficio_nombre": "Fontaneria",
            "proveedor_codigo": "P006",
            "proveedor_nombre": "Proveedor Inventado Seis",
            "marca_cif": None,
        },
    ],
    "oficios_catalogo": [],
}

#: El fichero de grupos vigentes de R98, inventado.
GRUPOS: dict = {"obra": "0677", "oficio": [["0033", "0133"], ["0060", "0160"]]}

#: Dos oficios de `auxofc` que **no** son de la obra (R128), inventados. El
#: segundo arrastra blancos al final en Sigrid.
FUERA_DE_LA_OBRA = "Oficio que salió"
RETIRADO = "Oficio retirado"

#: `CATALOGO` con `oficios_catalogo` lleno (`design.md` §10.5, «Cómo se
#: prueba»): los oficios de la obra y los dos de arriba. `CATALOGO` sigue con
#: la lista vacía, y por eso los tests de siempre no cambian.
CATALOGO_SIGRID: dict = {
    **copy.deepcopy(CATALOGO),
    "oficios_catalogo": [
        *(
            {"codigo": codigo, "nombre": nombre}
            for codigo, nombre in sorted(
                {(f["oficio_codigo"], f["oficio_nombre"]) for f in CATALOGO["oficios_obra"]}
            )
        ),
        {"codigo": "0901", "nombre": FUERA_DE_LA_OBRA},
        {"codigo": "0902", "nombre": RETIRADO + "   "},
    ],
}


def _errata(origen: int, n: int, nombre: str) -> str:
    """El mensaje literal de la errata (R97, `design.md` §10.5)."""
    return (
        f"fila {origen} del original, fila nueva {n}: el oficio «{nombre}» no es "
        "exactamente el nombre de ningún oficio de Sigrid, ni de la obra ni del "
        "resto del catálogo: corrígelo en la tabla de correcciones (R97)"
    )

#: El original: sin cabecera, la unidad solo en la primera fila (doc. 05).
ORIGINAL: tuple[tuple[str | None, str | None, str | None, str | None], ...] = (
    ("Villa 1", "SALA/ ESTUDIO", "Pintar paredes del baño", "Pintura"),
    (None, "SALA/ ESTUDIO", "Rematar agujero en pared", "Albañilería"),
    (
        None,
        "TERRAZA",
        "Enlechar azulejos, PELIGRO DE SEGURIDAD.",
        "Solados y alicatados",
    ),
    (None, "TERRAZA", "Sellar encuentro. Puerta suelta", None),
    (None, "GENERALES", "Repetida de otra fila", None),
    (None, "COCINA", "Grifo  gotea\n", None),
    (None, "Cocina", "Mampara floja", "Mamparas"),
)


def _nueva(
    descripcion: str,
    *,
    ubicacion: str | None = None,
    detalle: str | None = None,
    oficio: str | None = None,
    urgencia: str | None = None,
    listado: str | None = None,
) -> dict:
    return {
        "ubicacion": ubicacion,
        "descripcion": descripcion,
        "detalle": detalle,
        "oficio": oficio,
        "urgencia": urgencia,
        "listado": listado,
    }


#: La tabla de correcciones (el YAML de §10.2 ya cargado), inventada.
CORRECCIONES: dict = {
    "obra": "0677",
    "unidad": "Villa 1 inventada",
    "filas": [
        {
            "origen": 1,
            "ubicacion_original": "SALA/ ESTUDIO",
            "texto_original": "Pintar paredes del baño",
            "resultado": [
                _nueva(
                    "Pintar paredes del baño",
                    ubicacion="Sala/estudio",
                    oficio="Pintura",
                )
            ],
            "cambios": ["ubicación normalizada"],
            "propuesto": [],
        },
        {
            "origen": 2,
            "ubicacion_original": "SALA/ ESTUDIO",
            "texto_original": "Rematar agujero en pared",
            "resultado": [
                _nueva(
                    "Rematar agujero en pared",
                    ubicacion="Sala/estudio",
                    oficio="Albañilería",
                )
            ],
            "cambios": ["ubicación normalizada"],
        },
        {
            "origen": 3,
            "ubicacion_original": "TERRAZA",
            "texto_original": "Enlechar azulejos, PELIGRO DE SEGURIDAD.",
            "resultado": [
                _nueva(
                    "Enlechar azulejos",
                    ubicacion="Terraza 1",
                    oficio="Solados y Alicatados M.O.",
                    urgencia="seguridad",
                )
            ],
            "cambios": ["ubicación normalizada", "peligro de seguridad a Urgencia"],
            "propuesto": ["oficio", "ubicacion"],
        },
        {
            "origen": 4,
            "ubicacion_original": "TERRAZA",
            "texto_original": "Sellar encuentro. Puerta suelta",
            "resultado": [
                _nueva(
                    "Sellar encuentro",
                    ubicacion="Terraza 1",
                    oficio="Solados y Alicatados",
                ),
                _nueva("Puerta suelta", ubicacion="Terraza 1"),
            ],
            "cambios": ["dos defectos, separados"],
            "propuesto": ["oficio"],
        },
        {
            "origen": 5,
            "ubicacion_original": "GENERALES",
            "texto_original": "Repetida de otra fila",
            "resultado": [],
            "cambios": ["repite la fila 2"],
        },
        {
            "origen": 6,
            "ubicacion_original": "COCINA",
            "texto_original": "Grifo gotea",
            "resultado": [
                _nueva("Grifo gotea", ubicacion="Cocina", oficio="Fontanería")
            ],
            "cambios": ["ubicación normalizada"],
            "propuesto": ["oficio"],
        },
        {
            "origen": 7,
            "ubicacion_original": "Cocina",
            "texto_original": "Mampara floja",
            "resultado": [
                _nueva("Mampara floja", ubicacion="Cocina", oficio="Mamparas")
            ],
        },
    ],
}


def _original(
    filas: tuple[tuple[str | None, ...], ...] = ORIGINAL, *, otra_hoja: bool = False
) -> bytes:
    """El Excel actual, en memoria: A = unidad, D = ubicación, E = texto, F = oficio."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    for numero, (unidad, ubicacion, texto, oficio) in enumerate(filas, start=1):
        hoja.cell(row=numero, column=1, value=unidad)
        hoja.cell(row=numero, column=4, value=ubicacion)
        hoja.cell(row=numero, column=5, value=texto)
        hoja.cell(row=numero, column=6, value=oficio)
    if otra_hoja:
        libro.create_sheet("Otra").cell(row=1, column=4, value="no se lee")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _correcciones(cambio: Callable[[dict], None] | None = None) -> dict:
    tabla = copy.deepcopy(CORRECCIONES)
    if cambio is not None:
        cambio(tabla)
    return tabla


#: «Lo de siempre» en `_migrar`: `None` también es un valor que se prueba.
_DE_SIEMPRE = object()


def _migrar(
    *,
    original: bytes | None = None,
    correcciones: object = _DE_SIEMPRE,
    catalogo: object = _DE_SIEMPRE,
    grupos: object = GRUPOS,
    **puertos,
) -> mig.Migracion:
    return mig.migrar(
        original=_original() if original is None else original,
        correcciones=_correcciones() if correcciones is _DE_SIEMPRE else correcciones,
        catalogo=copy.deepcopy(CATALOGO) if catalogo is _DE_SIEMPRE else catalogo,
        grupos=copy.deepcopy(grupos),
        listas=CONFIG.listas,
        textos=CONFIG.textos,
        ahora=AHORA,
        **puertos,
    )


def _detenida(**kwargs) -> tuple[str, ...]:
    with pytest.raises(mig.MigracionDetenida) as error:
        _migrar(**kwargs)
    return error.value.problemas


def _celdas_v2(contenido: bytes) -> list[dict[str, str | None]]:
    """Las filas de «Incidencias» del v2, por nombre de columna."""
    hoja = abrir(contenido)["Incidencias"]
    cabecera = [c.value for c in hoja[1]]
    filas = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in fila):
            filas.append(dict(zip(cabecera, fila, strict=True)))
    return filas


# --------------------------------------------------------------------------
# El camino bueno: lo que sale
# --------------------------------------------------------------------------


def test_f036_r55_r56_camino_bueno_con_grupos() -> None:
    m = _migrar()
    assert m.obra == "0677"
    assert m.unidad == "Villa 1 inventada"
    assert m.con_grupos is True
    assert m.grupos_de_varios == 2
    assert [c.origen for c in m.compuestas] == [1, 2, 3, 4, 4, 6, 7]
    celdas = _celdas_v2(m.v2)
    assert [c["Descripción corta"] for c in celdas] == [
        "Pintar paredes del baño",
        "Rematar agujero en pared",
        "Enlechar azulejos",
        "Sellar encuentro",
        "Puerta suelta",
        "Grifo gotea",
        "Mampara floja",
    ]
    assert {c["Unidad"] for c in celdas} == {"Villa 1 inventada"}
    assert [c["Ubicación"] for c in celdas] == [
        "Sala/estudio",
        "Sala/estudio",
        "Terraza 1",
        "Terraza 1",
        "Terraza 1",
        "Cocina",
        "Cocina",
    ]
    assert [c["Urgencia"] for c in celdas] == [
        None,
        None,
        "Peligro para la seguridad",
        None,
        None,
        None,
        None,
    ]
    assert all(c["Listado"] is None and c["Detalle"] is None for c in celdas)
    assert all(c["Errores (lo rellena el sistema)"] is None for c in celdas)


def test_f036_r56_el_v2_pasa_el_importador_sin_errores() -> None:
    m = _migrar()
    libro = LectorPlantillaOpenpyxl().leer(contenido=m.v2)
    assert reconocer_plantilla(libro) == "0677"
    assert len(libro.filas) == 7


def test_f036_r55_el_v2_sale_del_mismo_generador_con_las_opciones_de_la_obra() -> None:
    llamadas: list[dict] = []

    class Espia:
        def generar(self, **kwargs) -> bytes:
            llamadas.append(kwargs)
            return GeneradorPlantillaOpenpyxl().generar(**kwargs)

    m = _migrar(generador=Espia())
    (llamada,) = llamadas
    assert llamada["catalogo"].obra_codigo == "0677"
    assert [u.codigo for u in llamada["catalogo"].unidades] == [
        "0677.03VILLA 1.",
        "0677.03VILLA 2.",
    ]
    assert [o.etiqueta for o in llamada["opciones_oficio"]] == [
        "Pintura",
        "Albañilería",
        "Solados y Alicatados",
        "Fontanería",
        "Mamparas",
    ]
    assert llamada["listas"] is CONFIG.listas
    assert llamada["textos"] is CONFIG.textos
    assert llamada["generada_at"] == AHORA
    assert llamada["importacion_origen"] is None
    assert len(llamada["filas"]) == 7
    # El libro trae los desplegables de la plantilla del portal (R55).
    catalogos = abrir(m.v2)["_catalogos"]
    oficios = [c.value for c in catalogos["C"][1:] if c.value is not None]
    assert oficios == [o.etiqueta for o in llamada["opciones_oficio"]]


def test_f036_r55_el_catalogo_sale_del_json_como_de_sigrid() -> None:
    m = _migrar()
    assert m.catalogo.obra_nombre == "Obra inventada"
    assert [o.codigo for o in m.catalogo.oficios] == [
        "0010",
        "0020",
        "0033",
        "0060",
        "0133",
        "0144",
        "0160",
    ]
    assert len(m.catalogo.proveedores) == 8
    assert m.catalogo.proveedores[4].proveedor_codigo is None


def test_f036_r55_el_primer_nombre_de_un_codigo_manda() -> None:
    def repetir(catalogo: dict) -> None:
        catalogo["oficios_obra"].append(
            {
                "oficio_codigo": "0010",
                "oficio_nombre": "Pintura bis",
                "proveedor_codigo": None,
                "proveedor_nombre": None,
            }
        )

    catalogo = copy.deepcopy(CATALOGO)
    repetir(catalogo)
    m = _migrar(catalogo=catalogo)
    assert m.catalogo.oficios[0].nombre == "Pintura"


# --------------------------------------------------------------------------
# R97 · el nombre exacto de Sigrid, su código y su grupo
# --------------------------------------------------------------------------


def test_f036_r97_con_grupos_se_escribe_la_etiqueta_del_grupo() -> None:
    m = _migrar()
    oficios = [c["Oficio"] for c in _celdas_v2(m.v2)]
    assert oficios == [
        "Pintura",
        "Albañilería",
        "Solados y Alicatados",  # 0133 va en el grupo de 0033 (R86)
        "Solados y Alicatados",
        None,
        "Fontanería",
        "Mamparas",
    ]
    tercera = m.compuestas[2]
    assert tercera.oficio_codigo == "0133"
    assert tercera.oficio_etiqueta == "Solados y Alicatados"
    assert tercera.codigos_en_obra == 2
    assert [c.codigos_en_obra for c in m.compuestas] == [1, 1, 2, 2, 0, 2, 1]


def test_f036_r97_sin_grupos_cada_codigo_es_su_grupo() -> None:
    m = _migrar(grupos=None)
    assert m.con_grupos is False
    assert m.grupos_de_varios == 0
    oficios = [c["Oficio"] for c in _celdas_v2(m.v2)]
    assert oficios == [
        "Pintura",
        "Albañilería",
        "Solados y Alicatados M.O.",
        "Solados y Alicatados",
        None,
        "Fontanería (0060)",  # choca con «Fontaneria» al plegar (R8)
        "Mamparas",
    ]
    assert [c.codigos_en_obra for c in m.compuestas] == [1, 1, 1, 1, 0, 1, 1]


@pytest.mark.parametrize(
    "nombre",
    [
        "Solados y alicatados",  # la muestra, que no es exactamente el de Sigrid
        "Pintura ",
        " Pintura",
        "pintura",
        "Pintura (0010)",  # una etiqueta del desplegable no es un nombre
        "Oficio que no existe",
        # R12-1: la comparación es tal cual frente a la obra y a todo Sigrid.
        "Solados y  Alicatados",  # blancos interiores (dos espacios)
        "Píntura",  # tildes por el lado tolerante, con ñ y sin ella
        "Albanileria",  # tildes por la suma de la obra (la ñ cuenta como tilde)
        "Pintura.",  # puntuación quitada por el lado tolerante
        "Solados-y-Alicatados",  # puntuación cambiada por espacio: ambos y tolerante
        "Solados y Alicatados M O ",  # puntuación cambiada por espacio: ambos y suma
        "SoladosyAlicatados",  # blancos todos fuera: la suma de los dos catálogos
        unicodedata.normalize("NFD", "Albañilería"),  # forma Unicode de la obra
    ],
)
@pytest.mark.parametrize("de_sigrid", [False, True], ids=["CATALOGO", "CATALOGO_SIGRID"])
def test_f036_r97_nombre_que_no_es_exactamente_uno_de_la_obra_para(
    nombre: str, de_sigrid: bool
) -> None:
    """La errata sigue parando, con `oficios_catalogo` vacío o lleno (R97, §10.5)."""

    def poner(tabla: dict) -> None:
        tabla["filas"][0]["resultado"][0]["oficio"] = nombre

    catalogo = copy.deepcopy(CATALOGO_SIGRID if de_sigrid else CATALOGO)
    problemas = _detenida(correcciones=_correcciones(poner), catalogo=catalogo)
    assert problemas == (_errata(1, 1, nombre),)


def test_f036_r97_el_nombre_se_compara_con_el_de_sigrid_recortado() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["oficios_obra"][0]["oficio_nombre"] = "Pintura  "
    m = _migrar(catalogo=catalogo)
    assert m.compuestas[0].oficio_codigo == "0010"


def test_f036_r97_nombre_de_dos_oficios_de_la_obra_para() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["oficios_obra"][-1]["oficio_nombre"] = "Fontanería"
    problemas = _detenida(catalogo=catalogo)
    assert problemas == (
        (
            "fila 6 del original, fila nueva 1: el oficio «Fontanería» es el nombre de "
            "varios oficios de la obra en Sigrid (0060, 0160): no se puede elegir"
        ),
    )


def test_f036_r97_un_oficio_sin_nombre_en_sigrid_no_casa_con_nada() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["oficios_obra"][0]["oficio_nombre"] = None
    problemas = _detenida(catalogo=catalogo)
    assert problemas[0].startswith(
        "fila 1 del original, fila nueva 1: el oficio «Pintura»"
    )


def test_f036_r97_todos_los_nombres_malos_se_dicen_a_la_vez() -> None:
    def poner(tabla: dict) -> None:
        tabla["filas"][0]["resultado"][0]["oficio"] = "Uno"
        tabla["filas"][3]["resultado"][1]["oficio"] = "Dos"

    problemas = _detenida(correcciones=_correcciones(poner))
    assert [p.split(":")[0] for p in problemas] == [
        "fila 1 del original, fila nueva 1",
        "fila 4 del original, fila nueva 2",
    ]


def test_f036_r97_fichero_de_grupos_de_otra_obra() -> None:
    problemas = _detenida(grupos={"obra": "0678", "oficio": []})
    assert problemas == (
        "el fichero de grupos es de la obra «0678» y la tabla, de la «0677»",
    )


def test_f036_r97_el_codigo_de_obra_de_los_grupos_se_normaliza() -> None:
    m = _migrar(grupos={"obra": " 0677 ", "oficio": [["0033", "0133"]]})
    assert m.grupos_de_varios == 1


@pytest.mark.parametrize(
    "grupos",
    [
        [],
        {"obra": "0677"},
        {"oficio": []},
        {"obra": "0677", "oficio": [], "proveedor": []},
        {"obra": 677, "oficio": []},
        {"obra": "06 77/", "oficio": []},
        {"obra": "0677", "oficio": {}},
        {"obra": "0677", "oficio": [["0033"]]},
        {"obra": "0677", "oficio": ["0033"]},
        {"obra": "0677", "oficio": [["0033", 133]]},
        {"obra": "0677", "oficio": [["0033", ""]]},
    ],
)
def test_f036_r97_fichero_de_grupos_mal_formado(grupos: object) -> None:
    problemas = _detenida(grupos=grupos)
    assert len(problemas) == 1
    assert problemas[0].startswith("el fichero de grupos no tiene la forma de R98")


@pytest.mark.parametrize(
    ("oficio", "repetido"),
    [
        ([["0033", "0133", "0033"]], "0033"),
        ([["0033", "0133"], ["0144", "0133"]], "0133"),
    ],
)
def test_f036_r97_codigo_repetido_en_los_grupos(oficio: list, repetido: str) -> None:
    problemas = _detenida(grupos={"obra": "0677", "oficio": oficio})
    assert problemas == (f"el fichero de grupos repite el código «{repetido}»",)


def test_f036_r97_codigos_de_los_grupos_que_no_son_de_la_obra_no_estorban() -> None:
    m = _migrar(grupos={"obra": "0677", "oficio": [["0033", "0133", "9999"]]})
    assert m.compuestas[2].oficio_etiqueta == "Solados y Alicatados"


def test_f036_r97_las_decisiones_del_fichero_solo_con_los_dos_codigos() -> None:
    decisiones = mig.decisiones_de_grupos(
        {"obra": "0677", "oficio": [["0133", "0033", "9999"]]}, "0677"
    )
    assert [(d.codigo_a, d.codigo_b) for d in decisiones] == [
        ("0033", "0133"),
        ("0033", "9999"),
        ("0133", "9999"),
    ]
    assert all(
        d.decision == MISMO and d.catalogo is Catalogo.OFICIO for d in decisiones
    )
    assert {d.obra_codigo for d in decisiones} == {"0677"}
    equivalencias = mig.EquivalenciasDeFichero(decisiones)
    vistas = equivalencias.ultimas_decisiones(
        catalogo=Catalogo.OFICIO, codigos={"0033", "0133"}
    )
    assert [(d.codigo_a, d.codigo_b) for d in vistas] == [("0033", "0133")]
    assert (
        equivalencias.ultimas_decisiones(
            catalogo=Catalogo.PROVEEDOR, codigos={"0033", "0133"}
        )
        == ()
    )
    with pytest.raises(RuntimeError, match="no registra"):
        equivalencias.registrar(decisiones=decisiones)


# --------------------------------------------------------------------------
# R91 · el proveedor, solo si es el único y resuelve al mismo oficio
# --------------------------------------------------------------------------


def test_f036_r91_con_grupos() -> None:
    m = _migrar()
    proveedores = [c["Proveedor"] for c in _celdas_v2(m.v2)]
    assert proveedores == [
        "Pintura · Proveedor Inventado Uno",  # un solo par → se completa
        None,  # dos proveedores: no se elige
        None,  # 0133: el único par del grupo es de 0033 → quedaría 0033
        "Solados y Alicatados · Proveedor Inventado Cuatro",  # 0033, su par
        None,  # sin oficio
        None,  # el único par junta 0060 y 0160: no resuelve a un código
        None,  # 0144: sin proveedor en la obra
    ]
    assert [c.proveedor_codigo for c in m.compuestas] == [
        "P001",
        None,
        None,
        "P004",
        None,
        None,
        None,
    ]


def test_f036_r91_sin_grupos() -> None:
    m = _migrar(grupos=None)
    proveedores = [c["Proveedor"] for c in _celdas_v2(m.v2)]
    assert proveedores == [
        "Pintura · Proveedor Inventado Uno",
        None,
        None,  # 0133 no tiene proveedor en la obra
        "Solados y Alicatados · Proveedor Inventado Cuatro",
        None,
        "Fontanería (0060) · Proveedor Inventado Seis",  # su grupo solo: resuelve
        None,
    ]


def test_f036_r91_el_proveedor_completado_resuelve_al_mismo_oficio_al_importar() -> (
    None
):
    m = _migrar()
    libro = LectorPlantillaOpenpyxl().leer(contenido=m.v2)
    validas, errores = validar_filas(
        libro.filas,
        catalogo=m.catalogo,
        opciones_oficio=m.opciones.oficios,
        opciones_proveedor=m.opciones.proveedores,
        listas=CONFIG.listas,
    )
    assert errores == ()
    assert [
        (v.oficio and v.oficio.codigo, v.proveedor and v.proveedor.codigo)
        for v in validas
    ] == [
        ("0010", "P001"),
        ("0020", None),
        (None, None),  # grupo de dos códigos en la obra, sin proveedor: ambiguo
        ("0033", "P004"),
        (None, None),
        (None, None),
        ("0144", None),
    ]


# --------------------------------------------------------------------------
# R54 · cotejo con el original
# --------------------------------------------------------------------------


def test_f036_r54_blancos_distintos_casan_y_nada_mas() -> None:
    # «Grifo  gotea\n» del original casa con «Grifo gotea» de la tabla.
    assert _migrar().originales[5].texto == "Grifo  gotea\n"


def test_f036_r54_falta_una_fila_del_original() -> None:
    problemas = _detenida(correcciones=_correcciones(lambda t: t["filas"].pop(4)))
    assert problemas == ("falta la fila 5 del original en la tabla de correcciones",)


def test_f036_r54_fila_repetida() -> None:
    def repetir(tabla: dict) -> None:
        tabla["filas"].insert(1, copy.deepcopy(tabla["filas"][0]))

    problemas = _detenida(correcciones=_correcciones(repetir))
    assert problemas == ("la fila 1 del original aparece 2 veces en la tabla",)


def test_f036_r54_fila_que_no_esta_en_el_original() -> None:
    def sobrar(tabla: dict) -> None:
        fila = copy.deepcopy(tabla["filas"][-1])
        fila["origen"] = 9
        tabla["filas"].append(fila)

    problemas = _detenida(correcciones=_correcciones(sobrar))
    assert problemas == ("la fila 9 de la tabla no es una fila con datos del original",)


def test_f036_r54_la_tabla_va_en_el_orden_del_original() -> None:
    def desordenar(tabla: dict) -> None:
        tabla["filas"][0], tabla["filas"][1] = tabla["filas"][1], tabla["filas"][0]

    problemas = _detenida(correcciones=_correcciones(desordenar))
    assert problemas == ("las filas de la tabla no van en el orden del original",)


def test_f036_r54_ubicacion_o_texto_que_no_casan() -> None:
    def cambiar(tabla: dict) -> None:
        tabla["filas"][0]["ubicacion_original"] = "SALA/ESTUDIO"
        tabla["filas"][1]["texto_original"] = "Rematar agujero en la pared"

    problemas = _detenida(correcciones=_correcciones(cambiar))
    assert problemas == (
        (
            "la fila 1 del original no casa: su ubicación es «SALA/ ESTUDIO» y la "
            "tabla dice «SALA/ESTUDIO»"
        ),
        (
            "la fila 2 del original no casa: su texto es «Rematar agujero en pared» y "
            "la tabla dice «Rematar agujero en la pared»"
        ),
    )


def test_f036_r54_mayusculas_distintas_no_casan() -> None:
    def cambiar(tabla: dict) -> None:
        tabla["filas"][6]["ubicacion_original"] = "COCINA"

    problemas = _detenida(correcciones=_correcciones(cambiar))
    assert problemas == (
        (
            "la fila 7 del original no casa: su ubicación es «Cocina» y la tabla dice "
            "«COCINA»"
        ),
    )


def test_f036_r54_ubicacion_vacia_en_el_original_casa_con_null() -> None:
    original = list(ORIGINAL)
    original[6] = (None, None, "Mampara floja", "Mamparas")

    def cambiar(tabla: dict) -> None:
        tabla["filas"][6]["ubicacion_original"] = None
        tabla["filas"][6]["cambios"] = ["ubicación puesta"]

    m = _migrar(
        original=_original(tuple(original)), correcciones=_correcciones(cambiar)
    )
    assert m.originales[6].ubicacion is None


def test_f036_r54_todos_los_problemas_a_la_vez() -> None:
    def romper(tabla: dict) -> None:
        tabla["filas"].pop(6)
        tabla["filas"][0]["texto_original"] = "otro"

    problemas = _detenida(correcciones=_correcciones(romper))
    assert problemas == (
        "falta la fila 7 del original en la tabla de correcciones",
        (
            "la fila 1 del original no casa: su texto es «Pintar paredes del baño» y la "
            "tabla dice «otro»"
        ),
    )


def test_f036_r54_filas_vacias_del_original_no_cuentan() -> None:
    original = (*ORIGINAL, (None, None, None, None), (None, "   ", None, None))
    m = _migrar(original=_original(original))
    assert [o.numero for o in m.originales] == [1, 2, 3, 4, 5, 6, 7]


def test_f036_r54_fila_en_blanco_en_medio_se_salta() -> None:
    original = (*ORIGINAL[:3], (None, None, None, None), *ORIGINAL[3:])

    def renumerar(tabla: dict) -> None:
        for fila in tabla["filas"][3:]:
            fila["origen"] += 1

    m = _migrar(original=_original(original), correcciones=_correcciones(renumerar))
    assert [o.numero for o in m.originales] == [1, 2, 3, 5, 6, 7, 8]


def test_f036_r54_solo_se_lee_la_primera_hoja_y_valores_que_no_son_texto() -> None:
    original = list(ORIGINAL)
    original[6] = (None, "Cocina", 12.5, "Mamparas")

    def cambiar(tabla: dict) -> None:
        tabla["filas"][6]["texto_original"] = "12.5"
        tabla["filas"][6]["cambios"] = ["texto"]

    m = _migrar(
        original=_original(tuple(original), otra_hoja=True),
        correcciones=_correcciones(cambiar),
    )
    assert m.originales[6].texto == "12.5"


def test_f036_r54_mas_de_una_unidad_en_el_original_para() -> None:
    original = list(ORIGINAL)
    original[3] = ("Villa 2", *ORIGINAL[3][1:])
    problemas = _detenida(original=_original(tuple(original)))
    assert problemas == (
        (
            "el original trae 2 unidades distintas en la columna A y la tabla de "
            "correcciones solo admite una"
        ),
    )


def test_f036_r54_la_misma_unidad_escrita_con_otros_blancos_es_una() -> None:
    original = list(ORIGINAL)
    original[3] = (" Villa  1 ", *ORIGINAL[3][1:])
    assert len(_migrar(original=_original(tuple(original))).originales) == 7


def test_f036_r54_el_original_se_abre_solo_para_leer_y_por_sus_valores(
    monkeypatch,
) -> None:
    llamadas: list[dict] = []
    abrir_de_verdad = mig.load_workbook

    def espia(fichero, **kwargs):
        llamadas.append(kwargs)
        return abrir_de_verdad(fichero, **kwargs)

    monkeypatch.setattr(mig, "load_workbook", espia)
    mig.leer_original(_original())
    assert llamadas == [{"read_only": True, "data_only": True}]


def test_f036_r54_mas_alla_de_la_columna_f_no_se_lee() -> None:
    libro = Workbook()
    hoja = libro.active
    hoja.cell(row=1, column=4, value="SALA/ ESTUDIO")
    hoja.cell(row=1, column=5, value="Texto")
    hoja.cell(row=1, column=7, value="columna G")
    hoja.cell(row=2, column=7, value="solo en la G")
    salida = io.BytesIO()
    libro.save(salida)
    filas = mig.leer_original(salida.getvalue())
    assert filas == (mig.FilaOriginal(1, None, "SALA/ ESTUDIO", "Texto", None),)


def test_f036_r54_original_que_no_es_un_excel() -> None:
    problemas = _detenida(original=b"esto no es un xlsx")
    assert problemas == ("el original no se puede abrir como un Excel (.xlsx)",)


def test_f036_r54_cambios_obligatorios_si_la_fila_cambia() -> None:
    def quitar(tabla: dict) -> None:
        del tabla["filas"][0]["cambios"]
        tabla["filas"][3]["cambios"] = []

    problemas = _detenida(correcciones=_correcciones(quitar))
    assert problemas == (
        "la fila 1 del original cambia y no dice por qué (`cambios`)",
        "la fila 4 del original cambia y no dice por qué (`cambios`)",
    )


def test_f036_r54_r57_descartar_exige_motivo() -> None:
    problemas = _detenida(
        correcciones=_correcciones(lambda t: t["filas"][4].update(cambios=[]))
    )
    assert problemas == ("la fila 5 del original cambia y no dice por qué (`cambios`)",)


@pytest.mark.parametrize(
    "cambio",
    [
        {"ubicacion": "Otra"},
        {"descripcion": "Mampara muy floja"},
        {"detalle": "algo"},
        {"oficio": None},
        {"urgencia": "urgente"},
        {"listado": "primero"},
    ],
)
def test_f036_r54_cualquier_campo_distinto_es_un_cambio(cambio: dict) -> None:
    def tocar(tabla: dict) -> None:
        tabla["filas"][6]["resultado"][0].update(cambio)

    problemas = _detenida(correcciones=_correcciones(tocar))
    assert "la fila 7 del original cambia y no dice por qué (`cambios`)" in problemas


def test_f036_r54_dos_filas_iguales_al_original_tambien_son_un_cambio() -> None:
    def duplicar(tabla: dict) -> None:
        resultado = tabla["filas"][6]["resultado"]
        resultado.append(copy.deepcopy(resultado[0]))

    problemas = _detenida(correcciones=_correcciones(duplicar))
    assert problemas == ("la fila 7 del original cambia y no dice por qué (`cambios`)",)


# --------------------------------------------------------------------------
# R57 · lo que el script hace cumplir de la tabla (§10.2)
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("campo", "texto", "problema"),
    [
        (
            "descripcion",
            "Mampara floja, PELIGRO DE SEGURIDAD",
            "«peligro de seguridad»",
        ),
        ("detalle", "peligro  de   seguridád", "«peligro de seguridad»"),
        ("descripcion", "Repasar como en los otros baños", "«como en los otros»"),
        ("detalle", "Como en las otras", "«como en las otras»"),
        ("detalle", "como en todos los baños", "«como en todos»"),
        ("detalle", "como en todas", "«como en todas»"),
        ("descripcion", "Igual que la anterior", "«igual que»"),
    ],
)
def test_f036_r57_textos_prohibidos(campo: str, texto: str, problema: str) -> None:
    def poner(tabla: dict) -> None:
        tabla["filas"][6]["resultado"][0][campo] = texto
        tabla["filas"][6]["cambios"] = ["texto"]

    problemas = _detenida(correcciones=_correcciones(poner))
    nombre = {"descripcion": "la descripción", "detalle": "el detalle"}[campo]
    assert problemas == (
        f"fila 7 del original, fila nueva 1: {nombre} lleva {problema}",
    )


def test_f036_r57_los_problemas_de_contenido_se_suman_a_los_del_cotejo() -> None:
    def romper(tabla: dict) -> None:
        tabla["filas"].pop(4)
        tabla["filas"][0]["resultado"][0]["detalle"] = "igual que arriba"

    problemas = _detenida(correcciones=_correcciones(romper))
    assert problemas == (
        "falta la fila 5 del original en la tabla de correcciones",
        "fila 1 del original, fila nueva 1: el detalle lleva «igual que»",
    )


# --------------------------------------------------------------------------
# La forma de la tabla de correcciones (§10.2)
# --------------------------------------------------------------------------


def _fila0(tabla: dict) -> dict:
    return tabla["filas"][0]


@pytest.mark.parametrize(
    ("romper", "problema"),
    [
        (lambda t: t.clear() or t.update(x=1), "la tabla de correcciones"),
        (lambda t: t.pop("obra"), "la tabla de correcciones: falta «obra»"),
        (lambda t: t.update(extra=1), "la tabla de correcciones: sobra «extra»"),
        (lambda t: t.update(obra=677), "la tabla de correcciones, obra"),
        (lambda t: t.update(unidad=" "), "la tabla de correcciones, unidad"),
        (lambda t: t.update(filas={}), "la tabla de correcciones, filas"),
        (lambda t: t.update(filas=[]), "la tabla de correcciones, filas"),
        (lambda t: t["filas"].append(3), "la tabla de correcciones, entrada 8"),
        (lambda t: _fila0(t).pop("resultado"), "entrada 1: falta «resultado»"),
        (lambda t: _fila0(t).update(nota="x"), "entrada 1: sobra «nota»"),
        (lambda t: _fila0(t).update(origen="1"), "entrada 1, origen"),
        (lambda t: _fila0(t).update(origen=True), "entrada 1, origen"),
        (lambda t: _fila0(t).update(origen=0), "entrada 1, origen"),
        (
            lambda t: _fila0(t).update(ubicacion_original=3),
            "entrada 1, ubicacion_original",
        ),
        (lambda t: _fila0(t).update(texto_original=""), "entrada 1, texto_original"),
        (lambda t: _fila0(t).update(resultado={}), "entrada 1, resultado"),
        (lambda t: _fila0(t).update(resultado=["x"]), "entrada 1, resultado 1"),
        (lambda t: _fila0(t).update(cambios="uno"), "entrada 1, cambios"),
        (lambda t: _fila0(t).update(cambios=[""]), "entrada 1, cambios"),
        (lambda t: _fila0(t).update(cambios=[1]), "entrada 1, cambios"),
        (lambda t: _fila0(t).update(propuesto=["proveedor"]), "entrada 1, propuesto"),
        (lambda t: _fila0(t).update(propuesto="oficio"), "entrada 1, propuesto"),
        (
            lambda t: _fila0(t)["resultado"][0].pop("descripcion"),
            "resultado 1: falta «descripcion»",
        ),
        (
            lambda t: _fila0(t)["resultado"][0].update(proveedor="x"),
            "resultado 1: sobra «proveedor»",
        ),
        (
            lambda t: _fila0(t)["resultado"][0].update(descripcion=" "),
            "resultado 1, descripcion",
        ),
        (
            lambda t: _fila0(t)["resultado"][0].update(detalle=""),
            "resultado 1, detalle",
        ),
        (lambda t: _fila0(t)["resultado"][0].update(oficio=10), "resultado 1, oficio"),
        (
            lambda t: _fila0(t)["resultado"][0].update(urgencia="Urgente"),
            "resultado 1, urgencia",
        ),
        (
            lambda t: _fila0(t)["resultado"][0].update(listado="Primer listado"),
            "resultado 1, listado",
        ),
    ],
)
def test_f036_forma_de_la_tabla_mal(
    romper: Callable[[dict], None], problema: str
) -> None:
    problemas = _detenida(correcciones=_correcciones(romper))
    assert len(problemas) == 1
    assert problema in problemas[0]


@pytest.mark.parametrize("crudo", [None, [], "texto"])
def test_f036_forma_de_la_tabla_que_no_es_un_mapping(crudo: object) -> None:
    problemas = _detenida(correcciones=crudo)
    assert problemas == (
        "la tabla de correcciones tiene que ser un mapping con sus campos",
    )


def test_f036_forma_de_la_tabla_opcionales() -> None:
    tabla = mig.leer_correcciones(_correcciones())
    fila = tabla.filas[1]
    assert fila.cambios == ("ubicación normalizada",)
    assert fila.propuesto == ()
    assert tabla.filas[6].cambios == ()
    minima = mig.leer_correcciones(
        {
            "obra": "0677",
            "unidad": "U",
            "filas": [
                {
                    "origen": 1,
                    "ubicacion_original": None,
                    "texto_original": None,
                    "resultado": [{"descripcion": "Algo"}],
                }
            ],
        }
    )
    (nueva,) = minima.filas[0].resultado
    assert nueva == mig.FilaNueva(
        ubicacion=None,
        descripcion="Algo",
        detalle=None,
        oficio=None,
        urgencia=None,
        listado=None,
    )


def test_f036_forma_de_la_tabla_urgencia_y_listado_por_codigo() -> None:
    tabla = mig.leer_correcciones(_correcciones())
    assert tabla.filas[2].resultado[0].urgencia == "seguridad"
    assert tabla.filas[2].propuesto == ("oficio", "ubicacion")


def test_f036_forma_de_la_tabla_propuesto_admite_todos_los_campos() -> None:
    def todos(tabla: dict) -> None:
        tabla["filas"][0]["propuesto"] = list(mig.CAMPOS_NUEVA)

    tabla = mig.leer_correcciones(_correcciones(todos))
    assert tabla.filas[0].propuesto == mig.CAMPOS_NUEVA


# --------------------------------------------------------------------------
# El catálogo (JSON de T2) y la unidad
# --------------------------------------------------------------------------


def test_f036_catalogo_de_otra_obra() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["obra"]["codigo"] = "0678"
    problemas = _detenida(catalogo=catalogo, grupos=None)
    assert problemas == ("el catálogo es de la obra «0678» y la tabla, de la «0677»",)


def test_f036_catalogo_obra_con_blancos_se_normaliza() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["obra"]["codigo"] = "0677  "
    assert _migrar(catalogo=catalogo).obra == "0677"


@pytest.mark.parametrize(
    "romper",
    [
        lambda c: c.pop("unidades"),
        lambda c: c.pop("oficios_obra"),
        lambda c: c["obra"].pop("codigo"),
        lambda c: c.update(obra="0677"),
        lambda c: c["oficios_obra"][0].pop("oficio_codigo"),
        lambda c: c["oficios_obra"][0].update(oficio_codigo=None),
        lambda c: c["unidades"].append("0677.04"),
        lambda c: c.update(obra={"codigo": 677, "nombre": None}),
        # R128: `oficios_catalogo`, una lista de `{codigo, nombre}` (§10.5).
        lambda c: c.pop("oficios_catalogo"),
        lambda c: c.update(oficios_catalogo=None),
        lambda c: c.update(oficios_catalogo="Pintura"),
        lambda c: c.update(oficios_catalogo={"codigo": "0901", "nombre": "X"}),
        lambda c: c["oficios_catalogo"].append("Pintura"),
        lambda c: c["oficios_catalogo"].append(["0901", "X"]),
        lambda c: c["oficios_catalogo"].append({"codigo": "0901"}),
        lambda c: c["oficios_catalogo"].append({"nombre": "X"}),
        lambda c: c["oficios_catalogo"].append({"codigo": "0901", "nombre": 5}),
    ],
)
def test_f036_catalogo_mal_formado(romper: Callable[[dict], None]) -> None:
    catalogo = copy.deepcopy(CATALOGO)
    romper(catalogo)
    problemas = _detenida(catalogo=catalogo, grupos=None)
    assert problemas == (
        "el JSON del catálogo no tiene la forma de T2 (`design.md` §10.3)",
    )


def test_f036_catalogo_que_no_es_un_mapping() -> None:
    problemas = _detenida(catalogo=[], grupos=None)
    assert problemas == (
        "el JSON del catálogo no tiene la forma de T2 (`design.md` §10.3)",
    )


def test_f036_catalogo_sin_unidades_para_con_el_motivo_del_dominio() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["unidades"] = []
    problemas = _detenida(catalogo=catalogo)
    assert problemas == (
        "ninguna obra con el código «0677» tiene unidades de posventa en Sigrid",
    )


def test_f036_tabla_con_codigo_de_obra_invalido() -> None:
    problemas = _detenida(correcciones=_correcciones(lambda t: t.update(obra="06/77")))
    assert problemas == ("la tabla de correcciones: el código de obra no es válido",)


def test_f036_unidad_que_no_es_una_etiqueta_de_la_obra() -> None:
    problemas = _detenida(
        correcciones=_correcciones(lambda t: t.update(unidad="Villa 1"))
    )
    assert problemas == (
        (
            "la unidad «Villa 1» no es la etiqueta de ninguna unidad de la obra en el "
            "desplegable de la plantilla"
        ),
    )


def test_f036_unidad_de_otra_villa_de_la_obra_vale() -> None:
    m = _migrar(
        correcciones=_correcciones(lambda t: t.update(unidad="Villa 2 inventada"))
    )
    assert {c["Unidad"] for c in _celdas_v2(m.v2)} == {"Villa 2 inventada"}


def test_f036_opciones_que_chocan_paran() -> None:
    catalogo = copy.deepcopy(CATALOGO)
    catalogo["oficios_obra"].append(
        {
            "oficio_codigo": "0061",
            "oficio_nombre": "Fontanería (0060)",
            "proveedor_codigo": None,
            "proveedor_nombre": None,
        }
    )
    problemas = _detenida(catalogo=catalogo, grupos=None)
    assert problemas == (
        (
            "las opciones de la plantilla no se pueden componer: dos opciones del "
            "desplegable dan la misma etiqueta"
        ),
    )


# --------------------------------------------------------------------------
# R56 · la ida y vuelta, y lo que para
# --------------------------------------------------------------------------


def test_f036_r56_errores_del_importador_paran_con_la_fila_del_original() -> None:
    def romper(tabla: dict) -> None:
        tabla["filas"][1]["resultado"][0]["ubicacion"] = "SALA/ ESTUDIO"
        tabla["filas"][3]["resultado"][1]["descripcion"] = "x" * 129

    problemas = _detenida(correcciones=_correcciones(romper))
    assert problemas == (
        (
            "fila 2 del original (fila 3 del v2), Ubicación: ese valor no está en la "
            "lista de ubicaciones: elígelo del desplegable (cuidado con mayúsculas, "
            "tildes y espacios)"
        ),
        (
            "fila 4 del original (fila 6 del v2), Descripción corta: la descripción "
            "corta tiene 129 caracteres y el máximo es 128: deja lo esencial y pasa el "
            "resto a «Detalle»"
        ),
    )


class _LectorFalso:
    def __init__(self, libro: LibroLeido | None = None, error: Exception | None = None):
        self.libro = libro
        self.error = error

    def leer(self, *, contenido: bytes) -> LibroLeido:
        if self.error is not None:
            raise self.error
        assert self.libro is not None
        return self.libro


def _libro_leido(m: mig.Migracion) -> LibroLeido:
    return LectorPlantillaOpenpyxl().leer(contenido=m.v2)


def test_f036_r56_el_v2_de_otra_obra_para() -> None:
    libro = replace(_libro_leido(_migrar()), obra_codigo="0678")
    problemas = _detenida(lector=_LectorFalso(libro))
    assert problemas == ("el v2 se reconoce como la plantilla de otra obra («0678»)",)


def test_f036_r56_el_v2_con_filas_de_menos_para() -> None:
    libro = _libro_leido(_migrar())
    libro = replace(libro, filas=libro.filas[:-1])
    problemas = _detenida(lector=_LectorFalso(libro))
    assert problemas == (
        "el v2 trae 6 filas válidas al volver a leerlo y se escribieron 7",
    )


@pytest.mark.parametrize(
    "error",
    [
        FicheroNoEsPlantilla("no_es_xlsx", "El fichero no es un Excel."),
        FicheroDemasiadoGrande("El fichero pasa de 2 MiB."),
    ],
)
def test_f036_r56_el_v2_que_el_lector_rechaza_para(error: Exception) -> None:
    problemas = _detenida(lector=_LectorFalso(error=error))
    assert problemas == (f"el v2 no pasa el lector del importador: {error.motivo}",)


def test_f036_r56_el_v2_que_no_se_reconoce_para() -> None:
    libro = replace(_libro_leido(_migrar()), identificador="otra cosa")
    problemas = _detenida(lector=_LectorFalso(libro))
    assert problemas[0].startswith("el v2 no pasa el lector del importador: ")


def test_f036_r56_el_generador_que_falla_para() -> None:
    class Generador:
        def generar(self, **kwargs) -> bytes:
            raise ValueError("no caben más de 1000 filas en la plantilla")

    problemas = _detenida(generador=Generador())
    assert problemas == (
        "el v2 no se puede generar: no caben más de 1000 filas en la plantilla",
    )


def test_f036_r56_avisos_del_importador_sin_los_del_proveedor() -> None:
    m = _migrar()
    assert [len(a) for a in m.avisos] == [0, 0, 1, 0, 0, 1, 0]
    assert m.avisos[2] == (
        (
            "el oficio «Solados y Alicatados» tiene varios códigos en la obra: se "
            "elegirá en la bandeja"
        ),
    )
    assert mig.avisos_sin_proveedor(
        ("el proveedor «Proveedor Inventado Uno» tiene varios códigos", "otro")
    ) == ("otro",)


# --------------------------------------------------------------------------
# R58 · el informe
# --------------------------------------------------------------------------


def _informe(**kwargs) -> str:
    return mig.informe(_migrar(**kwargs), sha256_original="abc123")


#: Los códigos de proveedor del `obrofc` del catálogo inventado (R116).
CODIGOS_DE_PROVEEDOR = tuple(
    sorted(
        {
            f["proveedor_codigo"]
            for f in CATALOGO["oficios_obra"]
            if f["proveedor_codigo"] is not None
        }
    )
)


def test_f036_r58_el_informe_no_lleva_nombres_de_proveedor() -> None:
    for grupos in (GRUPOS, None):
        texto = _informe(grupos=grupos)
        for nombre in NOMBRES_DE_PROVEEDOR:
            assert nombre not in texto
        assert "Inventado" not in texto
        # R58 pide además «sin datos personales», que precisa R116: ni códigos.
        for codigo in CODIGOS_DE_PROVEEDOR:
            assert codigo not in texto


def test_f036_r116_el_informe_no_contiene_ningun_codigo_de_proveedor() -> None:
    assert CODIGOS_DE_PROVEEDOR == ("P001", "P002", "P003", "P004", "P006")
    for grupos in (GRUPOS, None):
        m = _migrar(grupos=grupos)
        # El v2 sí lleva los proveedores completados: es la plantilla (R91).
        assert any(c.proveedor_codigo for c in m.compuestas)
        texto = mig.informe(m, sha256_original="abc123")
        for codigo in CODIGOS_DE_PROVEEDOR:
            assert codigo not in texto


def test_f036_r116_el_informe_marca_los_proveedores_completados_sin_codigo() -> None:
    marca = "Proveedor: proveedor completado (único de la obra para ese oficio) · "
    assert _informe().count(marca) == 2
    assert _informe(grupos=None).count(marca) == 3
    assert "Sin nombres ni códigos de proveedor (R58, R116)" in _informe()


def test_f036_r58_cabecera_y_recuentos() -> None:
    texto = _informe()
    assert texto.startswith("<!-- progress/migracion_F-036.md -->\n")
    assert "`sha256` del original: `abc123`" in texto
    assert "Obra `0677` · unidad «Villa 1 inventada»" in texto
    assert "fichero de grupos vigentes (2 grupos de varios códigos)" in texto
    assert "Ida y vuelta por el importador: **0 errores**" in texto
    for linea in (
        "| Filas del original | 7 |",
        "| Filas nuevas | 7 |",
        "| Filas del original separadas en varias | 1 |",
        "| Filas del original descartadas | 1 |",
        "| Oficios completados (la muestra no lo traía) | 2 |",
        "| Oficios de la muestra que no son exactamente un nombre de Sigrid de la obra | 1 |",
        "| Filas nuevas con un oficio cuyo grupo tiene varios códigos en la obra | 3 |",
        "| Filas nuevas con proveedor, completado por R91 (el único de la obra para ese oficio) | 2 |",
        "| Urgencias marcadas: Urgente | 0 |",
        "| Urgencias marcadas: Peligro para la seguridad | 1 |",
        "| Filas nuevas con algún aviso del importador | 2 |",
    ):
        assert linea in texto


def test_f036_r58_recuentos_sin_grupos() -> None:
    m = _migrar(grupos=None)
    assert m.recuentos == mig.Recuentos(
        filas_originales=7,
        filas_nuevas=7,
        separadas=1,
        descartadas=1,
        oficios_completados=2,
        oficios_no_exactos=1,
        oficios_en_grupo_de_varios=0,
        oficios_fuera_de_la_obra=0,
        proveedores_completados=3,
        urgencias={"urgente": 0, "seguridad": 1},
        con_avisos=0,
    )
    texto = mig.informe(m, sha256_original="x")
    assert "sin fichero de grupos: cada código es su propio grupo" in texto


def test_f036_r58_una_tabla_por_ubicacion_original() -> None:
    texto = _informe()
    titulos = [l for l in texto.splitlines() if l.startswith("## Ubicación original")]
    assert titulos == [
        "## Ubicación original «SALA/ ESTUDIO»",
        "## Ubicación original «TERRAZA»",
        "## Ubicación original «GENERALES»",
        "## Ubicación original «COCINA»",
        "## Ubicación original «Cocina»",
    ]


def test_f036_r58_cada_fila_con_su_propuesta_cambios_y_propuesto() -> None:
    texto = _informe()
    lineas = texto.splitlines()
    fila3 = next(l for l in lineas if l.startswith("| 3 |"))
    assert "Enlechar azulejos, PELIGRO DE SEGURIDAD." in fila3
    assert "Solados y alicatados" in fila3  # el oficio de la muestra
    assert (
        "Oficio: Solados y Alicatados M.O. (código `0133`; en el v2, «Solados y "
        "Alicatados», grupo con 2 códigos en la obra)"
    ) in fila3
    assert "Urgencia: Peligro para la seguridad" in fila3
    assert "ubicación normalizada<br>peligro de seguridad a Urgencia" in fila3
    assert fila3.endswith("| oficio, ubicacion |")
    assert "Aviso: el oficio «Solados y Alicatados» tiene varios códigos" in fila3
    fila4 = next(l for l in lineas if l.startswith("| 4 |"))
    assert "1. Ubicación: Terraza 1 · Descripción: Sellar encuentro" in fila4
    assert "<br>2. Ubicación: Terraza 1 · Descripción: Puerta suelta" in fila4
    assert fila4.endswith("| oficio, proveedor |")
    fila5 = next(l for l in lineas if l.startswith("| 5 |"))
    assert "descartada" in fila5 and "repite la fila 2" in fila5
    fila1 = next(l for l in lineas if l.startswith("| 1 |"))
    assert "Oficio: Pintura (código `0010`)" in fila1
    assert fila1.endswith("| proveedor |")
    fila7 = next(l for l in lineas if l.startswith("| 7 |"))
    assert fila7.endswith("| — | — |")


def test_f036_r58_el_oficio_dice_la_etiqueta_del_v2_cuando_no_es_su_nombre() -> None:
    con_grupos = _informe()
    assert (
        "Oficio: Solados y Alicatados (código `0033`; en el v2, «Solados y "
        "Alicatados», grupo con 2 códigos en la obra)"
    ) in con_grupos
    sin_grupos = _informe(grupos=None)
    assert (
        "Oficio: Fontanería (código `0060`; en el v2, «Fontanería (0060)»)"
        in sin_grupos
    )
    assert "Oficio: Solados y Alicatados M.O. (código `0133`) ·" in sin_grupos


def test_f036_r58_el_informe_escapa_lo_que_romperia_la_tabla() -> None:
    original = list(ORIGINAL)
    original[6] = (None, "Cocina", "A | B <x>\nC & D", "Mamparas")

    def cambiar(tabla: dict) -> None:
        tabla["filas"][6]["texto_original"] = "A | B <x> C & D"
        tabla["filas"][6]["resultado"][0]["detalle"] = "uno\ndos"
        tabla["filas"][6]["cambios"] = ["texto"]

    texto = _informe(
        original=_original(tuple(original)), correcciones=_correcciones(cambiar)
    )
    fila7 = next(l for l in texto.splitlines() if l.startswith("| 7 |"))
    assert "A \\| B &lt;x&gt;<br>C &amp; D" in fila7
    assert "Detalle: uno<br>dos" in fila7


def test_f036_r58_el_nombre_del_informe_va_en_su_cabecera() -> None:
    texto = mig.informe(_migrar(), sha256_original="x", nombre="otro.md")
    assert texto.startswith("<!-- progress/otro.md -->\n")


def test_f036_r58_listado_y_urgencia_por_codigo_salen_con_su_etiqueta() -> None:
    def marcar(tabla: dict) -> None:
        tabla["filas"][6]["resultado"][0].update(listado="segundo", urgencia="urgente")
        tabla["filas"][6]["cambios"] = ["listado y urgencia"]

    m = _migrar(correcciones=_correcciones(marcar))
    ultima = _celdas_v2(m.v2)[-1]
    assert (ultima["Listado"], ultima["Urgencia"]) == ("Segundo listado", "Urgente")
    fila7 = next(
        linea
        for linea in mig.informe(m, sha256_original="x").splitlines()
        if linea.startswith("| 7 |")
    )
    assert "Urgencia: Urgente · Listado: Segundo listado" in fila7
    assert "| Urgencias marcadas: Urgente | 1 |" in mig.informe(m, sha256_original="x")


# --------------------------------------------------------------------------
# La línea de órdenes: R53 y R97, con un disco en memoria
# --------------------------------------------------------------------------


class DiscoEnMemoria:
    """Un disco de mentira: ningún `.xlsx` toca el disco de verdad (R59)."""

    def __init__(self, ficheros: dict[str, bytes]):
        self.ficheros = {Path(r).resolve(): d for r, d in ficheros.items()}
        self.escritos: list[Path] = []
        self.leidos: list[Path] = []
        self.al_escribir: Callable[[DiscoEnMemoria], None] | None = None

    def leer(self, ruta: Path) -> bytes:
        self.leidos.append(ruta)
        if ruta not in self.ficheros:
            raise FileNotFoundError(str(ruta))
        return self.ficheros[ruta]

    def escribir(self, ruta: Path, datos: bytes, *, sobrescribir: bool) -> None:
        if not sobrescribir and ruta in self.ficheros:
            raise FileExistsError(str(ruta))
        self.ficheros[ruta] = datos
        self.escritos.append(ruta)
        if self.al_escribir is not None:
            self.al_escribir(self)

    def existe(self, ruta: Path) -> bool:
        return ruta in self.ficheros

    def de(self, ruta: str) -> bytes:
        return self.ficheros[Path(ruta).resolve()]


ORIGINAL_RUTA = "entrada/creacion_incidencias.xlsx"
V2_RUTA = "salida/creacion_incidencias_v2.xlsx"
INFORME_RUTA = "salida/migracion_F-036.md"


def _disco(
    *, grupos: object = GRUPOS, correcciones: object = None, catalogo: object = CATALOGO
) -> DiscoEnMemoria:
    ficheros = {
        ORIGINAL_RUTA: _original(),
        "entrada/catalogo.json": json.dumps(catalogo).encode("utf-8"),
        "entrada/correcciones.yaml": yaml.safe_dump(
            _correcciones() if correcciones is None else correcciones,
            allow_unicode=True,
        ).encode("utf-8"),
    }
    if grupos is not None:
        ficheros["entrada/grupos.json"] = json.dumps(grupos).encode("utf-8")
    return DiscoEnMemoria(ficheros)


def _argv(
    *extra: str, grupos: bool = True, salida: str = V2_RUTA, informe: str = INFORME_RUTA
) -> list[str]:
    argv = [
        "--original",
        ORIGINAL_RUTA,
        "--catalogo",
        "entrada/catalogo.json",
        "--correcciones",
        "entrada/correcciones.yaml",
        "--salida",
        salida,
        "--informe",
        informe,
    ]
    if grupos:
        argv += ["--grupos", "entrada/grupos.json"]
    return argv + list(extra)


def test_f036_r53_r56_linea_de_ordenes_escribe_v2_e_informe(capsys) -> None:
    disco = _disco()
    bytes_antes = disco.de(ORIGINAL_RUTA)
    antes = hashlib.sha256(bytes_antes).hexdigest()
    assert cli.main(_argv(), disco=disco) == 0
    assert disco.escritos == [Path(V2_RUTA).resolve(), Path(INFORME_RUTA).resolve()]
    assert disco.de(ORIGINAL_RUTA) == bytes_antes
    libro = LectorPlantillaOpenpyxl().leer(contenido=disco.de(V2_RUTA))
    assert reconocer_plantilla(libro) == "0677"
    informe = disco.de(INFORME_RUTA).decode("utf-8")
    assert f"`sha256` del original: `{antes}`" in informe
    salida = capsys.readouterr().out
    assert "Ida y vuelta por el importador: 0 errores" in salida
    assert f"sha256 del original antes:   {antes}" in salida
    assert f"sha256 del original después: {antes}" in salida
    assert "Filas del original: 7 · filas nuevas: 7" in salida
    assert "Proveedor Inventado" not in salida


def test_f036_r97_solo_informe_sin_grupos_no_escribe_el_v2(capsys) -> None:
    disco = _disco(grupos=None)
    assert cli.main(_argv("--solo-informe", grupos=False), disco=disco) == 0
    assert disco.escritos == [Path(INFORME_RUTA).resolve()]
    salida = capsys.readouterr().out
    assert "sin --grupos: cada código es su propio grupo" in salida
    assert "v2: no se escribe (--solo-informe)" in salida


def test_f036_r97_solo_informe_con_grupos_tampoco_escribe_el_v2() -> None:
    disco = _disco()
    assert cli.main(_argv("--solo-informe"), disco=disco) == 0
    assert disco.escritos == [Path(INFORME_RUTA).resolve()]


def test_f036_r97_sin_grupos_y_sin_solo_informe_se_niega_sin_leer_nada(capsys) -> None:
    disco = _disco(grupos=None)
    assert cli.main(_argv(grupos=False), disco=disco) == 2
    assert disco.leidos == [] and disco.escritos == []
    assert "sin --grupos solo se admite --solo-informe (R97)" in capsys.readouterr().err


@pytest.mark.parametrize(
    "salida",
    [
        ORIGINAL_RUTA,
        "entrada/../entrada/creacion_incidencias.xlsx",
        "./" + ORIGINAL_RUTA,
    ],
)
@pytest.mark.parametrize("sobrescribir", [(), ("--sobrescribir",)])
def test_f036_r53_salida_igual_al_original_se_niega(
    salida: str, sobrescribir: tuple[str, ...], capsys
) -> None:
    disco = _disco()
    assert cli.main(_argv(*sobrescribir, salida=salida), disco=disco) == 1
    assert disco.escritos == [] and disco.leidos == []
    assert (
        "--salida es el original: se niega a escribir encima (R53)"
        in capsys.readouterr().err
    )


def test_f036_r53_informe_igual_al_original_se_niega(capsys) -> None:
    disco = _disco()
    assert cli.main(_argv(informe=ORIGINAL_RUTA), disco=disco) == 1
    assert disco.escritos == []
    assert (
        "--informe es el original: se niega a escribir encima (R53)"
        in capsys.readouterr().err
    )


def test_f036_r53_informe_igual_a_la_salida_se_niega(capsys) -> None:
    disco = _disco()
    assert cli.main(_argv(informe=V2_RUTA), disco=disco) == 1
    assert disco.escritos == []
    assert "--informe y --salida son el mismo fichero" in capsys.readouterr().err


def test_f036_r53_salida_existente_sin_sobrescribir_se_niega(capsys) -> None:
    disco = _disco()
    disco.ficheros[Path(V2_RUTA).resolve()] = b"v2 anterior"
    assert cli.main(_argv(), disco=disco) == 1
    assert disco.escritos == [] and disco.leidos == []
    assert disco.de(V2_RUTA) == b"v2 anterior"
    assert (
        "--salida ya existe: pasa --sobrescribir para reemplazarla (R53)"
        in capsys.readouterr().err
    )


def test_f036_r53_salida_existente_tambien_con_solo_informe() -> None:
    disco = _disco()
    disco.ficheros[Path(V2_RUTA).resolve()] = b"v2 anterior"
    assert cli.main(_argv("--solo-informe"), disco=disco) == 1
    assert disco.escritos == []


def test_f036_r53_con_sobrescribir_se_reemplaza() -> None:
    disco = _disco()
    disco.ficheros[Path(V2_RUTA).resolve()] = b"v2 anterior"
    assert cli.main(_argv("--sobrescribir"), disco=disco) == 0
    assert disco.de(V2_RUTA).startswith(b"PK")


def test_f036_r53_el_informe_se_reescribe_siempre() -> None:
    disco = _disco()
    disco.ficheros[Path(INFORME_RUTA).resolve()] = b"informe anterior"
    assert cli.main(_argv(), disco=disco) == 0
    assert disco.de(INFORME_RUTA).startswith(b"<!-- progress/migracion_F-036.md -->")


def test_f036_r53_original_que_cambia_durante_la_ejecucion_falla(capsys) -> None:
    disco = _disco()

    def tocar(d: DiscoEnMemoria) -> None:
        d.ficheros[Path(ORIGINAL_RUTA).resolve()] = b"otro contenido"

    disco.al_escribir = tocar
    assert cli.main(_argv(), disco=disco) == 1
    err = capsys.readouterr().err
    assert "el original ha cambiado durante la migración" in err


def test_f036_r53_r54_si_para_no_se_escribe_nada(capsys) -> None:
    tabla = _correcciones(lambda t: t["filas"].pop(0))
    disco = _disco(correcciones=tabla)
    assert cli.main(_argv(), disco=disco) == 1
    assert disco.escritos == []
    err = capsys.readouterr().err
    assert "ERROR: la migración se ha parado sin escribir nada:" in err
    assert "  - falta la fila 1 del original en la tabla de correcciones" in err


@pytest.mark.parametrize(
    ("fichero", "contenido", "problema"),
    [
        (
            "entrada/catalogo.json",
            b"{no es json",
            "el JSON del catálogo no se puede leer",
        ),
        ("entrada/catalogo.json", b"\xff\xfe", "el JSON del catálogo no se puede leer"),
        ("entrada/grupos.json", b"[", "el fichero de grupos no se puede leer"),
        (
            "entrada/correcciones.yaml",
            b"a: [",
            "la tabla de correcciones no se puede leer",
        ),
        (
            "entrada/correcciones.yaml",
            b"\xff",
            "la tabla de correcciones no se puede leer",
        ),
    ],
)
def test_f036_ficheros_que_no_se_pueden_leer(
    fichero: str, contenido: bytes, problema: str, capsys
) -> None:
    disco = _disco()
    disco.ficheros[Path(fichero).resolve()] = contenido
    assert cli.main(_argv(), disco=disco) == 1
    assert disco.escritos == []
    assert problema in capsys.readouterr().err


def test_f036_catalogo_con_bom_se_lee() -> None:
    disco = _disco()
    ruta = Path("entrada/catalogo.json").resolve()
    disco.ficheros[ruta] = b"\xef\xbb\xbf" + disco.ficheros[ruta]
    assert cli.main(_argv(), disco=disco) == 0


def test_f036_fichero_que_falta(capsys) -> None:
    disco = _disco()
    del disco.ficheros[Path(ORIGINAL_RUTA).resolve()]
    assert cli.main(_argv(), disco=disco) == 1
    assert disco.escritos == []
    assert "ERROR: no se puede leer o escribir un fichero:" in capsys.readouterr().err


def test_f036_argumentos_obligatorios() -> None:
    for falta in (
        "--original",
        "--catalogo",
        "--correcciones",
        "--salida",
        "--informe",
    ):
        argv = _argv()
        posicion = argv.index(falta)
        del argv[posicion : posicion + 2]
        with pytest.raises(SystemExit) as salida:
            cli.main(argv, disco=_disco())
        assert salida.value.code == 2


def test_f036_sha256() -> None:
    assert cli.sha256(b"abc") == hashlib.sha256(b"abc").hexdigest()


def test_f036_disco_local(tmp_path: Path) -> None:
    """El disco de verdad, con un fichero de texto (nunca un `.xlsx`)."""
    disco = cli.DiscoLocal()
    ruta = tmp_path / "prueba.bin"
    assert disco.existe(ruta) is False
    disco.escribir(ruta, b"uno", sobrescribir=False)
    assert disco.existe(ruta) is True
    assert disco.leer(ruta) == b"uno"
    with pytest.raises(FileExistsError):
        disco.escribir(ruta, b"dos", sobrescribir=False)
    assert disco.leer(ruta) == b"uno"
    disco.escribir(ruta, b"dos", sobrescribir=True)
    assert disco.leer(ruta) == b"dos"


def test_f036_se_lanza_como_script_desde_cualquier_sitio(tmp_path: Path) -> None:
    resultado = subprocess.run(
        [sys.executable, str(SERVICIO / "scripts" / "migrar_excel_f036.py"), "--help"],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    assert resultado.returncode == 0, resultado.stderr
    assert "--solo-informe" in resultado.stdout
    assert list(tmp_path.iterdir()) == []


def test_f036_resultados_inmutables() -> None:
    m = _migrar()
    for objeto in (
        m,
        m.recuentos,
        m.compuestas[0],
        m.originales[0],
        m.tabla,
        m.tabla.filas[0],
        m.compuestas[0].nueva,
    ):
        with pytest.raises(AttributeError):
            objeto.__setattr__("x", 1)


# --------------------------------------------------------------------------
# R97, R128–R131 · un catálogo de Sigrid que ha cambiado (undécima enmienda,
# `design.md` §10.5)
# --------------------------------------------------------------------------


def _poner_oficio(*cambios: tuple[int, int, str | None]) -> Callable[[dict], None]:
    """En la entrada `i` de la tabla, fila nueva `j` (las dos desde 0), ese oficio."""

    def poner(tabla: dict) -> None:
        for i, j, nombre in cambios:
            tabla["filas"][i]["resultado"][j]["oficio"] = nombre

    return poner


#: La fila 2 del original («Albañilería») con un oficio que ha salido de la obra.
SALIO_EN_LA_2 = _poner_oficio((1, 0, FUERA_DE_LA_OBRA))

FILA_DE_RECUENTO = (
    "| Filas nuevas sin oficio porque ese oficio ya no está en la obra en Sigrid | {} |"
)
SECCION = "## Filas sin oficio porque ese oficio ya no está en la obra en Sigrid"
LINEA_DE_ORDENES = "Filas sin oficio porque ese oficio ya no está en la obra en Sigrid: {}"


def _catalogo_viejo() -> dict:
    """El de «antes»: el oficio que sale aún en la obra, con un único proveedor."""
    viejo = copy.deepcopy(CATALOGO_SIGRID)
    viejo["oficios_obra"].append(
        {
            "oficio_codigo": "0901",
            "oficio_nombre": FUERA_DE_LA_OBRA,
            "proveedor_codigo": "P007",
            "proveedor_nombre": "Proveedor Inventado Siete",
            "marca_cif": None,
        }
    )
    return viejo


def _catalogo_nuevo() -> dict:
    """El de «hoy»: ese oficio fuera de la obra y `Pintura` con otro proveedor único."""
    nuevo = copy.deepcopy(CATALOGO_SIGRID)
    nuevo["oficios_obra"][0].update(
        proveedor_codigo="P009", proveedor_nombre="Proveedor Inventado Nueve"
    )
    return nuevo


def _migrar_salio(**kwargs) -> mig.Migracion:
    kwargs.setdefault("catalogo", copy.deepcopy(CATALOGO_SIGRID))
    return _migrar(correcciones=_correcciones(SALIO_EN_LA_2), **kwargs)


def _validar_contra(v2: bytes, catalogo: dict) -> tuple:
    """El v2 leído y validado contra **otro** catálogo, como hace el entorno."""
    cat = mig.catalogo_de_la_obra(catalogo, "0677")
    opciones = opciones_de_la_obra(
        cat, mig.EquivalenciasDeFichero(mig.decisiones_de_grupos(GRUPOS, "0677"))
    )
    libro = LectorPlantillaOpenpyxl().leer(contenido=v2)
    return validar_filas(
        libro.filas,
        catalogo=cat,
        opciones_oficio=opciones.oficios,
        opciones_proveedor=opciones.proveedores,
        listas=CONFIG.listas,
    )


def test_f036_r128_oficio_que_ya_no_esta_en_la_obra_queda_vacio() -> None:
    m = _migrar_salio()
    salio = m.compuestas[1]
    assert salio.origen == 2
    assert salio.oficio_fuera_de_la_obra == FUERA_DE_LA_OBRA
    assert (
        salio.oficio_codigo,
        salio.oficio_etiqueta,
        salio.proveedor_codigo,
        salio.codigos_en_obra,
    ) == (None, None, None, 0)
    assert salio.fila.valores[OFICIO] is None
    assert salio.fila.valores[PROVEEDOR] is None
    # El resto de la fila no cambia (R128).
    assert salio.fila.valores[DESCRIPCION] == "Rematar agujero en pared"
    assert salio.fila.valores[UBICACION] == "Sala/estudio"
    # Las demás filas, igual que con `CATALOGO` y la tabla de siempre.
    de_siempre = _migrar()
    assert len(m.compuestas) == len(de_siempre.compuestas) == 7
    for i, (una, otra) in enumerate(
        zip(m.compuestas, de_siempre.compuestas, strict=True)
    ):
        if i != 1:
            assert una == otra
            assert una.oficio_fuera_de_la_obra is None


def test_f036_r128_el_v2_vuelve_con_cero_errores_y_la_fila_sin_oficio() -> None:
    m = _migrar_salio()
    validas, errores = _validar_contra(m.v2, CATALOGO_SIGRID)
    assert errores == ()
    assert len(validas) == 7
    assert (validas[1].oficio, validas[1].proveedor) == (None, None)
    celdas = _celdas_v2(m.v2)
    assert celdas[1]["Descripción corta"] == "Rematar agujero en pared"
    assert (celdas[1]["Oficio"], celdas[1]["Proveedor"]) == (None, None)
    assert not any(
        FUERA_DE_LA_OBRA in str(valor)
        for fila in celdas
        for valor in fila.values()
        if valor is not None
    )


def test_f036_r128_el_nombre_de_sigrid_se_compara_recortado() -> None:
    # «Oficio retirado   » en Sigrid casa con «Oficio retirado» en la tabla.
    m = _migrar(
        catalogo=copy.deepcopy(CATALOGO_SIGRID),
        correcciones=_correcciones(_poner_oficio((1, 0, RETIRADO))),
    )
    assert m.compuestas[1].oficio_fuera_de_la_obra == RETIRADO
    # La tabla, en cambio, se compara tal cual: blancos, capitalización,
    # tildes, puntuación o forma Unicode, errata (R128). Las tildes importan:
    # en `auxofc` hay gemelos que solo difieren en ellas (review 10, R10-1).
    # La puntuación y la forma Unicode, por lo mismo (review 11, R11-1 y
    # R11-2): «Solados y Alicatados M.O.» es de la obra en `CATALOGO`, y la
    # forma NFD de un nombre no es el nombre.
    for nombre in (
        FUERA_DE_LA_OBRA + " ",
        " " + FUERA_DE_LA_OBRA,
        FUERA_DE_LA_OBRA.lower(),
        RETIRADO + "   ",
        "Oficio que salio",
        "Albañileria",
        "Solados y Alicatados MO",
        unicodedata.normalize("NFD", FUERA_DE_LA_OBRA),
    ):
        problemas = _detenida(
            catalogo=copy.deepcopy(CATALOGO_SIGRID),
            correcciones=_correcciones(_poner_oficio((1, 0, nombre))),
        )
        assert problemas == (_errata(2, 1, nombre),)


def test_f036_r128_nombres_de_sigrid_recortados_sin_nulos_ni_vacios() -> None:
    datos = {
        "oficios_catalogo": [
            {"codigo": "1", "nombre": " Uno  "},
            {"codigo": "2", "nombre": None},
            {"codigo": "3", "nombre": "   "},
            {"codigo": "4", "nombre": ""},
            {"codigo": "5", "nombre": "Dos"},
        ]
    }
    assert mig.nombres_de_sigrid(datos) == frozenset({"Uno", "Dos"})
    assert mig.nombres_de_sigrid({"oficios_catalogo": []}) == frozenset()
    for crudo in ([], "texto", None):
        with pytest.raises(mig.MigracionDetenida) as error:
            mig.nombres_de_sigrid(crudo)
        assert error.value.problemas == (
            "el JSON del catálogo no tiene la forma de T2 (`design.md` §10.3)",
        )


def test_f036_r128_la_obra_manda() -> None:
    con_sigrid = _migrar(catalogo=copy.deepcopy(CATALOGO_SIGRID))
    de_siempre = _migrar()
    assert con_sigrid.compuestas == de_siempre.compuestas
    assert [c.oficio_codigo for c in con_sigrid.compuestas] == [
        "0010",
        "0020",
        "0133",
        "0033",
        None,
        "0060",
        "0144",
    ]
    assert con_sigrid.compuestas[0].proveedor_codigo == "P001"
    assert all(c.oficio_fuera_de_la_obra is None for c in con_sigrid.compuestas)
    assert con_sigrid.recuentos == de_siempre.recuentos


def test_f036_r128_nombre_de_varios_oficios_de_la_obra_sigue_parando() -> None:
    catalogo = copy.deepcopy(CATALOGO_SIGRID)
    catalogo["oficios_obra"][-1]["oficio_nombre"] = "Fontanería"
    problemas = _detenida(catalogo=catalogo)
    assert problemas == (
        (
            "fila 6 del original, fila nueva 1: el oficio «Fontanería» es el nombre de "
            "varios oficios de la obra en Sigrid (0060, 0160): no se puede elegir"
        ),
    )


def test_f036_r97_errata_y_fuera_de_la_obra_a_la_vez() -> None:
    problemas = _detenida(
        catalogo=copy.deepcopy(CATALOGO_SIGRID),
        correcciones=_correcciones(
            _poner_oficio((0, 0, "Errata inventada"), (1, 0, FUERA_DE_LA_OBRA))
        ),
    )
    assert problemas == (_errata(1, 1, "Errata inventada"),)


def test_f036_r129_el_proveedor_es_el_del_catalogo_con_que_se_ejecuta() -> None:
    """Invariante: nace en verde (el código de antes ya lo cumple, R129)."""
    uno = _migrar(catalogo=copy.deepcopy(CATALOGO_SIGRID))
    nueve = _migrar(catalogo=_catalogo_nuevo())
    # `migrar` ya ha exigido 0 errores en la ida y vuelta de los dos (R56).
    assert _celdas_v2(uno.v2)[0]["Proveedor"] == "Pintura · Proveedor Inventado Uno"
    assert _celdas_v2(nueve.v2)[0]["Proveedor"] == "Pintura · Proveedor Inventado Nueve"
    assert (
        uno.compuestas[0].proveedor_codigo,
        nueve.compuestas[0].proveedor_codigo,
    ) == ("P001", "P009")
    dos = copy.deepcopy(CATALOGO_SIGRID)
    dos["oficios_obra"].insert(
        1,
        {
            "oficio_codigo": "0010",
            "oficio_nombre": "Pintura",
            "proveedor_codigo": "P009",
            "proveedor_nombre": "Proveedor Inventado Nueve",
            "marca_cif": None,
        },
    )
    m = _migrar(catalogo=dos)
    assert _celdas_v2(m.v2)[0]["Proveedor"] is None
    assert m.compuestas[0].proveedor_codigo is None


def test_f036_r129_la_fila_sin_oficio_no_lleva_proveedor() -> None:
    antes = _migrar_salio(catalogo=_catalogo_viejo())
    assert antes.compuestas[1].proveedor_codigo == "P007"
    assert _celdas_v2(antes.v2)[1]["Proveedor"] == (
        f"{FUERA_DE_LA_OBRA} · Proveedor Inventado Siete"
    )
    hoy = _migrar_salio(catalogo=_catalogo_nuevo())
    assert hoy.compuestas[1].proveedor_codigo is None
    assert hoy.compuestas[1].fila.valores[PROVEEDOR] is None
    assert _celdas_v2(hoy.v2)[1]["Proveedor"] is None
    texto = mig.informe(hoy, sha256_original="x")
    assert "P007" not in texto
    assert "Siete" not in texto


def test_f036_r130_con_el_catalogo_de_siempre_recuento_cero_y_sin_seccion() -> None:
    for catalogo in (CATALOGO, CATALOGO_SIGRID):
        m = _migrar(catalogo=copy.deepcopy(catalogo))
        assert m.recuentos.oficios_fuera_de_la_obra == 0
        texto = mig.informe(m, sha256_original="x")
        lineas = texto.splitlines()
        # Siempre, detrás de la de grupos de varios códigos.
        i = lineas.index(
            "| Filas nuevas con un oficio cuyo grupo tiene varios códigos en la obra | 3 |"
        )
        assert lineas[i + 1] == FILA_DE_RECUENTO.format(0)
        assert SECCION not in texto
        assert "ya no está en la obra en Sigrid:" not in texto


def test_f036_r130_recuento_y_seccion() -> None:
    m = _migrar_salio()
    assert m.recuentos.oficios_fuera_de_la_obra == 1
    texto = mig.informe(m, sha256_original="x")
    assert FILA_DE_RECUENTO.format(1) in texto
    lineas = texto.splitlines()
    i = next(
        n
        for n, linea in enumerate(lineas)
        if linea.startswith("| Filas nuevas con algún aviso del importador |")
    )
    # Justo después de los recuentos.
    assert lineas[i + 1 : i + 8] == [
        "",
        SECCION,
        "",
        "| Fila del original | Fila nueva | Oficio en la tabla de correcciones |",
        "|---:|---:|---|",
        f"| 2 | 1 | {FUERA_DE_LA_OBRA} |",
        "",
    ]
    assert lineas[i + 8].startswith("## Ubicación original")
    fila2 = next(l for l in lineas if "Rematar agujero en pared" in l)
    assert (
        f"Oficio: — (ya no está en la obra en Sigrid: «{FUERA_DE_LA_OBRA}») · " in fila2
    )
    # La fila a la que la tabla no pone oficio no cuenta ni cambia su texto.
    assert m.compuestas[4].nueva.oficio is None
    assert m.compuestas[4].oficio_fuera_de_la_obra is None
    assert "Descripción: Puerta suelta · Oficio: — · Urgencia" in texto
    # Ni los códigos de `auxofc` que no son de la obra (R116, §10.5).
    assert "0901" not in texto and "0902" not in texto


def test_f036_r130_varias_filas_y_su_numero_dentro_de_la_original() -> None:
    def cambiar(tabla: dict) -> None:
        _poner_oficio((1, 0, FUERA_DE_LA_OBRA), (3, 1, RETIRADO))(tabla)
        tabla["filas"][3]["resultado"].append(
            _nueva("Marco rozado", ubicacion="Terraza 1")
        )

    m = _migrar(
        catalogo=copy.deepcopy(CATALOGO_SIGRID), correcciones=_correcciones(cambiar)
    )
    assert [c.oficio_fuera_de_la_obra for c in m.compuestas] == [
        None,
        FUERA_DE_LA_OBRA,
        None,
        None,
        RETIRADO,
        None,  # «Marco rozado»: la tabla no le pone oficio, no cuenta
        None,
        None,
    ]
    assert m.recuentos.oficios_fuera_de_la_obra == 2
    texto = mig.informe(m, sha256_original="x")
    assert FILA_DE_RECUENTO.format(2) in texto
    assert (
        f"| 2 | 1 | {FUERA_DE_LA_OBRA} |\n| 4 | 2 | {RETIRADO} |\n\n## Ubicación"
        in texto
    )
    assert f"Oficio: — (ya no está en la obra en Sigrid: «{RETIRADO}»)" in texto


def test_f036_r130_el_nombre_de_la_seccion_se_escapa() -> None:
    raro = "Oficio | <raro>"
    catalogo = copy.deepcopy(CATALOGO_SIGRID)
    catalogo["oficios_catalogo"].append({"codigo": "0903", "nombre": raro})
    m = _migrar(
        catalogo=catalogo, correcciones=_correcciones(_poner_oficio((1, 0, raro)))
    )
    texto = mig.informe(m, sha256_original="x")
    assert "| 2 | 1 | Oficio \\| &lt;raro&gt; |" in texto


def test_f036_r130_la_linea_de_ordenes_lo_dice(capsys) -> None:
    assert cli.main(_argv(), disco=_disco()) == 0
    assert LINEA_DE_ORDENES.format(0) in capsys.readouterr().out
    disco = _disco(
        catalogo=_catalogo_nuevo(), correcciones=_correcciones(SALIO_EN_LA_2)
    )
    assert cli.main(_argv(), disco=disco) == 0
    salida = capsys.readouterr().out
    lineas = salida.splitlines()
    i = next(n for n, l in enumerate(lineas) if l.startswith("Filas del original:"))
    assert lineas[i + 1] == LINEA_DE_ORDENES.format(1)
    assert "Proveedor Inventado" not in salida


def test_f036_r130_sin_proveedores_en_el_informe(capsys) -> None:
    catalogos = (_catalogo_viejo(), _catalogo_nuevo())
    proveedores = {
        (f["proveedor_codigo"], f["proveedor_nombre"])
        for c in catalogos
        for f in c["oficios_obra"]
        if f["proveedor_codigo"] is not None
    }
    assert ("P007", "Proveedor Inventado Siete") in proveedores
    assert ("P009", "Proveedor Inventado Nueve") in proveedores
    for catalogo in catalogos:
        disco = _disco(catalogo=catalogo, correcciones=_correcciones(SALIO_EN_LA_2))
        assert cli.main(_argv(), disco=disco) == 0
        for texto in (disco.de(INFORME_RUTA).decode("utf-8"), capsys.readouterr().out):
            assert "Inventado" not in texto
            for codigo, nombre in proveedores:
                assert codigo not in texto
                assert nombre not in texto


def test_f036_r131_un_v2_de_otro_catalogo_da_errores_y_el_regenerado_no() -> None:
    """El incidente de T29 en pequeño (R131): el control es de procedimiento."""
    viejo = _migrar_salio(catalogo=_catalogo_viejo())
    _, con_error = _validar_contra(viejo.v2, _catalogo_nuevo())
    filas = {e.fila for fila in con_error for e in fila.errores}
    # Justo la de `Pintura` (su par ya no es de la obra) y la del oficio que salió.
    assert filas == {PRIMERA_FILA, PRIMERA_FILA + 1}
    regenerado = _migrar_salio(catalogo=_catalogo_nuevo())
    validas, con_error = _validar_contra(regenerado.v2, _catalogo_nuevo())
    assert con_error == ()
    assert len(validas) == 7
